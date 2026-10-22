"""Spec 024, T13: metodología del PIDE sin DIgSILENT (mercado, diagnóstico, plan, formatos)."""
from __future__ import annotations

import random

import pytest
from openpyxl import load_workbook

from igea_dgs import pide

LIM = {'tension_min_pu': 0.95, 'tension_max_pu': 1.05, 'cargabilidad_max_pct': 100,
       'transformador_max_pct': 80}


def _lectura(v_mt=0.97, v_bt=0.96, linea=60.0, trafo=50.0, sin=(), fp=None, perdidas=20.0):
    return {
        'topologia': {'barras': {'B_MT': {'nombre': 'B_MT', 'kv': 10.0},
                                 'B_BT': {'nombre': 'B_BT', 'kv': 0.38}}},
        'flujo': {
            'balance': {'unidad': 'kW', 'valores': {'cabeceras': {'F1': 1000.0}},
                        'fp_cabeceras': {'F1': fp} if fp is not None else {}},
            'perdidas': {'unidad': 'kW', 'valores': {'total': perdidas}},
            'cargabilidad': {'unidad': '%', 'valores': {'lineas': {'Red\\L1': linea},
                                                        'transformadores': {'Red\\T1': trafo}}},
            'tensiones': {'unidad': 'p.u.', 'valores': {'B_MT': v_mt, 'B_BT': v_bt},
                          'sin_tension': list(sin)},
            'confiabilidad': {'ausente': 'x'}},
    }


def test_crecimiento_y_proyeccion_a_20_anios():
    """024:RF-60 — tasa compuesta de la historia y proyección del año 00 al 20."""
    assert pide.tasa_de_crecimiento([100, 110, 121]) == pytest.approx(0.10)
    serie = pide.proyectar(5.0, 0.03)
    assert len(serie) == 21 and serie[0] == 5.0 and serie[20] == pytest.approx(5.0 * 1.03 ** 20)
    with pytest.raises(pide.PideError):
        pide.tasa_de_crecimiento([100])


def test_densidad_por_cuadriculas_de_la_tabla_2():
    """024:RF-60 — cuadrículas UTM alineadas a múltiplos del lado, demanda total y
    concentrada en MW y zona de la Tabla N.º 2."""
    puntos = [(400_100, 8_400_200, 3000, False), (400_900, 8_400_900, 1500, True),
              (401_100, 8_400_100, 100, False)]
    celdas = pide.cuadriculas(puntos, 1000)
    a = celdas[(400_000, 8_400_000)]
    assert a['mw'] == pytest.approx(4.5) and a['mw_concentradas'] == pytest.approx(1.5)
    assert a['zona'] == 'Muy Alta' and a['color'] == 'Rojo'
    assert celdas[(401_000, 8_400_000)]['zona'] == 'Baja'
    assert pide.zona_de_densidad(2.0) == ('Alta 2', 'Azul')
    with pytest.raises(pide.PideError):
        pide.cuadriculas(puntos, 300)


@pytest.mark.parametrize(('lectura', 'causa'), [
    (_lectura(v_mt=0.90), 'caida_de_tension'),
    (_lectura(v_mt=0.90, fp=0.85), 'factor_de_potencia'),
    (_lectura(v_bt=0.90, trafo=95.0), 'transformador_sobrecargado'),
    (_lectura(linea=120.0), 'conductor_insuficiente'),
    (_lectura(sin=['x', 'y']), 'isla_sin_fuente'),
    (_lectura(v_mt=1.08), 'sobretension'),
    (_lectura(perdidas=120.0), 'perdidas_sobre_reconocidas'),
])
def test_diagnostico_con_su_causa(lectura, causa):
    """024:RF-49 / 024:RF-50 / 024:RF-61 — cada problema con su ubicación y la causa que
    sostiene la evidencia del flujo."""
    problemas = pide.diagnosticar(lectura, perdidas_reconocidas_pct=8.0, fp_minimo=0.95, **LIM)
    assert causa in {p.causa for p in problemas}
    assert all(p.elemento and p.evidencia for p in problemas)


def test_red_que_cumple_no_tiene_problemas():
    """024:RF-49 — sin incumplimientos, el diagnóstico está vacío."""
    assert pide.diagnosticar(_lectura(), **LIM) == []


def test_anio_de_incumplimiento_con_la_demanda():
    """024:RF-51 — el primer año en que la caída o la carga pasan el límite."""
    assert pide.anio_de_incumplimiento(0.97, 0.95, 'tension', 0.05) == 11
    assert pide.anio_de_incumplimiento(70.0, 100, 'carga', 0.04) == 10
    assert pide.anio_de_incumplimiento(70.0, 100, 'carga', 0.0) is None
    assert pide.anio_de_incumplimiento(0.94, 0.95, 'tension', 0.03) == 0


def test_problemas_futuros_con_su_anio():
    """024:RF-51 — lo que hoy cumple y deja de cumplir dentro de los 20 años."""
    futuros = pide.problemas_futuros(_lectura(v_mt=0.97, linea=70.0), tasa=0.05, horizonte=20,
                                     tension_min_pu=0.95, cargabilidad_max_pct=100,
                                     transformador_max_pct=80)
    anios = {(p.tipo, p.elemento): p.anio for p in futuros}
    assert anios[('tension', 'B_MT')] == 11 and anios[('cargabilidad', 'L1')] == 8


def test_plan_por_etapas_sin_adelantar_obras():
    """024:RF-52 — una obra por causa en el año en que hace falta, la de menor costo entre
    las que la resuelven; lo de hoy entra en el primer año del PIDE."""
    problemas = [pide.Problema('tension', 'B1', 0.93, 0.95, 'caida_de_tension', 'e', 0),
                 pide.Problema('tension', 'B2', 0.94, 0.95, 'caida_de_tension', 'e', 0),
                 pide.Problema('cargabilidad', 'L7', 90, 100, 'conductor_insuficiente', 'e', 9)]
    obras = pide.plan_por_etapas(problemas, anios_pide=[4, 5, 6, 7])
    assert [(o.anio, o.alternativa) for o in obras] == [(4, 'bancos'), (9, 'cambio_seccion')]
    assert len(obras[0].problemas) == 2 and obras[0].rubro == 'mejoramiento'
    barato = pide.plan_por_etapas(problemas, anios_pide=[4], costo_relativo={'reguladores': 1, 'bancos': 5})
    assert barato[0].alternativa == 'reguladores'


def test_sorteo_de_redes_el_plan_cubre_cada_problema():
    """024:RF-52 — con problemas sorteados, cada causa resoluble tiene su obra y ninguna
    obra va antes del año de su problema."""
    semilla = random.randrange(1 << 30)
    rnd = random.Random(semilla)
    causas = list(pide.ALTERNATIVAS_POR_CAUSA)
    problemas = [pide.Problema('cargabilidad', f'E{i}', 1, 1, rnd.choice(causas), 'e', rnd.randint(0, 20))
                 for i in range(rnd.randint(1, 40))]
    obras = pide.plan_por_etapas(problemas, anios_pide=[4, 5, 6, 7])
    for p in problemas:
        o = next(o for o in obras if o.causa == p.causa and any(x.startswith(p.elemento + ':') for x in o.problemas))
        assert o.anio <= max(p.anio, 4), f'semilla {semilla}'


def test_formatos_01_a_15(tmp_path):
    """024:RF-63 — un libro con los quince formatos; el que no tiene datos lo dice."""
    ruta = pide.escribir_formatos(tmp_path / 'pide.xlsx',
                                  {'F01 Proyección': [['S1', 0, 100.0, 5.0, 3.0]]})
    libro = load_workbook(ruta)
    assert len(libro.sheetnames) == 15
    assert libro['F01 Proyección'].cell(2, 1).value == 'S1'
    assert libro['F09 Evaluación económica'].cell(2, 1).value == 'sin datos en esta ejecución'
