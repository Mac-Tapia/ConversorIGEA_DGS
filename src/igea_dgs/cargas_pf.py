"""Cargas en un proyecto existente de DIgSILENT: lectura de los libros y plan (spec 004).

Este módulo es aparte de la conversión. Antes, el plan se armaba con el modelo del TXT
o la MDB cargado en la web: sin esas entradas no se podía actualizar ni una demanda, y
la SED «no existía» si el proyecto venía de otra conversión. Ahora la web solo valida
el Excel; qué SED coincide y a qué barra se conecta una carga nueva lo decide el guion
``tools/cargas_pf.py`` mirando el propio proyecto de DIgSILENT (spec 004, RF-1).

Una fila con error no detiene las demás: se informa y el resto sigue. Una SED que no
coincide en el proyecto se omite allí, no aquí (RF-7).
"""

from __future__ import annotations

import csv
import math
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

from .loads import (
    LoadTemplateError, SUPPORTED_SUFFIXES, _normalise, _number, _power_conflict,
    _resolve_power, canonical_columns, read_sheet,
)

SCHEMA = 'igea-dgs-cargas-pf-v1'

DEMANDA_COLUMNS = ('SED', 'Kw', 'Kvar', '(kVA)', 'FP', 'accion', 'observaciones')
NUEVA_COLUMNS = (
    'SED', 'alimentador', 'barra', 'Lat', 'Lon', 'CoordX', 'CoordY', 'kVA_instalado',
    'Kw', 'Kvar', '(kVA)', 'FP', 'tipo_linea', 'tension_bt_kv', 'observaciones',
)

# Columnas propias de la carga nueva, con sus variantes de grafía (cabecera normalizada).
NUEVA_ALIASES = {
    'alimentador': 'alimentador', 'feeder': 'alimentador',
    'barra': 'barra', 'nodo': 'barra', 'nodoconexion': 'barra', 'nodeid': 'barra',
    'terminal': 'barra',
    'lat': 'Lat', 'latitud': 'Lat', 'latitude': 'Lat', 'gpslat': 'Lat',
    'lon': 'Lon', 'long': 'Lon', 'longitud': 'Lon', 'longitude': 'Lon', 'gpslon': 'Lon',
    'coordx': 'CoordX', 'x': 'CoordX', 'este': 'CoordX', 'easting': 'CoordX',
    'coordy': 'CoordY', 'y': 'CoordY', 'norte': 'CoordY', 'northing': 'CoordY',
    'kvainstalado': 'kVA_instalado', 'kvainst': 'kVA_instalado', 'trafo': 'kVA_instalado',
    'tipolinea': 'tipo_linea', 'conductor': 'tipo_linea', 'cable': 'tipo_linea',
    'tensionbtkv': 'tension_bt_kv', 'tensionbt': 'tension_bt_kv', 'btkv': 'tension_bt_kv',
    'kvbt': 'tension_bt_kv',
    'observaciones': 'observaciones', 'obs': 'observaciones',
    # Libro de planificación (SI213, 2026-10-12): la carga no tiene SED aún y se nombra
    # por su descripción; el alimentador va como «TRONCAL».
    'descripcion': 'nombre', 'nombre': 'nombre', 'proyecto': 'nombre', 'cliente': 'nombre',
    'troncal': 'alimentador',
}

# Potencias normalizadas de transformador (IEC 60076-1, serie R10) para deducir el kVA
# instalado de una carga que solo trae su demanda (RF-52). Se repite ×10 hasta cubrirla.
SERIE_R10_KVA = (10, 12.5, 16, 20, 25, 31.5, 40, 50, 63, 80)


def _columna_por_patron(clave: str) -> str:
    """Nombre canónico de una cabecera que no es un alias exacto (RF-51).

    «DEM PROY (KW)» o «COORDENADAS UTM (X)» no se pueden listar todas: se reconoce la
    unidad o el eje al final. El kVA no, porque puede ser la demanda o el instalado.
    """
    if clave.endswith('kvar'):
        return 'Kvar'
    if clave.endswith('kw'):
        return 'Kw'
    if 'utm' in clave or clave.startswith('coord'):
        if clave.endswith(('x', 'este')):
            return 'CoordX'
        if clave.endswith(('y', 'norte')):
            return 'CoordY'
    return ''


def kva_normalizado(kva: float) -> float:
    """La menor potencia de la serie R10 que cubre ``kva`` (RF-52)."""
    escala = 1.0
    while True:
        for valor in SERIE_R10_KVA:
            if valor * escala >= kva - 1e-9:
                return valor * escala
        escala *= 10


@dataclass(frozen=True)
class Demanda:
    """Demanda pedida para una SED. ``hoja`` orienta si la SED está en dos alimentadores."""

    sed: str
    hoja: str
    kw: float
    kvar: float
    fp: float

    def para_pf(self) -> dict[str, Any]:
        return {
            'sed': self.sed, 'hoja': self.hoja, 'kw': self.kw, 'kvar': self.kvar,
            'fp': self.fp, 'plini_mw': self.kw / 1000.0, 'qlini_mvar': self.kvar / 1000.0,
        }


@dataclass(frozen=True)
class CargaNueva:
    """Carga nueva: su demanda y lo necesario para conectarla en el proyecto."""

    sed: str
    hoja: str
    alimentador: str
    barra: str
    lat: float | None
    lon: float | None
    kva_instalado: float
    kw: float
    kvar: float
    fp: float
    tipo_linea: str = ''
    tension_bt_kv: float | None = None

    def para_pf(self) -> dict[str, Any]:
        datos = asdict(self)
        datos.update(plini_mw=self.kw / 1000.0, qlini_mvar=self.kvar / 1000.0)
        return datos


@dataclass
class Lectura:
    filas: list = field(default_factory=list)
    errores: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)        # valores deducidos (RF-52, RF-53)
    omitidas: list[str] = field(default_factory=list)      # accion = omitir
    sin_datos: list[str] = field(default_factory=list)     # ni kW ni kVA/FP


# --------------------------------------------------------------------------- libros
def _vacia(fila: Sequence[Any] | None) -> bool:
    return not fila or not any(c is not None and str(c).strip() for c in fila)


def _hojas(path: Path | str) -> dict[str, tuple[list[Any], list[list[Any]], int]]:
    """``{hoja: (cabecera, filas, fila_cabecera)}`` de un .xlsx (una por hoja) o un .csv (una).

    La cabecera es la primera fila con algo, no la fila 1: el libro de planificación de
    Electro Dunas (NuevosclientesMayoresv1, 2026-10-14) la trae en la fila 3 tras dos
    vacías, y tomando la fila 1 no se reconocía ninguna columna y sus 58 cargas se
    descartaban sin un solo aviso. ``fila_cabecera`` es su número en Excel, para que los
    mensajes citen la fila que el operador ve.
    """
    src = Path(path)
    if not src.is_file():
        raise LoadTemplateError(f'No se encuentra el fichero: {src}')
    suffix = src.suffix.lower()
    if suffix == '.csv':
        with src.open('r', encoding='utf-8-sig', newline='') as fh:
            sample = fh.read(4096)
            fh.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
            except csv.Error:
                dialect = csv.excel
                dialect.delimiter = ';'
            reader = csv.reader(fh, dialect=dialect)
            fila_cabecera = 0
            header: list[Any] | None = None
            for header in reader:
                fila_cabecera += 1
                if not _vacia(header):
                    break
            if header is None or _vacia(header):
                raise LoadTemplateError(f'{src.name}: fichero vacío.')
            body = [row for row in reader if not _vacia(row)]
        return {src.stem: (header, body, fila_cabecera)}
    if suffix != '.xlsx':
        raise LoadTemplateError(
            f'Extensión no soportada: {src.suffix!r}. Use ' + ' o '.join(SUPPORTED_SUFFIXES))
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise LoadTemplateError(
            'Leer .xlsx necesita openpyxl. Instale con: pip install -r requirements.txt') from exc
    wb = load_workbook(src, data_only=True, read_only=True)
    out: dict[str, tuple[list[Any], list[list[Any]], int]] = {}
    try:
        for name in wb.sheetnames:
            it = wb[name].iter_rows(values_only=True)
            fila_cabecera = 0
            header = None
            for fila in it:
                fila_cabecera += 1
                if not _vacia(fila):
                    header = list(fila)
                    break
            if header is None:
                continue
            body = [list(v) for v in it if not _vacia(v)]
            if body:
                out[name] = (header, body, fila_cabecera)
    finally:
        wb.close()
    if not out:
        raise LoadTemplateError(f'{src.name}: ninguna hoja con datos.')
    return out


def leer_demandas(path: Path | str) -> Lectura:
    """Demandas por SED de todas las hojas. Sin modelo: valida cada fila por sí misma."""
    lectura = Lectura()
    vistas: dict[str, str] = {}
    for hoja, (header, body, _fila_cabecera) in _hojas(path).items():
        try:
            read = read_sheet(header, body, hoja)
        except LoadTemplateError as exc:
            lectura.errores.append(str(exc))
            continue
        lectura.errores += read.errors
        lectura.omitidas += read.skipped
        lectura.sin_datos += read.no_data
        for row in read.rows:
            if row.no_data:
                continue
            clave = f'{hoja}\x00{row.sed_code}'
            if clave in vistas:
                continue
            vistas[clave] = hoja
            lectura.filas.append(Demanda(row.sed_code, hoja, row.kw, row.kvar, row.fp))
    return lectura


def leer_nuevas(path: Path | str, *, source_crs: str = '') -> Lectura:
    """Cargas nuevas de todas las hojas, con su posición ya en WGS84.

    La posición puede venir como ``Lat``/``Lon`` o como ``CoordX``/``CoordY`` en el CRS
    del perfil de empresa (``source_crs``): DIgSILENT compara con el GPS de sus barras,
    así que el guion la recibe siempre en grados.
    """
    lectura = Lectura()
    for hoja, (header, body, fila_cabecera) in _hojas(path).items():
        filas, errores = leer_hoja_nuevas(header, body, hoja, source_crs=source_crs,
                                          avisos=lectura.avisos, fila_cabecera=fila_cabecera)
        lectura.filas += filas
        lectura.errores += errores
    return lectura


def leer_hoja_nuevas(
    header: Sequence[Any], body: Sequence[Sequence[Any]], hoja: str, *, source_crs: str = '',
    avisos: list[str] | None = None, fila_cabecera: int = 1,
) -> tuple[list[CargaNueva], list[str]]:
    """Filas de una hoja de cargas nuevas. Lo deducido se anota en ``avisos`` (RF-52, RF-53)."""
    avisos = [] if avisos is None else avisos
    mapping = dict(canonical_columns(header))
    for index, raw in enumerate(header):
        clave = _normalise(raw)
        if clave in NUEVA_ALIASES:
            mapping[index] = NUEVA_ALIASES[clave]
    for index, raw in enumerate(header):
        nombre = '' if index in mapping else _columna_por_patron(_normalise(raw))
        if nombre and nombre not in mapping.values():
            mapping[index] = nombre
    # Sin columna de nombre ni de demanda, cada fila se saltaría como «resto de la hoja»
    # y el libro entero daría cero cargas sin decir por qué.
    if body and not {'SED', 'nombre', 'Kw', '(kVA)'} & set(mapping.values()):
        cabecera = ', '.join(str(c).strip() for c in header if c is not None and str(c).strip())
        return [], [f'hoja {hoja}, fila {fila_cabecera}: no se reconoce la cabecera '
                    f'({cabecera or "vacía"}); hace falta SED o descripción, y Kw o (kVA).']
    filas: list[CargaNueva] = []
    errores: list[str] = []
    vistas: dict[str, int] = {}
    pendientes_xy: list[tuple[int, float, float]] = []
    for offset, values in enumerate(body):
        fila = offset + fila_cabecera + 1
        cell = {name: values[i] if i < len(values) else None for i, name in mapping.items()}
        texto = {k: ('' if v is None else str(v)).strip() for k, v in cell.items()}
        # RF-53: sin SED, la carga se nombra por su descripción y, sin ella, por su fila;
        # una fila sin nombre ni demanda es un resto de la hoja y se salta.
        sed = texto.get('SED') or texto.get('nombre', '')
        if not sed:
            if not any(texto.get(c) for c in ('Kw', 'Kvar', '(kVA)')):
                continue
            sed = f'{hoja}_F{fila}'
            avisos.append(f'hoja {hoja}, fila {fila}: sin SED ni descripción; se nombra {sed} (RF-53)')
        donde = f'hoja {hoja}, fila {fila}: {sed}'
        if sed in vistas:
            errores.append(f'{donde} repetida (fila {vistas[sed]})')
            continue
        vistas[sed] = fila
        numeros: dict[str, float | None] = {}
        malo = ''
        for nombre in ('Lat', 'Lon', 'CoordX', 'CoordY', 'kVA_instalado', 'Kw', 'Kvar',
                       '(kVA)', 'FP', 'tension_bt_kv'):
            valor = _number(cell.get(nombre))
            if valor is not None and math.isnan(valor):
                malo = f'{nombre}={cell.get(nombre)!r} no es un número'
                break
            if valor is not None and valor < 0 and nombre not in ('Lat', 'Lon', 'CoordX', 'CoordY'):
                malo = f'{nombre}={valor} negativo'
                break
            numeros[nombre] = valor
        if not malo and numeros.get('FP') is not None and not 0 <= numeros['FP'] <= 1:
            malo = f'FP={numeros["FP"]} fuera de [0, 1]'
        lat, lon = numeros.get('Lat'), numeros.get('Lon')
        if not malo and (lat is None) != (lon is None):
            malo = 'Lat y Lon van juntas'
        if not malo and lat is not None and not (-90 <= lat <= 90 and -180 <= lon <= 180):
            malo = f'Lat/Lon fuera de rango ({lat}, {lon})'
        x, y = numeros.get('CoordX'), numeros.get('CoordY')
        tiene_xy = x is not None and y is not None
        if not malo and lat is None and tiene_xy and not source_crs:
            malo = 'CoordX/CoordY sin CRS del perfil: use Lat/Lon o configure el CRS'
        if not malo and not texto.get('barra') and lat is None and not tiene_xy:
            malo = 'sin barra ni coordenadas: indique la barra o Lat/Lon (o CoordX/CoordY)'
        if malo:
            errores.append(f'{donde}: {malo}')
            continue
        kw, kvar, _kva, fp, _derivada, sin_datos = _resolve_power(
            numeros.get('Kw'), numeros.get('Kvar'), numeros.get('(kVA)'), numeros.get('FP'))
        conflicto = _power_conflict(numeros.get('(kVA)'), numeros.get('FP'), kw, kvar)
        if conflicto:
            errores.append(f'{donde}: {conflicto}. Deje un solo par: Kw y Kvar, Kw y FP, o (kVA) y FP.')
            continue
        if sin_datos:
            errores.append(f'{donde}: sin demanda (Kw, o (kVA) y FP)')
            continue
        kva_instalado = numeros.get('kVA_instalado')
        if not kva_instalado:
            # RF-52: un libro de planificación da la demanda, no el transformador; se toma
            # el normalizado que la cubre y se anota, en vez de rechazar la carga.
            kva_instalado = kva_normalizado(math.hypot(kw, kvar))
            avisos.append(f'{donde}: sin kVA_instalado; se toma {kva_instalado:g} kVA, el '
                          f'normalizado que cubre {math.hypot(kw, kvar):.1f} kVA de demanda (RF-52)')
        if lat is None and tiene_xy:
            pendientes_xy.append((len(filas), x, y))
        filas.append(CargaNueva(
            sed=sed, hoja=hoja, alimentador=texto.get('alimentador', ''),
            barra=texto.get('barra', ''), lat=lat, lon=lon,
            kva_instalado=float(kva_instalado), kw=kw, kvar=kvar, fp=fp,
            tipo_linea=texto.get('tipo_linea', ''), tension_bt_kv=numeros.get('tension_bt_kv'),
        ))
    if pendientes_xy:
        gps = a_wgs84([(x, y) for _i, x, y in pendientes_xy], source_crs)
        for (index, _x, _y), punto in zip(pendientes_xy, gps):
            fila = filas[index]
            if punto is None:
                if not fila.barra:
                    errores.append(f'hoja {hoja}: {fila.sed}: CoordX/CoordY no se pudieron '
                                   f'pasar a GPS con {source_crs}')
                    filas[index] = None  # type: ignore[call-overload]
                continue
            filas[index] = CargaNueva(**{**asdict(fila), 'lat': punto[0], 'lon': punto[1]})
    return [f for f in filas if f is not None], errores


def a_wgs84(puntos: Sequence[tuple[float, float]], source_crs: str) -> list[tuple[float, float] | None]:
    """(x, y) en ``source_crs`` → (lat, lon). ``None`` si un punto no se puede pasar."""
    from .geography import make_transformer

    try:
        transformer = make_transformer(source_crs, 'EPSG:4326')
    except ImportError:
        return [None] * len(puntos)
    out: list[tuple[float, float] | None] = []
    for x, y in puntos:
        try:
            lon, lat = transformer.transform(x, y)
        except Exception:  # noqa: BLE001 - un punto malo no tumba el libro
            out.append(None)
            continue
        ok = math.isfinite(lat) and math.isfinite(lon) and -90 <= lat <= 90 and -180 <= lon <= 180
        out.append((lat, lon) if ok else None)
    return out


# --------------------------------------------------------------------------- plan
def armar_plan(
    *,
    project: str,
    feeders: Sequence[str],
    demandas: Lectura | None,
    nuevas: Lectura | None,
    costos: dict | None = None,
    tec: dict | None = None,
    ahora: datetime | None = None,
    carpeta: str | None = None,
    base: dict | None = None,
    minima_m: float | None = None,
) -> dict[str, Any]:
    """El plan que recibe ``tools/cargas_pf.py``. Igual para uno que para varios (RF-5).

    ``project`` es la ruta del proyecto dentro de ``carpeta``, la de importación de
    DIgSILENT (spec 005, RF-16): dos proyectos con el mismo nombre en subcarpetas
    distintas son dos proyectos.

    ``base`` es el ajuste a la máxima sobre el que se integra (spec 007, RF-1): sus
    variaciones, una por alimentador ajustado, forman con las demandas del envío el
    estado sin proyecto. La lee la web del trabajo del ajuste (``services.base_del_ajuste``).

    ``minima_m`` es la distancia mínima de la spec 016 (RF-20, RF-73): una carga cuyo punto
    de la troncal queda más cerca de una barra se engancha en ella sin partir el tramo.
    """
    if minima_m is not None and (not isinstance(minima_m, (int, float)) or minima_m < 0):
        raise LoadTemplateError(f'La distancia mínima tiene que ser un número de metros, 0 o más: '
                                f'{minima_m!r}.')
    elegidos = list(dict.fromkeys(f.strip() for f in feeders if f and f.strip()))
    if not project.strip():
        raise LoadTemplateError('Elija el proyecto de DIgSILENT.')
    if not elegidos:
        raise LoadTemplateError('Elija al menos un alimentador.')
    if demandas is None and nuevas is None:
        raise LoadTemplateError('Suba el libro de demandas, el de cargas nuevas o los dos.')
    stamp = (ahora or datetime.now()).strftime('%Y%m%d_%H%M%S')
    errores = list((demandas.errores if demandas else []) + (nuevas.errores if nuevas else []))
    plan = {
        'schema': SCHEMA,
        'project': project.strip(),
        'folder': carpeta,
        'feeders': elegidos,
        'stamp': stamp,
        # El guion añade un sufijo si ya existe (spec 004, RF-4). Las demandas van en su
        # propia variación y cada carga nueva en la suya, que el guion nombra con su SED y
        # su alimentador cuando sabe a cuál se conecta (spec 007, RF-9, RF-10, RF-19).
        'variacion': f'Demandas_{stamp}',
        'variacion_demandas': f'Demandas_{stamp}',
        'demandas': [d.para_pf() for d in (demandas.filas if demandas else [])],
        'nuevas': [n.para_pf() for n in (nuevas.filas if nuevas else [])],
        'omitidas_por_el_operador': list(demandas.omitidas) if demandas else [],
        'sin_datos': list(demandas.sin_datos) if demandas else [],
        'row_errors': errores,
        'costos': costos,
        'tec': tec,
        'minima_m': minima_m,
    }
    plan['aplicable'] = bool(plan['demandas'] or plan['nuevas'])
    plan['base'] = base or None
    plan['avisos'] = _avisos_de_la_base(base) + list(nuevas.avisos if nuevas else [])
    plan['calibradas_cambiadas'] = _calibradas_cambiadas(base, plan['demandas'])
    return plan


def _avisos_de_la_base(base: dict | None) -> list[str]:
    """Lo que el operador debe saber de la base antes de aplicar (spec 007, RF-5, RF-34, RF-37)."""
    if not base:
        return ['Sin base del ajuste: se aplica sobre las variaciones que estén activas en el '
                'proyecto (spec 007, RF-5).']
    avisos = []
    for a in base.get('ajustes') or []:
        if not a.get('cuadrado'):
            desvio = (a.get('desvio') or {}).get('kw')
            error = f' (error en la cabecera {desvio:+.2%})' if isinstance(desvio, (int, float)) else ''
            avisos.append(f"El ajuste de {a['feeder']} no cuadró{error}: se usa igual como base (RF-34).")
    for x in base.get('excluidos') or []:
        avisos.append(f"{x['feeder']} queda fuera de la base: {x['motivo']} (RF-37).")
    return avisos


def _calibradas_cambiadas(base: dict | None, demandas: list[dict]) -> list[dict]:
    """SED que el ajuste calibró y que el libro de demandas cambia (spec 007, RF-20, RF-21)."""
    calibradas = (base or {}).get('sed') or {}
    out = []
    for d in demandas:
        c = calibradas.get(d['sed'])
        if c is not None:
            out.append({'sed': d['sed'], 'feeder': c.get('feeder'),
                        'calibrada': {'kw': c.get('kw'), 'kvar': c.get('kvar')},
                        'libro': {'kw': d['kw'], 'kvar': d['kvar']}})
    return out


# --------------------------------------------------------------------------- plantillas
def escribir_plantilla(tipo: str, destino: Path) -> Path:
    """Plantilla vacía de demandas (``tipo='demandas'``) o de cargas nuevas (``'nuevas'``)."""
    columnas = {'demandas': DEMANDA_COLUMNS, 'nuevas': NUEVA_COLUMNS}[tipo]
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    if destino.suffix.lower() == '.csv':
        with destino.open('w', encoding='utf-8-sig', newline='') as fh:
            csv.writer(fh, delimiter=';').writerow(columnas)
        return destino
    from openpyxl import Workbook

    wb = Workbook()
    hoja = wb.active
    hoja.title = 'cargas'
    hoja.append(list(columnas))
    wb.save(destino)
    return destino
