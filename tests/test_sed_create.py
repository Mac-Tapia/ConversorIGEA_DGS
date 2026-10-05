"""Creación de SED desde coordenadas: nodo más cercano y conductor de la derivación.

El operador da **dónde está** la carga, no a qué nodo se conecta. El módulo:

1. busca el **nodo más cercano** del modelo (las coordenadas están en el CRS proyectado
   del export, en metros, así que la distancia euclidiana ya es metros);
2. usa esa distancia como **longitud de la derivación**;
3. elige el **conductor aéreo** con dos criterios: ampacidad con margen, y caída de
   tensión por debajo del límite.

Un resultado medido que conviene tener presente y que estas pruebas fijan: para una SED
de distribución (15–630 kVA) la corriente es de 0,4 a 16 A a 22,9 kV, y el conductor más
pequeño de un catálogo real ya es de 112 A. **Ninguno de los dos criterios discrimina**,
y sale siempre el menor del catálogo. Lo que decide en la práctica es la normalización
de la empresa, y para eso está la columna ``conductor``.
"""

from __future__ import annotations

import math

import pytest

from igea_dgs.loads_create import (
    AMPACITY_MARGIN,
    CREATE_COLUMNS,
    LoadTemplateError,
    build_create_plan,
    create_plan_to_payload,
    find_nearest_node,
    read_create_sheet,
    select_conductor,
    single_new_load,
    write_create_template,
)
from igea_dgs.model import FeederModel, LineType, Load, Node, Sed


def _model(nominal_kv: float = 22.9) -> FeederModel:
    """Tres nodos en línea, separados 100 m, y un catálogo de tres conductores."""
    nodes = {
        'N1': Node('N1', 500000.0, 8500000.0),
        'N2': Node('N2', 500100.0, 8500000.0),
        'N3': Node('N3', 500200.0, 8500000.0),
    }
    types = {
        'LINE:CHICO': LineType('LINE:CHICO', 'CHICO', 'LINE', 1.0, 1.5, 0.45, 1.3, 0, 0, 100.0),
        'LINE:MEDIO': LineType('LINE:MEDIO', 'MEDIO', 'LINE', 0.5, 0.8, 0.42, 1.2, 0, 0, 200.0),
        'LINE:GRANDE': LineType('LINE:GRANDE', 'GRANDE', 'LINE', 0.2, 0.4, 0.38, 1.1, 0, 0, 400.0),
        # DEFAULT es el comodín del catálogo, no un conductor: no debe elegirse solo.
        'LINE:DEFAULT': LineType('LINE:DEFAULT', 'DEFAULT', 'LINE', 0.1, 0.2, 0.2, 0.6, 0, 0, 1000.0),
        # Subterráneo: tampoco debe salir en una derivación aérea.
        'CABLE:SUB': LineType('CABLE:SUB', 'SUB', 'CONCENTRIC NEUTRAL CABLE', 0.3, 0.5, 0.1, 0.3, 0, 0, 150.0),
    }
    sed = Sed(code='SE_EXISTE', loc_name='SE_EXISTE', node_id='N2', design_kva=100.0,
              section_id='SEC_1', device_number='DEV_1', load_key=('SEC_1', 'DEV_1'))
    load = Load('SEC_1', 'DEV_1', 'CUST', '1', 'N2', 0.01, 0.003, 0.95, 100.0, 0.0, 'ABC',
                sed_code='SE_EXISTE', display_name='SE_EXISTE')
    return FeederModel(
        name='AL01', network_id='NET_AL01', nominal_kv=nominal_kv, source_node='N1',
        nodes=nodes, lines=[], loads=[load], devices=[], line_types=types, seds=[sed],
    )


class TestNearestNode:
    def test_exact_node_is_at_zero_distance(self):
        node, distance = find_nearest_node(_model(), 500100.0, 8500000.0)
        assert node == 'N2'
        assert distance == pytest.approx(0.0)

    def test_closest_of_several_is_chosen(self):
        node, distance = find_nearest_node(_model(), 500180.0, 8500000.0)
        assert node == 'N3'
        assert distance == pytest.approx(20.0)

    def test_distance_is_metres_because_the_crs_is_projected(self):
        """3-4-5: la distancia euclidiana en un CRS métrico ya son metros."""
        _node, distance = find_nearest_node(_model(), 500100.0 + 30.0, 8500000.0 + 40.0)
        assert distance == pytest.approx(50.0)

    def test_model_without_coordinates_is_reported(self):
        model = _model()
        model.nodes = {'N1': Node('N1', None, None)}
        with pytest.raises(LoadTemplateError, match='ningún nodo'):
            find_nearest_node(model, 1.0, 2.0)


class TestConductorSelection:
    def test_small_sed_takes_the_smallest_conductor(self):
        """El caso real: ni ampacidad ni caída discriminan."""
        choice = select_conductor(_model(), kva=160, length_m=100.0, fp=0.95)
        assert choice.code == 'CHICO'
        assert choice.binding == 'minimo'
        assert choice.current_a == pytest.approx(160 / (math.sqrt(3) * 22.9))
        assert choice.voltage_drop_pct < 0.1

    def test_ampacity_binds_for_a_very_large_load(self):
        # 20 MVA a 22,9 kV son ~504 A: solo el mayor del catálogo (400 A) se acerca.
        choice = select_conductor(_model(), kva=20000, length_m=100.0, fp=0.95)
        assert choice.code == 'GRANDE'
        assert choice.binding == 'ampacidad'
        assert 'ningún conductor' in choice.reason

    def test_ampacity_margin_is_applied(self):
        """Se exige ampacidad ≥ I × margen, no solo ≥ I."""
        model = _model()
        # Corriente que deja CHICO (100 A) justo por debajo con el margen.
        kva = 100.0 / AMPACITY_MARGIN * math.sqrt(3) * model.nominal_kv * 1.01
        choice = select_conductor(model, kva=kva, length_m=10.0, fp=0.95)
        assert choice.code != 'CHICO'

    def test_voltage_drop_binds_on_a_long_derivation(self):
        """Con una caída máxima muy exigente, gana el de menor resistencia."""
        choice = select_conductor(
            _model(), kva=2000, length_m=3000.0, fp=0.95, max_drop_pct=0.9,
        )
        assert choice.code in ('MEDIO', 'GRANDE')
        assert choice.voltage_drop_pct <= 0.9

    def test_default_is_never_chosen_on_its_own(self):
        """DEFAULT es el comodín del catálogo, no una sección real."""
        for kva in (15, 100, 630, 5000):
            assert select_conductor(_model(), kva=kva, length_m=100.0).code != 'DEFAULT'

    def test_underground_types_are_not_used_for_an_overhead_derivation(self):
        for kva in (15, 630):
            assert select_conductor(_model(), kva=kva, length_m=100.0).code != 'SUB'

    def test_forced_conductor_wins_without_discussion(self):
        choice = select_conductor(_model(), kva=15, length_m=100.0, forced_code='GRANDE')
        assert choice.code == 'GRANDE'
        assert choice.binding == 'impuesto'

    def test_unknown_forced_conductor_is_reported(self):
        with pytest.raises(LoadTemplateError, match='no está entre los tipos aéreos'):
            select_conductor(_model(), kva=15, length_m=100.0, forced_code='NO_EXISTE')

    def test_choice_carries_the_numbers_that_justify_it(self):
        choice = select_conductor(_model(), kva=400, length_m=250.0, fp=0.9)
        assert choice.length_m == pytest.approx(250.0)
        assert choice.current_a > 0
        assert choice.reason and choice.binding
        # Sin longitud no hay caída: la fórmula debe respetarlo.
        assert select_conductor(_model(), kva=400, length_m=0.0).voltage_drop_pct == 0.0


class TestCreateFromCoordinates:
    def test_point_resolves_to_node_distance_and_conductor(self):
        rows, errors = single_new_load(
            sed_code='SE_NUEVA', coord_x=500130.0, coord_y=8500040.0,
            installed_kva=160, kva=45, fp=0.95, feeder='AL01',
        )
        plan = build_create_plan(_model(), rows, errors)
        assert plan.is_applicable
        item = plan.create[0]
        assert item.node_id == 'N2'
        assert item.distance_m == pytest.approx(50.0)
        assert item.nearest_from_coords is True
        assert item.new_node_id == 'NODE_SE_NUEVA'
        assert item.new_section_id == 'SEC_SE_NUEVA'
        assert item.conductor.code == 'CHICO'

    def test_explicit_node_overrides_the_nearest(self):
        """Si el operador da nodo y coordenadas, manda el nodo."""
        rows, errors = single_new_load(
            sed_code='SE_NUEVA', node_id='N3', coord_x=500100.0, coord_y=8500000.0,
            installed_kva=100, kw=10, feeder='AL01',
        )
        plan = build_create_plan(_model(), rows, errors)
        item = plan.create[0]
        assert item.node_id == 'N3'
        assert item.nearest_from_coords is False
        assert item.distance_m == pytest.approx(100.0), 'la longitud es al nodo elegido'

    def test_a_far_away_point_is_blocked_as_a_coordinate_error(self):
        rows, errors = single_new_load(
            sed_code='SE_LEJOS', coord_x=550000.0, coord_y=8500000.0,
            installed_kva=100, kw=10,
        )
        plan = build_create_plan(_model(), rows, errors)
        assert plan.is_applicable is False
        assert 'más de' in plan.row_errors[0]
        assert 'CRS' in plan.row_errors[0]

    def test_neither_point_nor_node_is_reported(self):
        rows, errors = single_new_load(sed_code='SE_X', installed_kva=100, kw=10)
        plan = build_create_plan(_model(), rows, errors)
        assert plan.is_applicable is False
        assert any('CoordX' in e for e in errors + plan.row_errors)

    def test_existing_sed_goes_to_already_exists(self):
        rows, errors = single_new_load(
            sed_code='SE_EXISTE', coord_x=500100.0, coord_y=8500000.0,
            installed_kva=100, kw=10,
        )
        plan = build_create_plan(_model(), rows, errors)
        assert plan.already_exists == ['SE_EXISTE']
        assert plan.create == []

    def test_missing_transformer_rating_is_reported(self):
        _rows, errors = single_new_load(
            sed_code='SE_X', coord_x=500100.0, coord_y=8500000.0, installed_kva=0, kw=10,
        )
        assert any('kVA_instalado' in e for e in errors)


class TestCreateTemplate:
    def test_template_has_the_coordinate_and_conductor_columns(self, tmp_path):
        from openpyxl import load_workbook

        path = write_create_template([], tmp_path / 'crear.xlsx', feeder='AL01')
        header = [c.value for c in load_workbook(path).active[1]]
        assert header == list(CREATE_COLUMNS)
        assert 'CoordX' in header and 'CoordY' in header and 'conductor' in header

    def test_prefilled_load_keeps_only_its_own_pair(self, tmp_path):
        """Precargada desde (kVA)/FP, se puede corregir el kVA sin chocar con un Kw viejo."""
        from openpyxl import load_workbook

        from igea_dgs.loads import SedLoad, read_sheet, SHEET_COLUMNS

        desde_kva = read_sheet(list(SHEET_COLUMNS), [['SE_KVA', '', '', '50', '0.92']], 'AL01').rows[0]
        desde_kw = SedLoad('SE_KW', 'AL01', kw=28.0, kvar=9.5, kva=math.hypot(28, 9.5), fp=0.947)
        path = write_create_template([desde_kva, desde_kw], tmp_path / 'crear.xlsx', feeder='AL01')
        wb = load_workbook(path)
        ws = wb.active
        filas = {r[0]: dict(zip(CREATE_COLUMNS, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
        assert filas['SE_KVA']['Kw'] is None and filas['SE_KVA']['(kVA)'] == pytest.approx(50.0)
        assert filas['SE_KW']['(kVA)'] is None and filas['SE_KW']['Kw'] == pytest.approx(28.0)

        ws['B2'], ws['C2'], ws['G2'] = 500130, 8500040, 160      # SE_KVA: punto y trafo
        ws['J2'] = 80                                              # corrige el kVA
        ws['B3'], ws['C3'], ws['G3'] = 500150, 8500040, 100
        wb.save(path)
        values = list(load_workbook(path, read_only=True).active.iter_rows(values_only=True))
        rows, errors = read_create_sheet(list(values[0]), [list(v) for v in values[1:]], 'AL01')
        assert errors == []
        assert rows[0].kw == pytest.approx(80 * 0.92)

    def test_node_help_sheet_is_added(self, tmp_path):
        from openpyxl import load_workbook

        path = write_create_template(
            [], tmp_path / 'crear.xlsx', feeder='AL01', node_choices=['N1', 'N2'],
        )
        wb = load_workbook(path)
        assert 'nodos_validos' in wb.sheetnames
        assert [r[0].value for r in wb['nodos_validos'].iter_rows(min_row=2)] == ['N1', 'N2']

    def test_sheet_reads_coordinates(self):
        rows, errors = read_create_sheet(
            list(CREATE_COLUMNS),
            [['SE_A', '500130', '8500040', '', '', '1', '160', '', '', '45', '0.95', '', '']],
            'AL01',
        )
        assert errors == []
        assert rows[0].coord_x == pytest.approx(500130.0)
        assert rows[0].coord_y == pytest.approx(8500040.0)

    def test_negative_coordinates_are_valid(self):
        """Husos y hemisferios dan coordenadas negativas: no son un error."""
        rows, errors = read_create_sheet(
            list(CREATE_COLUMNS),
            [['SE_A', '-500130', '-8500040', '', '', '1', '160', '', '', '45', '0.95', '', '']],
            'AL01',
        )
        assert errors == []
        assert rows[0].coord_x == pytest.approx(-500130.0)

    def test_non_numeric_coordinates_are_reported(self):
        _rows, errors = read_create_sheet(
            list(CREATE_COLUMNS),
            [['SE_A', 'x', 'y', '', '', '1', '160', '', '', '45', '0.95', '', '']],
            'AL01',
        )
        assert any('CoordX' in e for e in errors)

    def test_column_aliases_are_tolerated(self):
        rows, errors = read_create_sheet(
            ['SED', 'Este', 'Norte', 'trafo', 'S kVA', 'cos phi'],
            [['SE_A', '500130', '8500040', '160', '45', '0.95']],
            'AL01',
        )
        assert errors == []
        assert rows[0].coord_x == pytest.approx(500130.0)
        assert rows[0].installed_kva == pytest.approx(160.0)


class TestPayloadForPowerFactory:
    def _payload(self, **kwargs):
        rows, errors = single_new_load(
            sed_code='SE_NUEVA', coord_x=500130.0, coord_y=8500040.0,
            installed_kva=160, kva=45, fp=0.95, feeder='AL01', **kwargs,
        )
        plan = build_create_plan(_model(), rows, errors)
        return create_plan_to_payload(plan, nominal_kv=22.9, source_crs='EPSG:32718')

    def test_payload_carries_everything_the_script_needs(self):
        item = self._payload()['create'][0]
        for key in (
            'sed_code', 'connect_to_node', 'distance_m', 'new_node_id', 'new_section_id',
            'conductor_code', 'conductor_ampacity_a', 'length_km', 'strn_mva',
            'plini_mw', 'qlini_mvar', 'slini_mva',
        ):
            assert key in item, key
        assert item['length_km'] == pytest.approx(0.050)
        assert item['strn_mva'] == pytest.approx(0.160)

    def test_gps_is_computed_from_the_projected_coordinates(self):
        pytest.importorskip('pyproj')
        item = self._payload()['create'][0]
        # EPSG:32718 en esa zona cae en la costa peruana.
        assert item['gps_lat'] is not None and item['gps_lon'] is not None
        assert -19 < item['gps_lat'] < -3
        assert -82 < item['gps_lon'] < -68

    def test_payload_without_crs_leaves_gps_empty(self):
        rows, errors = single_new_load(
            sed_code='SE_N', coord_x=500130.0, coord_y=8500040.0,
            installed_kva=160, kw=40, feeder='AL01',
        )
        plan = build_create_plan(_model(), rows, errors)
        item = create_plan_to_payload(plan, nominal_kv=22.9)['create'][0]
        assert item['gps_lat'] is None, 'sin CRS no se inventa GPS'

    def test_payload_is_json_serialisable(self):
        import json

        assert json.loads(json.dumps(self._payload()))['summary']['create'] == 1


class TestRealFeederCreation:
    """Sobre un alimentador real, con su catálogo y su geometría."""

    def test_create_from_a_point_near_the_feeder(self, ds):
        from igea_dgs.model import build_feeder_model

        try:
            model = build_feeder_model(ds, 'PA217', strict=False)
        except KeyError:
            pytest.skip('PA217 no está en el export cargado')
        node_with_xy = next(n for n in model.nodes.values() if n.x is not None)
        rows, errors = single_new_load(
            sed_code='SE_PRUEBA_9', coord_x=node_with_xy.x + 40.0,
            coord_y=node_with_xy.y + 30.0, installed_kva=160, kva=45, fp=0.95,
        )
        plan = build_create_plan(model, rows, errors)
        assert plan.is_applicable, plan.row_errors
        item = plan.create[0]
        assert item.distance_m <= 50.0 + 1e-6
        assert item.conductor.code in {t.code for t in model.line_types.values()}
        assert item.conductor.current_a == pytest.approx(160 / (math.sqrt(3) * model.nominal_kv))


def _create_tool():
    import importlib.util
    from pathlib import Path

    path = Path(__file__).resolve().parents[1] / 'tools' / 'create_sed_loads.py'
    spec = importlib.util.spec_from_file_location('create_sed_loads', path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class TestDrawingInPowerFactory:
    """La SED nueva se dibuja en el unifilar, a la escala con que se importó el DGS.

    PowerFactory no guarda esa escala, así que se mide con los nodos cercanos: lo que
    se prueba aquí es esa medida, que no necesita PowerFactory.
    """

    def test_plan_carries_the_reference_nodes(self):
        model = _model()
        rows, errors = single_new_load(
            sed_code='SE_N', coord_x=500130.0, coord_y=8500040.0, installed_kva=160, kw=40,
        )
        payload = create_plan_to_payload(build_create_plan(model, rows, errors), nominal_kv=22.9)
        refs = payload['create'][0]['reference_nodes']
        assert refs[0]['node_id'] == 'N2', 'primero el nodo de enganche'
        assert {r['node_id'] for r in refs} == {'N1', 'N2', 'N3'}
        assert refs[0]['x'] == pytest.approx(500100.0)

    def test_scale_and_offset_are_measured_from_drawn_nodes(self):
        tool = _create_tool()
        # Diagrama a 2,08 u/m con origen desplazado, como la escala NA205.
        s, bx, by = 2.08, -1_040_000.0, -17_680_000.0
        refs = [(x, y, s * x + bx, s * y + by)
                for x, y in ((500000, 8500000), (500100, 8500000), (500100, 8500300))]
        to_diagram = tool.diagram_transform(refs)
        gx, gy = to_diagram(500130.0, 8500040.0)
        assert gx == pytest.approx(s * 500130 + bx)
        assert gy == pytest.approx(s * 8500040 + by)

    def test_a_straight_street_still_gives_the_scale(self):
        """Todos los nodos con la misma y: un ajuste por eje fallaría en y."""
        tool = _create_tool()
        refs = [(x, 100.0, 3 * x, 300.0) for x in (0.0, 50.0, 120.0)]
        assert tool.diagram_transform(refs)(10.0, 110.0) == pytest.approx((30.0, 330.0))

    def test_without_two_distinct_points_there_is_no_scale(self):
        tool = _create_tool()
        assert tool.diagram_transform([(1.0, 1.0, 5.0, 5.0)]) is None
        assert tool.diagram_transform([(1.0, 1.0, 5.0, 5.0), (1.0, 1.0, 5.0, 5.0)]) is None

    def test_without_scale_the_new_bus_is_set_apart_from_the_anchor(self):
        tool = _create_tool()
        xy = tool.new_sed_positions({'coord_x': None, 'coord_y': None}, (10.0, 10.0), None, 4.0)
        assert xy == (14.0, 6.0)

    def test_line_rotation_matches_the_converter(self):
        tool = _create_tool()
        assert tool.line_rotation((0, 0), (1, 1)) == 45
        assert tool.line_rotation((0, 0), (-1, 0)) == 180


class TestSplitCreateByFeeder:
    def _far_model(self):
        """Otro alimentador 5 km al este del de ``_model``."""
        m = _model()
        nodes = {f'F{i}': type(n)(f'F{i}', n.x + 5000.0, n.y) for i, n in enumerate(m.nodes.values())}
        return type(m)(name='AL02', network_id='NET_AL02', nominal_kv=22.9, source_node='F0',
                       nodes=nodes, lines=[], loads=[], devices=[], line_types=m.line_types, seds=[])

    def test_sheet_named_as_a_feeder_wins(self):
        from igea_dgs.loads_create import split_create_by_feeder

        rows, _ = single_new_load(sed_code='SE_A', coord_x=505100.0, coord_y=8500000.0,
                                  installed_kva=100, kw=10)
        out, sueltos = split_create_by_feeder({'AL01': (rows, [])}, {'AL01': _model(), 'AL02': self._far_model()})
        assert [r.sed_code for r in out['AL01'][0]] == ['SE_A'] and sueltos == []

    def test_coordinates_pick_the_nearest_feeder(self):
        from igea_dgs.loads_create import split_create_by_feeder

        cerca_1, _ = single_new_load(sed_code='SE_1', coord_x=500120.0, coord_y=8500010.0, installed_kva=100, kw=10)
        cerca_2, _ = single_new_load(sed_code='SE_2', coord_x=505110.0, coord_y=8500010.0, installed_kva=100, kw=10)
        out, sueltos = split_create_by_feeder(
            {'NUEVAS': (cerca_1 + cerca_2, [])}, {'AL01': _model(), 'AL02': self._far_model()})
        assert sueltos == []
        assert [r.sed_code for r in out['AL01'][0]] == ['SE_1']
        assert [r.sed_code for r in out['AL02'][0]] == ['SE_2']
        assert out['AL02'][0][0].feeder == 'AL02'

    def test_node_id_picks_its_feeder(self):
        from igea_dgs.loads_create import split_create_by_feeder

        rows, _ = single_new_load(sed_code='SE_N', node_id='F1', installed_kva=100, kw=10)
        out, _ = split_create_by_feeder({'NUEVAS': (rows, [])}, {'AL01': _model(), 'AL02': self._far_model()})
        assert [r.sed_code for r in out['AL02'][0]] == ['SE_N']

    def test_coordinate_tie_is_ambiguous_not_dictionary_order(self):
        from igea_dgs.loads_create import split_create_by_feeder

        same_a = _model()
        same_b = _model()
        rows, _ = single_new_load(
            sed_code='SE_TIE', coord_x=500100.0, coord_y=8500000.0,
            installed_kva=100, kw=10,
        )
        out, sueltos = split_create_by_feeder(
            {'NUEVAS': (rows, [])}, {'F-A': same_a, 'F-B': same_b},
        )
        assert out['F-A'][0] == [] and out['F-B'][0] == []
        assert len(sueltos) == 1
        assert 'ambigua' in sueltos[0].lower()
        assert 'F-A' in sueltos[0] and 'F-B' in sueltos[0]

    def test_row_without_any_clue_is_reported(self):
        from igea_dgs.loads_create import NewSedLoad, split_create_by_feeder

        out, sueltos = split_create_by_feeder(
            {'NUEVAS': ([NewSedLoad('SE_X', installed_kva=100)], [])}, {'AL01': _model()})
        assert out['AL01'][0] == [] and 'SE_X' in sueltos[0]


def _lote_tool():
    import importlib.util
    import sys
    from pathlib import Path

    tools = Path(__file__).resolve().parents[1] / 'tools'
    sys.path.insert(0, str(tools))
    try:
        spec = importlib.util.spec_from_file_location('aplicar_lote_cargas', tools / 'aplicar_lote_cargas.py')
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod
    finally:
        sys.path.remove(str(tools))


class TestBatchTool:
    def test_scenario_and_variation_names_are_valid_in_powerfactory(self):
        tool = _lote_tool()
        assert tool.nombre_escenario('PA217', '20261002_0930') == 'Cargas_PA217_20261002_0930'
        assert tool.nombre_variacion('PA 2/17', '20261002_0930') == 'SED_nuevas_PA_2_17_20261002_0930'
        assert len(tool.nombre_variacion('X' * 60, '20261002_0930')) <= tool.MAX_NOMBRE

    def test_undrawn_is_a_warning_and_not_found_is_a_failure(self):
        tool = _lote_tool()
        assert not tool.con_fallos({'creacion': {'created': [1], 'undrawn': ['SE: sin diagrama']}})
        assert tool.con_fallos({'actualizacion': {'not_found': ['SE1']}})
        assert tool.con_fallos({'creacion': {'error': 'no se pudo crear la etapa'}})
        assert tool.con_fallos({'error': 'boom'})
