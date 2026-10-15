"""Escritura en Word de los informes de estudio eléctrico y de factibilidad (spec 017).

Los dos informes son esquemas generales: sirven para cualquier alimentador, estudio o
cliente nuevo. Del modelo de la empresa se toma solo la forma (cabecera, pie, logo,
estilos y orden de las secciones); su cuerpo se vacía y de sus datos no queda nada: ni el
n.º ni la referencia de la cabecera, ni el total de páginas escrito a mano en el pie
(RF-3, RF-6). Lo que no llega del análisis o del formulario queda como un espacio vacío
resaltado para rellenarlo en Word, y su motivo va al JSON (RF-23). Nada nombra los casos
de DIgSILENT ni el orden de evaluación de las obras: el Word va al destinatario tal cual.
"""
from __future__ import annotations

import io
import json
import os
import re
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Mapping, Sequence

from .estudios_pf import Limites
from .informes import (
    ETAPAS, NOMBRE_ETAPA, SOLICITUDES_VIGENTES, TRIFASICA, Caso, ConfigInforme, Devanado,
    Encabezado, Etapa, InformeError, PerdidasAnuales, Solicitud, _pu, antecedentes_factibilidad,
    conclusiones_estudio, conclusiones_factibilidad, descripcion_factibilidad, describir_obra,
    etapa_recomendada, numero, obras_a_ejecutar, redactar_cargabilidad, redactar_tension,
    texto_transformador, veredicto_factibilidad)
from .libros_informe import DatosPerdidas

#: Imágenes de cada estado: vista de ubicación y los dos coloreados (RF-10, RF-46).
IMAGENES = ('ubicacion', 'tension', 'cargabilidad')
_NOMBRE_IMAGEN = {'ubicacion': 'vista topológica y de ubicación', 'tension': 'coloreado por tensión',
                  'cargabilidad': 'coloreado por cargabilidad'}
#: Ancho de las figuras: el del cuerpo de una página A4 con los márgenes del modelo.
_ANCHO_FIGURA_CM = 16.0
_ORDINALES = ('PRIMERA', 'SEGUNDA', 'TERCERA', 'CUARTA', 'QUINTA', 'SEXTA', 'SÉPTIMA', 'OCTAVA')


@dataclass(frozen=True)
class DatosEstudio:
    """Lo que necesita el informe de estudio: lo medido en DIgSILENT más el formulario."""

    alimentadores: Sequence[str]
    encabezado: Encabezado
    actual: Caso
    etapas: Sequence[Etapa]
    #: Por estado (``actual`` o la clave de una etapa), la ruta de cada imagen de IMAGENES.
    imagenes: Mapping[str, Mapping[str, Path | None]] = field(default_factory=dict)
    numero: str | None = None
    antecedentes: Sequence[str] = ()
    descripcion: Sequence[str] = ()
    #: Bancos de condensadores del alimentador como están en DIgSILENT (nombre, barra,
    #: kvar, en_servicio): entran en la calibración y el informe los muestra.
    bancos: Sequence[Mapping[str, Any]] = ()
    #: Clientes fijos de la última calibración con su FP real (EA/ERI) y el corregido
    #: (spec 015, RF-65).
    fijas: Sequence[Mapping[str, Any]] = ()


@dataclass(frozen=True)
class DatosFactibilidad:
    """Lo que necesita el informe de factibilidad de una carga nueva (RF-47 a RF-72)."""

    encabezado: Encabezado
    actual: Caso
    proyectado: Caso
    solicitud: Solicitud | None = None
    #: Datos en el punto de diseño del estado proyectado (RF-52): u_pu, kv_ll, kv_ln, i_a, kva, kw, kvar.
    punto_diseno: Mapping[str, float] | None = None
    #: Estructura de media tensión del punto de diseño, con su código de la red de origen (RF-53).
    codigo_punto: str | None = None
    #: Si el estado proyectado incluye cargas ingresadas en bloque (RF-50).
    solicitudes_vigentes: bool = False
    perdidas_actual: PerdidasAnuales | None = None
    perdidas_proyectado: PerdidasAnuales | None = None
    libro: DatosPerdidas | None = None
    devanado: Devanado | None = None
    caso_transmision: str = ''
    #: Por estado (``actual``, ``proyectado``) sus imágenes; ``transformador`` con la de la SET.
    imagenes: Mapping[str, Mapping[str, Path | None]] = field(default_factory=dict)
    fuentes: Sequence[str] = ()
    numero: str | None = None


# --------------------------------------------------------------------------- modelo
def abrir_modelo(ruta: Path):
    """El modelo como documento de ``python-docx``, aunque Word lo haya dejado dañado.

    El modelo de Casa Blanca, reparado con Word, trae imágenes de la cabecera y del pie
    que apuntan a «../NULL», una parte que no existe, y ``python-docx`` no lo abre. Esas
    relaciones se quitan en memoria; el fichero de la empresa no se toca.
    """
    import docx

    ruta = Path(ruta)
    try:
        with zipfile.ZipFile(ruta) as origen:
            nombres = set(origen.namelist())
            memoria = io.BytesIO()
            with zipfile.ZipFile(memoria, 'w', zipfile.ZIP_DEFLATED) as copia:
                for info in origen.infolist():
                    datos = origen.read(info)
                    if info.filename.endswith('.rels'):
                        datos = _sin_relaciones_rotas(info.filename, datos, nombres)
                    copia.writestr(info, datos)
        memoria.seek(0)
        return docx.Document(memoria)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise InformeError(f'No se pudo abrir el modelo del informe {ruta.name}: {exc}. '
                           'Ábralo en Word y guárdelo como .docx.') from exc


def _sin_relaciones_rotas(rels: str, datos: bytes, nombres: set[str]) -> bytes:
    """Quita las relaciones internas cuyo destino no está en el paquete."""
    carpeta = rels.replace('_rels/', '')
    base = carpeta.rsplit('/', 1)[0] if '/' in carpeta else ''

    def destino_existe(m: re.Match) -> bool:
        etiqueta = m.group(0)
        objetivo = re.search(rb'Target="([^"]*)"', etiqueta)
        if b'TargetMode="External"' in etiqueta or objetivo is None:
            return True
        destino = objetivo.group(1).decode()
        partes = destino.split('/') if destino.startswith('/') else (base + '/' + destino).split('/')
        pila: list[str] = []
        for p in partes:
            if p == '..':
                if pila:
                    pila.pop()
            elif p not in ('', '.'):
                pila.append(p)
        return '/'.join(pila) in nombres

    return re.sub(rb'<Relationship [^>]*/>', lambda m: m.group(0) if destino_existe(m) else b'', datos)


def _vaciar_cuerpo(documento) -> None:
    """Quita el cuerpo del modelo y deja una sola sección: la primera del modelo.

    Los dos modelos de Electro Dunas tienen dos secciones y la tabla de la cabecera (n.º
    de informe y referencia) es de la primera; la segunda, la de los anexos, tiene otra
    cabecera sin ella. El informe nuevo empieza donde empieza el modelo, así que se queda
    con la primera: su cabecera, su pie, su tamaño de página y sus márgenes. Las
    referencias que la primera no trae las hereda de la última, como hace Word.
    """
    import copy

    from docx.oxml.ns import qn

    cuerpo = documento.element.body
    final = cuerpo.find(qn('w:sectPr'))
    primera = cuerpo.find('.//' + qn('w:pPr') + '/' + qn('w:sectPr'))
    if final is not None and primera is not None:
        nueva = copy.deepcopy(primera)
        for etiqueta in ('w:headerReference', 'w:footerReference'):
            tiene = {r.get(qn('w:type')) for r in nueva.findall(qn(etiqueta))}
            for ref in final.findall(qn(etiqueta)):
                if ref.get(qn('w:type')) not in tiene:
                    nueva.insert(0, copy.deepcopy(ref))
        final.addprevious(nueva)
        cuerpo.remove(final)
    for hijo in list(cuerpo):
        if not hijo.tag.endswith('}sectPr'):
            cuerpo.remove(hijo)


def _celdas(partes):
    for parte in partes:
        for tabla in parte.tables:
            for fila in tabla.rows:
                yield from fila.cells


def _rellenar_cabecera(documento, rotulos: Mapping[str, str] | None, valores: Mapping[str, str | None]) -> None:
    """Deja en las celdas de la cabecera solo su rótulo y, debajo, el dato de este informe.

    El n.º y la referencia del modelo («GT-019-PRN-2024», «CAIDA DE TENSIÓN EN
    CASABLANCA…») no pueden pasar al informe nuevo: sin dato, la celda queda con el rótulo
    y el hueco. Los rótulos son del perfil de empresa.
    """
    if not rotulos:
        return
    cabeceras = [s.header for s in documento.sections] + [s.first_page_header for s in documento.sections]
    for celda in _celdas(cabeceras):
        texto = celda.text.strip()
        for clave, rotulo in rotulos.items():
            if texto.upper().startswith(rotulo.upper()):
                primera = next((x for x in texto.split('\n') if x.strip()), texto)
                _texto_celda(celda, primera, valores.get(clave))


def _texto_celda(celda, rotulo: str, valor: str | None) -> None:
    """El rótulo en la primera línea y el valor (o un hueco) en la segunda, con el formato del modelo."""
    parrafos = celda.paragraphs
    for p in parrafos[1:]:
        p._element.getparent().remove(p._element)
    p = parrafos[0]
    for r in p.runs[1:]:
        r._element.getparent().remove(r._element)
    if p.runs:
        p.runs[0].text = rotulo
    else:
        p.add_run(rotulo)
    segunda = celda.add_paragraph(style=p.style)
    if valor:
        segunda.add_run(valor)
    else:
        _vacio(segunda)


_TOTAL_FIJO = re.compile(r'^(\s*de\s*)\d+\s*$')


def _total_de_paginas(documento) -> None:
    """Cambia el «de 6» escrito a mano junto al n.º de página por el campo de páginas."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pies = [s.footer for s in documento.sections] + [s.first_page_footer for s in documento.sections]
    parrafos = [p for pie in pies for p in pie.paragraphs]
    parrafos += [p for c in _celdas(pies) for p in c.paragraphs]
    for p in parrafos:
        if 'NUMPAGES' in p._p.xml:
            continue
        runs = p.runs
        for i, run in enumerate(runs):
            m = _TOTAL_FIJO.match(run.text)
            if m:
                run.text = m.group(1)
            # En el modelo EMAPICA « de » y «8» son dos fragmentos distintos.
            elif (run.text.strip().isdigit() and i > 0
                  and re.fullmatch(r'\s*de\s*', runs[i - 1].text or '')):
                run.text = ''
            else:
                continue
            anterior = run._element
            for tipo, texto in (('begin', None), (None, ' NUMPAGES '), ('separate', None), (None, '1'),
                                ('end', None)):
                nuevo = OxmlElement('w:r')
                if run._element.rPr is not None:
                    nuevo.append(run._element.rPr.__copy__())
                if tipo:
                    fld = OxmlElement('w:fldChar')
                    fld.set(qn('w:fldCharType'), tipo)
                    nuevo.append(fld)
                elif texto.strip().isdigit():
                    t = OxmlElement('w:t')
                    t.text = texto
                    nuevo.append(t)
                else:
                    instr = OxmlElement('w:instrText')
                    instr.set(qn('xml:space'), 'preserve')
                    instr.text = texto
                    nuevo.append(instr)
                anterior.addnext(nuevo)
                anterior = nuevo


def _preparar(modelo: Path, rotulos, valores) :
    documento = abrir_modelo(modelo)
    _vaciar_cuerpo(documento)
    _rellenar_cabecera(documento, rotulos, valores)
    _total_de_paginas(documento)
    return documento


# --------------------------------------------------------------------------- escritura
def _vacio(parrafo) -> None:
    """Un espacio vacío resaltado, para que el ingeniero lo rellene en Word."""
    from docx.enum.text import WD_COLOR_INDEX

    parrafo.add_run('\u00a0' * 24).font.highlight_color = WD_COLOR_INDEX.YELLOW


class _Escritor:
    """El cuerpo del informe; lo que falta queda vacío y su motivo en ``huecos`` (RF-23)."""

    def __init__(self, documento) -> None:
        self.d = documento
        self.huecos: list[str] = []
        estilos = {s.name for s in documento.styles}
        self.estilo_tabla = 'Table Grid' if 'Table Grid' in estilos else None
        self.estilo_lista = 'List Bullet' if 'List Bullet' in estilos else None

    def titulo(self, texto: str) -> None:
        # El modelo escribe los títulos en estilo «Normal», en negrita y con su número.
        self.d.add_paragraph().add_run(texto).bold = True

    def parrafo(self, texto: str) -> None:
        self.d.add_paragraph(texto)

    def campo(self, rotulo: str, valor: str | None, motivo: str) -> None:
        """Un campo del formulario: «ASUNTO : …», vacío si el operador no lo dio."""
        p = self.d.add_paragraph(rotulo)
        if valor:
            p.add_run(valor)
        else:
            _vacio(p)
            self.huecos.append(motivo)

    def hueco(self, motivo: str) -> None:
        _vacio(self.d.add_paragraph())
        self.huecos.append(motivo)

    def vinieta(self, texto: str) -> None:
        self.d.add_paragraph(texto, style=self.estilo_lista)

    def imagen(self, ruta: Path | None, que: str, pie: str | None = None) -> None:
        from docx.shared import Cm

        if ruta is None or not Path(ruta).is_file():
            self.hueco(f'imagen de {que}' + (f' ({Path(ruta).name} no existe)' if ruta else ''))
            return
        self.d.add_picture(str(ruta), width=Cm(_ANCHO_FIGURA_CM))
        if pie:
            self.parrafo(pie)

    def tabla(self, cabecera: Sequence[str], filas: Sequence[Sequence[str]]) -> None:
        t = self.d.add_table(rows=1, cols=len(cabecera))
        if self.estilo_tabla:
            t.style = self.estilo_tabla
        for celda, texto in zip(t.rows[0].cells, cabecera):
            celda.text = texto
            for r in celda.paragraphs[0].runs:
                r.bold = True
        for fila in filas:
            for celda, texto in zip(t.add_row().cells, fila):
                celda.text = texto


def _f(valor: float | None, decimales: int) -> str:
    return '' if valor is None else numero(valor, decimales)


def _leyenda(bandas: Sequence[float] | None, que: str) -> str | None:
    """La leyenda de un coloreado con las bandas del perfil (RF-113, RF-114)."""
    if not bandas:
        return None
    tramos = [f'menos de {numero(bandas[0], 0)} %']
    tramos += [f'de {numero(a, 0)} a {numero(b, 0)} %' for a, b in zip(bandas, bandas[1:])]
    tramos.append(f'más de {numero(bandas[-1], 0)} %')
    return f'Leyenda del coloreado por {que}: ' + '; '.join(tramos) + '.'


def _encabezado(e: _Escritor, enc: Encabezado) -> None:
    """Lugar y fecha, asunto y destinatario, como el modelo (RF-22)."""
    e.parrafo(enc.fecha_texto)
    e.campo('ASUNTO\t:\t', enc.asunto, 'asunto del formulario')
    lineas = [x for x in (enc.destinatario or '').split('\n') if x.strip()]
    e.campo('PARA\t:\t', lineas[0] if lineas else None, 'destinatario del formulario')
    for linea in lineas[1:]:
        e.parrafo(f'\t\t{linea}')


def _datos_fuente(e: _Escritor, caso: Caso, rotulo: str) -> None:
    """La tabla «Carga Total» de la fuente y de las pérdidas (RF-15)."""
    e.parrafo(rotulo)
    filas = []
    for nombre, p in (('Potencia desde la fuente', caso.fuente), ('Pérdidas en las líneas', caso.perdidas)):
        if p is not None:
            filas.append([nombre, _f(p.kw, 2), _f(p.kvar, 2), _f(p.kva, 2), _f(p.fp_pct, 2)])
    e.tabla(['Carga total', 'kW', 'kVAR', 'kVA', 'FP (%)'], filas)


def _tabla_tension(e: _Escritor, caso: Caso) -> None:
    """Tensión mínima por fase, o una sola trifásica, en % de la nominal (RF-17, RF-18, RF-118)."""
    if caso.tension_min:
        e.tabla(['Caída de tensión', '%', 'Nodo'],
                [['Trifásica' if fase == TRIFASICA else fase, f'{numero(x.valor * 100, 2)} %', x.donde]
                 for fase, x in caso.tension_min.items()])


def _analisis_tension(e: _Escritor, titulo: str, caso: Caso, imagen: Path | None, estado: str,
                      config: ConfigInforme, tabla: bool = True) -> None:
    e.titulo(titulo)
    leyenda = _leyenda(config.bandas_tension_pct, 'tensión')
    if leyenda:
        e.parrafo(leyenda)
    e.parrafo(redactar_tension(caso, config.limites, config.norma))
    e.imagen(imagen, f'{_NOMBRE_IMAGEN["tension"]} del {estado}')
    if tabla:
        _tabla_tension(e, caso)


def _analisis_carga(e: _Escritor, titulo: str, caso: Caso, imagen: Path | None, estado: str,
                    config: ConfigInforme) -> None:
    e.titulo(titulo)
    leyenda = _leyenda(config.bandas_cargabilidad_pct, 'cargabilidad')
    if leyenda:
        e.parrafo(leyenda)
    e.parrafo(redactar_cargabilidad(caso, config.limites))
    e.imagen(imagen, f'{_NOMBRE_IMAGEN["cargabilidad"]} del {estado}')


def _sin_flujo(e: _Escritor, caso: Caso) -> bool:
    """Si el flujo no convergió, la sección lo dice y no trae valores (RF-24)."""
    if caso.converge:
        return False
    e.parrafo('El flujo de carga no converge: no se obtienen resultados para este escenario.')
    return True


def _guardar(documento, destino: Path, resumen: Mapping[str, Any]) -> None:
    """``.docx`` y JSON en ``.tmp`` y renombrados al final: nada a medias si algo falla (RF-25, RF-27)."""
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    tmp_docx = destino.with_name(destino.name + '.tmp')
    tmp_json = destino.with_suffix('.json.tmp')
    try:
        documento.save(tmp_docx)
        tmp_json.write_text(json.dumps(resumen, ensure_ascii=False, indent=2, default=str) + '\n',
                            encoding='utf-8')
        os.replace(tmp_json, destino.with_suffix('.json'))
        os.replace(tmp_docx, destino)
    finally:
        for tmp in (tmp_docx, tmp_json):
            tmp.unlink(missing_ok=True)


def _rutas(imagenes: Mapping[str, Mapping[str, Path | None]]) -> dict:
    return {k: {i: str(r) if r else None for i, r in v.items()} for k, v in imagenes.items()}


def _limites(config: ConfigInforme, tipo: str) -> None:
    config.modelo(tipo)  # dice qué clave del perfil falta (RF-5, RF-14)


# --------------------------------------------------------------------------- estudio
def _resumen_obras(etapas: Sequence[Etapa]) -> str:
    nombres = [NOMBRE_ETAPA[x.clave].lower() for x in etapas if x.obras]
    if not nombres:
        return 'las obras de mejora'
    return nombres[0] if len(nombres) == 1 else ', '.join(nombres[:-1]) + ' y ' + nombres[-1]


def _conclusiones(datos: DatosEstudio, rec, config: ConfigInforme) -> list[str]:
    return conclusiones_estudio(list(datos.alimentadores), datos.actual, rec, list(datos.etapas),
                                config.limites, config.norma)


def escribir_informe_estudio(config: ConfigInforme, destino: Path, datos: DatosEstudio) -> list[str]:
    """Escribe el informe de estudio en ``destino`` y su JSON al lado; devuelve los huecos.

    El orden es el del modelo de Casa Blanca (RF-3), sin el recibo de luz (RF-9).
    """
    _limites(config, 'estudio')
    enc = datos.encabezado
    documento = _preparar(config.modelo('estudio'), config.rotulos_cabecera,
                          {'numero': datos.numero, 'referencia': enc.asunto})
    e = _Escritor(documento)
    lista = ', '.join(datos.alimentadores)
    _encabezado(e, enc)

    e.titulo('1. ANTECEDENTES')
    antecedentes = list(datos.antecedentes) or ([enc.antecedentes] if enc.antecedentes else [])
    for texto in antecedentes:
        e.parrafo(texto)
    if not antecedentes:
        e.hueco('antecedentes del formulario')

    e.titulo('2. DESCRIPCIÓN')
    for texto in datos.descripcion:
        e.parrafo(texto)
    if not datos.descripcion:
        e.hueco('descripción del formulario')
    actual_img = datos.imagenes.get('actual') or {}
    e.titulo('2.1 VISTA TOPOLÓGICA Y UBICACIÓN')
    e.imagen(actual_img.get('ubicacion'), f'{_NOMBRE_IMAGEN["ubicacion"]} del estado actual')

    e.titulo(f'2.2 ANÁLISIS DE CARGABILIDAD Y CAÍDA DE TENSIÓN DEL ALIMENTADOR {lista} EN ESTADO ACTUAL')
    e.titulo(f'2.2.1 ESCENARIO ACTUAL DEL ALIMENTADOR {lista}')
    if not _sin_flujo(e, datos.actual):
        _datos_fuente(e, datos.actual, 'Datos de la fuente en MD:')
        _bancos(e, datos.bancos, lista)
        _fijas(e, datos.fijas, lista)
        _analisis_tension(e, '2.2.1.1 ANÁLISIS DE CAÍDA DE TENSIÓN', datos.actual, actual_img.get('tension'),
                          'estado actual', config)
        _analisis_carga(e, '2.2.1.2 ANÁLISIS DE CARGABILIDAD', datos.actual, actual_img.get('cargabilidad'),
                        'estado actual', config)

    etapas = sorted(datos.etapas, key=lambda x: ETAPAS.index(x.clave))
    rec = etapa_recomendada(datos.actual, etapas, config.limites)
    a_ejecutar = obras_a_ejecutar(rec, etapas)
    if not rec.actual_cumple and a_ejecutar:
        _proyectado(e, datos, a_ejecutar, config, lista)

    e.titulo('3. CONCLUSIONES Y RECOMENDACIONES')
    obras = {describir_obra(o) for x in a_ejecutar for o in x.obras}
    for texto in _conclusiones(datos, rec, config):
        if texto in obras:
            e.vinieta(texto)
        else:
            e.parrafo(texto)

    _guardar(documento, destino, {
        'tipo': 'estudio', 'alimentadores': list(datos.alimentadores), 'numero': datos.numero,
        'encabezado': {**asdict(enc), 'fecha': enc.fecha.isoformat()},
        'actual': asdict(datos.actual), 'etapas': [asdict(x) for x in etapas],
        'recomendada': rec.etapa.clave if rec.etapa else None, 'cumplen': dict(rec.cumplen),
        'obras_propuestas': [x.clave for x in a_ejecutar], 'bancos': [dict(b) for b in datos.bancos],
        'fijas': [dict(f) for f in datos.fijas],
        'imagenes': _rutas(datos.imagenes), 'huecos': e.huecos,
    })
    return e.huecos


def _bancos(e: _Escritor, bancos: Sequence[Mapping[str, Any]], lista: str) -> None:
    """Los bancos de condensadores existentes, tal como están en DIgSILENT."""
    if not bancos:
        e.parrafo(f'El alimentador {lista} no tiene bancos de condensadores en el modelo.')
        return
    e.parrafo(f'Bancos de condensadores existentes en el alimentador {lista}:')
    e.tabla(['Banco', 'Nodo', 'kvar', 'Estado'],
            [[str(b.get('nombre') or ''), str(b.get('barra') or ''), _f(b.get('kvar'), 0),
              'En servicio' if b.get('en_servicio') else 'Fuera de servicio'] for b in bancos])


def _fijas(e: _Escritor, fijas: Sequence[Mapping[str, Any]], lista: str) -> None:
    """Clientes fijos con su FP real (EA/ERI) y el corregido de la calibración (015:RF-65).

    El usuario, 2026-10-15: «poner factor de potencia real y el corregido en el informe
    modelo». En CA101 los clientes iban de 0,43 a 0,98 y se escribieron a 0,98 por su banco
    propio; el lector tiene que ver los dos valores.
    """
    if not fijas:
        return
    e.parrafo(f'Clientes importantes del alimentador {lista} (cargas fijas de la calibración), '
              'con su factor de potencia real, calculado de su energía activa y reactiva, y '
              'el corregido con que se modelan:')
    e.tabla(['SED', 'Cliente', 'P (kW)', 'FP real', 'FP corregido', 'Q real (kvar)',
             'Q modelada (kvar)'],
            [[str(f.get('sed') or ''), str(f.get('cliente') or ''), _f(f.get('p_kw'), 1),
              _f(f.get('fp_real'), 3), _f(f.get('fp_corregido'), 3), _f(f.get('q_real_kvar'), 1),
              _f(f.get('q_kvar'), 1)] for f in fijas])
    compensados = [f for f in fijas if f.get('compensado')]
    if compensados:
        minimo = min(float(f.get('fp_corregido') or 0) for f in compensados)
        e.parrafo(f'Los {len(compensados)} clientes con factor de potencia real por debajo de '
                  f'{_f(minimo, 2)} se modelan compensados por su banco de condensadores propio '
                  f'hasta {_f(minimo, 2)}: ' + ', '.join(
                      f"{f.get('sed')} ({_f(f.get('fp_real'), 3)} → {_f(f.get('fp_corregido'), 3)})"
                      for f in compensados) + '.')


def _proyectado(e: _Escritor, datos: DatosEstudio, a_ejecutar: list[Etapa], config: ConfigInforme,
                lista: str) -> None:
    """Escenario proyectado con las obras propuestas y su detalle (RF-30, RF-44 a RF-46)."""
    final = a_ejecutar[-1]
    caso = final.caso
    resumen = _resumen_obras(a_ejecutar)
    imagenes = datos.imagenes.get(final.clave) or {}
    e.titulo(f'2.3 ANÁLISIS DE CARGABILIDAD Y CAÍDA DE TENSIÓN – ESCENARIO CON {resumen.upper()} '
             '– ESTADO PROYECTADO')
    texto = (f'Se realizó una evaluación de flujo de carga en el software DIgSILENT PowerFactory, '
             f'considerando {resumen}.')
    antes, despues = datos.actual.minima, caso.minima if caso.converge else None
    if antes is not None and despues is not None:
        cambio = 'una mejora' if despues.valor > antes.valor else 'una variación'
        texto += (f' Los resultados de la simulación indican {cambio} de la tensión mínima del alimentador '
                  f'{lista}, de {_pu(antes.valor)} a {_pu(despues.valor)}, según se detalla a continuación:')
    e.parrafo(texto)
    if not _sin_flujo(e, caso):
        _datos_fuente(e, caso, 'Datos de la fuente:')
        _analisis_tension(e, '2.3.1 ANÁLISIS DE CAÍDA DE TENSIÓN', caso, imagenes.get('tension'),
                          'estado proyectado', config)
        _analisis_carga(e, '2.3.2 ANÁLISIS DE CARGABILIDAD', caso, imagenes.get('cargabilidad'),
                        'estado proyectado', config)

    e.titulo(f'2.4 DETALLE DE LAS OBRAS DE MEJORA DEL ALIMENTADOR {lista}')
    n = 0
    for x in a_ejecutar:
        if not x.obras:
            continue
        ordinal = _ORDINALES[n] if n < len(_ORDINALES) else f'{n + 1}.ª'
        n += 1
        e.titulo(f'2.4.{n} {ordinal} ETAPA' + (f' – AÑO {x.anio}' if x.anio else '') + ':')
        for obra in x.obras:
            e.parrafo(describir_obra(obra))
    e.titulo(f'2.4.{n + 1} ESCENARIO ACTUAL Y PROYECTADO DEL ALIMENTADOR {lista}')
    e.parrafo(f'ESCENARIO ACTUAL: el alimentador {lista} en su configuración actual.')
    e.imagen((datos.imagenes.get('actual') or {}).get('ubicacion'),
             f'{_NOMBRE_IMAGEN["ubicacion"]} del estado actual')
    e.parrafo(f'ESCENARIO PROYECTADO: el alimentador {lista} con {resumen}.')
    e.imagen(imagenes.get('ubicacion'), f'{_NOMBRE_IMAGEN["ubicacion"]} del estado proyectado')


# --------------------------------------------------------------------------- factibilidad
def escribir_informe_factibilidad(config: ConfigInforme, destino: Path, datos: DatosFactibilidad) -> list[str]:
    """Escribe el informe de factibilidad en ``destino`` y su JSON al lado; devuelve los huecos.

    El orden es el de RF-106, el del modelo EMAPICA.
    """
    _limites(config, 'factibilidad')
    enc, s = datos.encabezado, datos.solicitud
    documento = _preparar(config.modelo('factibilidad'), config.rotulos_cabecera,
                          {'numero': datos.numero, 'referencia': enc.asunto})
    e = _Escritor(documento)
    alimentador = s.alimentador if s else ''
    _encabezado(e, enc)

    e.titulo('1. ANTECEDENTES')
    if s is None:
        e.hueco('datos de la solicitud del formulario')
    else:
        for texto in antecedentes_factibilidad(s):
            e.parrafo(texto)

    e.titulo('2. DESCRIPCIÓN')
    if s is None or not datos.fuentes:
        e.hueco('fuentes de información del análisis')
    else:
        for texto in descripcion_factibilidad(s, list(datos.fuentes)):
            e.parrafo(texto)
    actual_img = datos.imagenes.get('actual') or {}
    proy_img = datos.imagenes.get('proyectado') or {}
    e.titulo('2.1 VISTA TOPOLÓGICA Y UBICACIÓN')
    e.imagen(proy_img.get('ubicacion'), 'ubicación del cliente nuevo')

    e.titulo(f'2.2 ANÁLISIS DE CARGABILIDAD Y CAÍDA DE TENSIÓN DEL ALIMENTADOR {alimentador} EN ESTADO ACTUAL'
             .replace('  ', ' '))
    e.titulo(f'2.2.1 ESCENARIO ACTUAL DEL ALIMENTADOR {alimentador}'.rstrip())
    if not _sin_flujo(e, datos.actual):
        _datos_fuente(e, datos.actual, 'Datos de la fuente en MD:')
        _analisis_tension(e, '2.2.1.1 ANÁLISIS DE CAÍDA DE TENSIÓN', datos.actual, actual_img.get('tension'),
                          'estado actual', config)
        _analisis_carga(e, '2.2.1.2 ANÁLISIS DE CARGABILIDAD', datos.actual, actual_img.get('cargabilidad'),
                        'estado actual', config)

    e.titulo('2.3 ESTADO PROYECTADO')
    proyectado = datos.proyectado
    if not _sin_flujo(e, proyectado):
        _datos_fuente(e, proyectado, 'Datos de la fuente en MD:')
        e.titulo('2.3.1 ANÁLISIS DE CAÍDA DE TENSIÓN')
        e.parrafo('Datos en el punto de diseño:')
        p = datos.punto_diseno
        if p:
            e.tabla(['Vp.u.', 'kVLL', 'kVLN', 'i (A)', 'kVA', 'kW', 'kVAR'],
                    [[_f(p.get('u_pu'), 3), _f(p.get('kv_ll'), 2), _f(p.get('kv_ln'), 2), _f(p.get('i_a'), 1),
                      _f(p.get('kva'), 0), _f(p.get('kw'), 0), _f(p.get('kvar'), 0)]])
        else:
            e.hueco('datos en el punto de diseño')
        e.titulo('2.3.2 CARGA PARA CONSIDERAR:')
        if datos.solicitudes_vigentes:
            e.parrafo(SOLICITUDES_VIGENTES)
        _analisis_tension(e, '2.3.3 ANÁLISIS DE CAÍDA DE TENSIÓN EN EL ESTADO PROYECTADO', proyectado,
                          proy_img.get('tension'), 'estado proyectado', config, tabla=False)
        _analisis_carga(e, '2.3.4 ANÁLISIS DE CARGABILIDAD', proyectado, proy_img.get('cargabilidad'),
                        'estado proyectado', config)
        _tabla_tension(e, proyectado)

    _perdidas(e, datos, config, alimentador)

    e.titulo('2.5 ANÁLISIS DEL TRANSFORMADOR DE POTENCIA')
    d = datos.devanado
    if s is None:
        e.hueco('datos de la SET del formulario')
    else:
        e.parrafo(texto_transformador(s, d, datos.caso_transmision, config.proyecto_transmision or ''))
        pie = None
        if d is not None and d.con_pct is not None:
            pie = (f'Figura 1: Cargabilidad de {numero(d.con_pct, 2)} % en el devanado de '
                   f'{numero(d.kv, 0) if float(d.kv).is_integer() else numero(d.kv, 1)} kV de la SET '
                   f'{s.set_nombre} – {s.transformador}')
        e.imagen((datos.imagenes.get('transformador') or {}).get('cargabilidad'),
                 'cargabilidad del transformador de la SET', pie)

    e.titulo('3. CONCLUSIONES')
    veredicto = veredicto_factibilidad(proyectado, [d] if d else None, config.limites)
    if s is None:
        e.hueco('conclusiones: faltan los datos de la solicitud')
    else:
        for texto in conclusiones_factibilidad(s, veredicto, datos.codigo_punto or '', d):
            e.parrafo(texto)

    _guardar(documento, destino, {
        'tipo': 'factibilidad', 'numero': datos.numero,
        'encabezado': {**asdict(enc), 'fecha': enc.fecha.isoformat()},
        'solicitud': ({**asdict(s), 'fecha_solicitud': s.fecha_solicitud.isoformat()} if s else None),
        'actual': asdict(datos.actual), 'proyectado': asdict(proyectado),
        'punto_diseno': dict(datos.punto_diseno or {}), 'codigo_punto': datos.codigo_punto,
        'perdidas': {'actual': asdict(datos.perdidas_actual) if datos.perdidas_actual else None,
                     'proyectado': asdict(datos.perdidas_proyectado) if datos.perdidas_proyectado else None},
        'libro': asdict(datos.libro) if datos.libro else None,
        'transformador': asdict(d) if d else None, 'factible': veredicto.factible,
        'excesos': list(veredicto.excesos), 'imagenes': _rutas(datos.imagenes), 'huecos': e.huecos,
    })
    return e.huecos


def _perdidas(e: _Escritor, datos: DatosFactibilidad, config: ConfigInforme, alimentador: str) -> None:
    """Pérdidas reconocidas del libro y pérdidas de cada estado (RF-54 a RF-64)."""
    e.titulo('2.4 ANÁLISIS DE PÉRDIDAS TÉCNICAS')
    regulador = f' reconocido por {config.regulador}' if config.regulador else ' reconocido'
    e.parrafo(f'Porcentajes de pérdidas por sistema eléctrico{regulador}.')
    libro = datos.libro
    e.tabla(['Alimentador', 'Nombre sist. eléct.', 'Código sist. eléct.', 'Pérdidas téc. kW (%)',
             'Pérdidas téc. kWh (%)'],
            [[alimentador, (libro.sistema_electrico or '') if libro else '',
              (libro.codigo_sistema or '') if libro else '',
              _f(libro.reconocido_kw_pct if libro else None, 2),
              _f(libro.reconocido_kwh_pct if libro else None, 2)]])
    if libro is None:
        e.huecos.append('libro de pérdidas de la empresa')
    else:
        e.huecos.extend(libro.faltan)
    e.parrafo('Pérdidas obtenidas en los escenarios del alimentador.')
    filas = []
    for nombre, p in (('ESTADO ACTUAL', datos.perdidas_actual), ('ESTADO PROYECTADO', datos.perdidas_proyectado)):
        if p is None:
            e.huecos.append(f'pérdidas del {nombre.lower()}')
            filas.append([f'{nombre} {alimentador}'.strip(), '', '', '', '', ''])
            continue
        filas.append([f'{nombre} {alimentador}'.strip(), _f(p.maxima_kw, 2), _f(p.energia_kwh, 2),
                      _f(p.perdidas_kwh, 2), f'{numero(p.perdidas_pct, 2)} %', _f(p.no_reconocida_soles, 2)])
    e.tabla(['', 'MD kW', 'MD kWh', 'Pérdidas téc. kWh - año', 'Pérdidas téc. kWh (%)',
             'Pérdida/año no reconocida (S/.)'], filas)
