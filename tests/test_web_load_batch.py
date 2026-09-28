from __future__ import annotations

import io
import json
import sys
import time
from pathlib import Path

import pytest

pytest.importorskip('fastapi')
pytest.importorskip('httpx')

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from synthetic_export import ExportSpec, write_export  # noqa: E402
from igea_dgs.web.app import create_app  # noqa: E402
from igea_dgs.web import services  # noqa: E402
from igea_dgs.web.jobs import Job, JobCancelled, JobContext  # noqa: E402


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('IGEA_WEB_FRONTEND', str(tmp_path / 'sin_front'))
    with TestClient(create_app(tmp_path / 'web')) as value:
        yield value


def _wait(client: TestClient, job: dict, timeout: float = 30.0) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        state = client.get(f"/api/jobs/{job['id']}").json()
        if state['status'] in ('done', 'failed', 'cancelled'):
            return state
        time.sleep(0.02)
    raise AssertionError(f"job {job['id']} no terminó")


def _loaded(client: TestClient, tmp_path: Path) -> tuple[str, list[str]]:
    red, loads, equipment = write_export(
        ExportSpec(feeders=2, sections_per_feeder=3), tmp_path / 'export',
    )
    wid = client.post('/api/workspaces').json()['id']
    for slot, path in (('red', red), ('loads', loads), ('equipment', equipment)):
        with path.open('rb') as stream:
            response = client.post(
                f'/api/workspaces/{wid}/inputs/{slot}',
                files={'file': (path.name, stream)},
            )
        assert response.status_code == 200, response.text
    client.put(f'/api/workspaces/{wid}/options', json={'include_geography': False})
    state = _wait(client, client.post(f'/api/workspaces/{wid}/load').json())
    assert state['status'] == 'done', state
    feeders = [
        row['feeder']
        for row in client.get(f'/api/workspaces/{wid}/feeders').json()['feeders']
    ]
    return wid, feeders


def _multipart(feeders: list[str], filename: str, content: bytes):
    return [
        *(('feeder', (None, feeder)) for feeder in feeders),
        ('file', (filename, content)),
    ]


def _create_plan(client: TestClient, wid: str, feeders: list[str]) -> dict:
    template = client.get(
        f'/api/workspaces/{wid}/load-template',
        params=[*(('feeder', feeder) for feeder in feeders), ('format', 'xlsx')],
    )
    assert template.status_code == 200, template.text
    response = client.post(
        f'/api/workspaces/{wid}/load-batch-plan',
        files=_multipart(feeders, 'cargas.xlsx', template.content),
    )
    assert response.status_code == 200, response.text
    return response.json()


def _create_pf_plan(client: TestClient, wid: str, feeders: list[str]) -> dict:
    ws = client.app.state.store.get(wid)
    for feeder in feeders:
        ws.pf_projects[feeder] = f'PF_{feeder}'
    return _create_plan(client, wid, feeders)


def _fake_pf(monkeypatch, tmp_path: Path) -> list[list[str]]:
    commands = []
    monkeypatch.setattr(services, '_require_pf', lambda: (tmp_path, Path(sys.executable)))

    def run_process(self, cmd, **_kwargs):
        commands.append(cmd)
        output = Path(cmd[cmd.index('--output-json') + 1])
        plan = json.loads(Path(cmd[cmd.index('--plan') + 1]).read_text(encoding='utf-8'))
        dry_run = '--dry-run' in cmd
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps({
            'batch_id': plan['batch_id'],
            'dry_run': dry_run,
            'status': 'PASS',
            'feeders': [
                {
                    'feeder': item['feeder'], 'status': 'PASS',
                    'written': 0 if dry_run else len(item.get('updates') or []),
                    'errors': [], 'details': [],
                }
                for item in plan['feeders']
            ],
        }), encoding='utf-8')
        return 0

    monkeypatch.setattr(JobContext, 'run_process', run_process)
    return commands


def test_bulk_template_contains_all_selected_feeders(client, tmp_path):
    wid, feeders = _loaded(client, tmp_path)

    response = client.get(
        f'/api/workspaces/{wid}/load-template',
        params=[*(('feeder', feeder) for feeder in feeders), ('format', 'xlsx')],
    )

    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    try:
        assert workbook.sheetnames == sorted(feeders)
    finally:
        workbook.close()


def test_bulk_plan_processes_every_sheet(client, tmp_path):
    wid, feeders = _loaded(client, tmp_path)
    body = _create_plan(client, wid, feeders)
    assert body['token']
    assert body['feeders'] == sorted(feeders)
    assert set(body['summary_by_feeder']) == set(feeders)
    assert all(value['updates'] > 0 for value in body['summary_by_feeder'].values())


def test_bulk_csv_requires_feeder_for_multiple_models(client, tmp_path):
    wid, feeders = _loaded(client, tmp_path)
    csv_without_feeder = b'SED;Kw;Kvar;(kVA);FP\nSE1001;10;3;;\n'

    response = client.post(
        f'/api/workspaces/{wid}/load-batch-plan',
        files=_multipart(feeders, 'cargas.csv', csv_without_feeder),
    )

    assert response.status_code == 400
    assert 'Alimentador' in response.json()['detail']


def test_plan_rows_are_paginated_without_truncating_totals(client, tmp_path):
    wid, feeders = _loaded(client, tmp_path)
    plan = _create_plan(client, wid, feeders)
    endpoint = f"/api/workspaces/{wid}/load-batch-plans/{plan['token']}/rows"

    first = client.get(endpoint, params={'offset': 0, 'limit': 1}).json()
    second = client.get(endpoint, params={'offset': 1, 'limit': 1}).json()

    assert first['total'] > 1
    assert first['total'] == second['total']
    assert len(first['rows']) == len(second['rows']) == 1
    assert first['rows'][0] != second['rows'][0]


def test_input_change_invalidates_batch_plan(client, tmp_path):
    wid, feeders = _loaded(client, tmp_path)
    plan = _create_plan(client, wid, feeders)
    ws = client.app.state.store.get(wid)
    Path(ws.inputs['loads']['path']).write_text('cambio externo', encoding='utf-8')

    response = client.get(
        f"/api/workspaces/{wid}/load-batch-plans/{plan['token']}/rows"
    )

    assert response.status_code == 400
    assert 'obsoleto' in response.json()['detail'].lower()


def test_dgs_hash_change_invalidates_batch_plan(client, tmp_path):
    wid, feeders = _loaded(client, tmp_path)
    state = _wait(
        client,
        client.post(f'/api/workspaces/{wid}/convert', json={'feeders': feeders}).json(),
    )
    assert state['status'] == 'done', state
    plan = _create_plan(client, wid, feeders)
    ws = client.app.state.store.get(wid)
    conversion = next(iter(ws.conversions.values()))
    dgs = Path(conversion['dgs'])
    dgs.write_bytes(dgs.read_bytes() + b'\n# cambio externo\n')

    response = client.get(
        f"/api/workspaces/{wid}/load-batch-plans/{plan['token']}/rows"
    )

    assert response.status_code == 400
    assert 'obsoleto' in response.json()['detail'].lower()


def test_project_mapping_change_invalidates_batch_plan(client, tmp_path):
    wid, feeders = _loaded(client, tmp_path)
    ws = client.app.state.store.get(wid)
    ws.pf_projects[feeders[0]] = 'PROYECTO_A'
    plan = _create_plan(client, wid, feeders)
    ws.pf_projects[feeders[0]] = 'PROYECTO_B'

    response = client.get(
        f"/api/workspaces/{wid}/load-batch-plans/{plan['token']}/rows"
    )

    assert response.status_code == 400
    assert 'obsoleto' in response.json()['detail'].lower()


def test_apply_requires_successful_dry_run(client, tmp_path):
    wid, feeders = _loaded(client, tmp_path)
    plan = _create_pf_plan(client, wid, feeders)

    response = client.post(
        f"/api/workspaces/{wid}/load-batch-plans/{plan['token']}/apply"
    )

    assert response.status_code == 400
    assert 'dry-run' in response.json()['detail'].lower()


def test_dry_run_uses_powerfactory_lane(client, tmp_path, monkeypatch):
    commands = _fake_pf(monkeypatch, tmp_path)
    wid, feeders = _loaded(client, tmp_path)
    plan = _create_pf_plan(client, wid, feeders)

    response = client.post(
        f"/api/workspaces/{wid}/load-batch-plans/{plan['token']}/dry-run"
    )

    assert response.status_code == 202, response.text
    assert response.json()['lane'] == 'powerfactory'
    state = _wait(client, response.json())
    assert state['status'] == 'done', state
    assert '--dry-run' in commands[0]


def test_changed_plan_after_dry_run_is_rejected(client, tmp_path, monkeypatch):
    _fake_pf(monkeypatch, tmp_path)
    wid, feeders = _loaded(client, tmp_path)
    plan = _create_pf_plan(client, wid, feeders)
    dry = client.post(
        f"/api/workspaces/{wid}/load-batch-plans/{plan['token']}/dry-run"
    ).json()
    assert _wait(client, dry)['status'] == 'done'
    ws = client.app.state.store.get(wid)
    path = Path(ws.plans[plan['token']]['path'])
    path.write_text(path.read_text(encoding='utf-8') + '\n', encoding='utf-8')

    response = client.post(
        f"/api/workspaces/{wid}/load-batch-plans/{plan['token']}/apply"
    )

    assert response.status_code == 400
    assert 'obsoleto' in response.json()['detail'].lower()


def test_single_feeder_apply_uses_batch_engine(client, tmp_path, monkeypatch):
    commands = _fake_pf(monkeypatch, tmp_path)
    wid, feeders = _loaded(client, tmp_path)
    feeder = feeders[0]
    ws = client.app.state.store.get(wid)
    ws.pf_projects[feeder] = f'PF_{feeder}'
    template = client.get(
        f'/api/workspaces/{wid}/feeders/{feeder}/load-template',
        params={'format': 'csv'},
    )
    plan = client.post(
        f'/api/workspaces/{wid}/feeders/{feeder}/load-plan',
        files={'file': (f'{feeder}.csv', template.content)},
    ).json()
    dry = client.post(
        f"/api/workspaces/{wid}/load-batch-plans/{plan['token']}/dry-run"
    )
    assert _wait(client, dry.json())['status'] == 'done'

    apply = client.post(f"/api/workspaces/{wid}/plans/{plan['token']}/apply")
    assert _wait(client, apply.json())['status'] == 'done'

    assert all('apply_sed_load_batch.py' in ' '.join(cmd) for cmd in commands)


def test_audit_contains_before_requested_verified_and_status(client, tmp_path, monkeypatch):
    _fake_pf(monkeypatch, tmp_path)
    wid, feeders = _loaded(client, tmp_path)
    plan = _create_pf_plan(client, wid, feeders)
    job = client.post(
        f"/api/workspaces/{wid}/load-batch-plans/{plan['token']}/dry-run"
    ).json()

    state = _wait(client, job)

    assert state['status'] == 'done', state
    artifacts = state['result']['artifacts']
    assert set(artifacts) == {
        'input_manifest.json', 'plan.json', 'preview.csv', 'result.json',
        'result.csv', 'rollback.json', 'logs.jsonl',
    }
    ws = client.app.state.store.get(wid)
    result_csv = (ws.out_dir / artifacts['result.csv']).read_text(encoding='utf-8-sig')
    header = result_csv.splitlines()[0]
    assert {'p_before_mw', 'p_requested_mw', 'p_verified_mw', 'status'} <= set(
        header.split(';')
    )


def test_each_run_uses_a_new_directory(client, tmp_path, monkeypatch):
    _fake_pf(monkeypatch, tmp_path)
    wid, feeders = _loaded(client, tmp_path)
    plan = _create_pf_plan(client, wid, feeders)

    results = []
    for _ in range(2):
        job = client.post(
            f"/api/workspaces/{wid}/load-batch-plans/{plan['token']}/dry-run"
        ).json()
        results.append(_wait(client, job)['result'])

    assert results[0]['run_id'] != results[1]['run_id']
    assert Path(results[0]['report']).parent != Path(results[1]['report']).parent


def test_cancel_before_write_is_clean(client, tmp_path, monkeypatch):
    _fake_pf(monkeypatch, tmp_path)
    wid, feeders = _loaded(client, tmp_path)
    plan = _create_pf_plan(client, wid, feeders)
    ws = client.app.state.store.get(wid)
    job = Job('test', 'test', wid, 'powerfactory', lambda _ctx: None)
    job.cancel.set()
    ctx = JobContext(job, lambda *_args: None)

    with pytest.raises(JobCancelled):
        services.dry_run_load_batch(ws, ctx, plan['token'])

    assert not (ws.out_dir / 'cargas').exists()
