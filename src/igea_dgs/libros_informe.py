"""Libro de pérdidas de la empresa y tabla de costes de las obras (spec 017).

El informe de factibilidad compara las pérdidas técnicas del flujo con las que reconoce
el regulador, y esas vienen de un libro Excel de la empresa (``informeModeloperdidas.xlsx``
en Electro Dunas) que cada empresa ordena a su manera: por eso las hojas y columnas se
buscan por los nombres del perfil, nunca por posición (RF-63). Las etapas 2 y 3 del
informe de estudio (bancos y reguladores) necesitan costes; si la empresa no tiene
cotizaciones, se parte de una tabla de referencia con precios públicos, marcada
«referencial» para que el informe lo diga (RF-39d, RF-39e).
"""
from __future__ import annotations

import json
import math
import re
import unicodedata
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, Mapping

#: Tabla de costes de referencia versionada, con precios públicos de la web (RF-39d).
COSTES_REFERENCIALES = Path(__file__).resolve().parents[2] / 'config' / 'costes_referenciales.example.json'
#: Clave del perfil (dentro de ``informes``) con los nombres del libro de pérdidas.
CLAVE_LIBRO = 'libro_perdidas'
_COLUMNAS = ('alimentador', 'sistema_electrico', 'reconocido_kw_pct', 'reconocido_kwh_pct')
#: Columnas que el libro puede no traer; el informe deja la celda en blanco.
_OPCIONALES = ('codigo_sistema',)
#: Filas en las que se busca la cabecera: los libros de la empresa traen título encima.
_FILAS_CABECERA = 30


class LibroError(ValueError):
    """El libro o la tabla no tienen la forma que dice el perfil; el mensaje dice qué falta."""


def _clave(texto: Any) -> str:
    """Nombre comparable: sin tildes, sin mayúsculas y con los espacios colapsados.

    En los libros de la empresa la misma columna aparece como «% Reconocido (kW)» o con
    dos espacios en medio según quién la escribió.
    """
    sin_tildes = unicodedata.normalize('NFKD', str(texto)).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', sin_tildes).strip().casefold()


# --------------------------------------------------------------------------- libro de pérdidas
@dataclass(frozen=True)
class ColumnasPerdidas:
    """Dónde está cada dato en el libro de pérdidas, con los nombres del perfil (RF-63).

    El precio de compra puede ser una celda fija de otra hoja (``celda``) o una columna de
    la fila del alimentador (``precio_columna``).
    """

    hoja: str
    columnas: Mapping[str, str]
    precio_hoja: str | None = None
    precio_celda: str | None = None
    precio_columna: str | None = None

    @classmethod
    def de_perfil(cls, datos: Mapping[str, Any]) -> 'ColumnasPerdidas':
        prefijo = f'informes.{CLAVE_LIBRO}'
        seccion = (datos.get('informes') or {}).get(CLAVE_LIBRO)
        if not isinstance(seccion, Mapping):
            raise LibroError(f'Falta {prefijo} en el perfil de empresa: hoja y columnas del libro de pérdidas.')
        hoja = seccion.get('hoja')
        columnas = seccion.get('columnas') or {}
        faltan = ([f'{prefijo}.hoja'] if not isinstance(hoja, str) or not hoja.strip() else [])
        faltan += [f'{prefijo}.columnas.{c}' for c in _COLUMNAS
                   if not isinstance(columnas.get(c), str) or not columnas[c].strip()]
        precio = seccion.get('precio_compra') or {}
        if not (precio.get('columna') or (precio.get('hoja') and precio.get('celda'))):
            faltan.append(f'{prefijo}.precio_compra (columna, o hoja y celda)')
        if faltan:
            raise LibroError(f'Falta en el perfil de empresa: {", ".join(faltan)}.')
        return cls(hoja=hoja, columnas={c: columnas[c] for c in (*_COLUMNAS, *_OPCIONALES)
                                        if isinstance(columnas.get(c), str) and columnas[c].strip()},
                   precio_hoja=precio.get('hoja'), precio_celda=precio.get('celda'),
                   precio_columna=precio.get('columna'))


@dataclass(frozen=True)
class DatosPerdidas:
    """Lo que el libro dice de un alimentador; lo que no está queda en None con su motivo (RF-64)."""

    alimentador: str
    sistema_electrico: str | None
    reconocido_kw_pct: float | None
    reconocido_kwh_pct: float | None
    precio_compra: float | None
    faltan: tuple[str, ...]
    codigo_sistema: str | None = None


def _numero(celda) -> float | None:
    """Valor numérico de una celda; un % con formato de porcentaje se guarda como fracción."""
    valor = celda.value
    if isinstance(valor, str):
        try:
            valor = float(valor.replace('%', '').replace(',', '.').strip())
        except ValueError:
            return None
    if isinstance(valor, bool) or not isinstance(valor, (int, float)) or not math.isfinite(valor):
        return None
    return float(valor)


def _porcentaje(celda) -> float | None:
    valor = _numero(celda)
    if valor is not None and '%' in (celda.number_format or ''):
        valor *= 100.0
    return valor


def leer_libro_perdidas(ruta: Path, alimentador: str, columnas: ColumnasPerdidas) -> DatosPerdidas:
    """Sistema eléctrico, pérdidas reconocidas y precio de compra de ``alimentador`` (RF-62).

    Una hoja o columna que no está es un error que nombra lo que falta y lo que hay, para
    que el operador corrija el perfil (RF-94). Un alimentador que no está no lo es: el
    informe sale con las pérdidas calculadas y el hueco de lo reconocido (RF-64).
    """
    from openpyxl import load_workbook

    try:
        # data_only: el valor que Excel calculó, no la fórmula del libro.
        libro = load_workbook(Path(ruta), data_only=True, read_only=False)
    except Exception as exc:  # openpyxl lanza tipos distintos según lo que esté roto
        raise LibroError(f'No se pudo abrir el libro de pérdidas {Path(ruta).name}: {exc}') from exc
    try:
        hojas = {_clave(n): n for n in libro.sheetnames}

        def hoja(nombre: str):
            if _clave(nombre) not in hojas:
                raise LibroError(f'El libro de pérdidas no tiene la hoja {nombre!r}; '
                                 f'tiene {", ".join(repr(n) for n in libro.sheetnames)}.')
            return libro[hojas[_clave(nombre)]]

        datos = hoja(columnas.hoja)
        buscadas = dict(columnas.columnas)
        if columnas.precio_columna:
            buscadas['precio_compra'] = columnas.precio_columna
        cabecera, indices = None, {}
        for fila in datos.iter_rows(min_row=1, max_row=_FILAS_CABECERA):
            nombres = {_clave(c.value): c.column - 1 for c in fila if c.value not in (None, '')}
            if _clave(buscadas['alimentador']) in nombres:
                cabecera = fila
                indices = {k: nombres.get(_clave(v)) for k, v in buscadas.items()
                           if k not in _OPCIONALES or _clave(v) in nombres}
                break
        if cabecera is None or None in indices.values():
            faltan = [repr(buscadas[k]) for k, v in indices.items() if v is None] or [repr(buscadas['alimentador'])]
            vistas = [str(c.value).strip() for c in (cabecera or []) if c.value not in (None, '')]
            raise LibroError(f'La hoja {datos.title!r} del libro de pérdidas no tiene la columna '
                             f'{", ".join(faltan)}'
                             + (f'; tiene {", ".join(repr(v) for v in vistas)}.' if vistas else '.'))

        faltan: list[str] = []
        precio = None
        if columnas.precio_celda:
            precio = _numero(hoja(columnas.precio_hoja)[columnas.precio_celda])
            if precio is None:
                faltan.append(f'precio de compra en {columnas.precio_hoja}!{columnas.precio_celda}')

        buscado = _clave(alimentador)
        for fila in datos.iter_rows(min_row=cabecera[0].row + 1):
            if _clave(fila[indices['alimentador']].value or '') != buscado:
                continue
            sistema = fila[indices['sistema_electrico']].value
            valores = {k: _porcentaje(fila[indices[k]]) for k in ('reconocido_kw_pct', 'reconocido_kwh_pct')}
            faltan += [f'{columnas.columnas[k]} de {alimentador.strip()}' for k, v in valores.items() if v is None]
            if columnas.precio_columna:
                precio = _numero(fila[indices['precio_compra']])
                if precio is None:
                    faltan.append(f'{columnas.precio_columna} de {alimentador.strip()}')
            codigo = fila[indices['codigo_sistema']].value if 'codigo_sistema' in indices else None
            return DatosPerdidas(alimentador.strip(), str(sistema).strip() if sistema not in (None, '') else None,
                                 valores['reconocido_kw_pct'], valores['reconocido_kwh_pct'], precio,
                                 tuple(faltan),
                                 str(codigo).strip() if codigo not in (None, '') else None)
        faltan.insert(0, f'el alimentador {alimentador.strip()} no está en la hoja {datos.title!r}')
        return DatosPerdidas(alimentador.strip(), None, None, None, precio, tuple(faltan))
    finally:
        libro.close()


# --------------------------------------------------------------------------- tabla de costes
@dataclass(frozen=True)
class Coste:
    """Una fila de la tabla de costes: un banco (``kvar``), un regulador (``tipo``) o el precio."""

    inversion: float | None
    moneda: str
    fuente: str
    fecha: str
    referencial: bool
    mantenimiento_anual: float | None = None
    kvar: float | None = None
    tension_kv: float | None = None
    tipo: str | None = None
    clase_kv: float | None = None
    valor: float | None = None
    unidad: str | None = None
    nota: str | None = None


def _fila(datos: Any, donde: str, requeridos: tuple[str, ...]) -> Coste:
    if not isinstance(datos, Mapping):
        raise LibroError(f'{donde} de la tabla de costes tiene que ser un objeto.')
    faltan = [c for c in ('moneda', 'fuente', 'fecha', *requeridos)
              if datos.get(c) in (None, '')]
    if not isinstance(datos.get('referencial'), bool):
        faltan.append('referencial (true o false)')
    if faltan:
        raise LibroError(f'{donde} de la tabla de costes no trae {", ".join(faltan)}: '
                         'cada precio va con su fuente y su fecha (RF-39b).')
    campos = set(Coste.__dataclass_fields__)
    desconocidos = set(datos) - campos
    if desconocidos:
        raise LibroError(f'{donde} de la tabla de costes tiene campos desconocidos: {", ".join(sorted(desconocidos))}.')
    # El precio de la energía no tiene inversión: la fila lo trae en ``valor``.
    return Coste(**{'inversion': None, **{k: datos[k] for k in campos if k in datos}})


@dataclass(frozen=True)
class Costes:
    bancos: tuple[Coste, ...]
    reguladores: tuple[Coste, ...]
    precio_energia: Coste | None

    @property
    def origen(self) -> str:
        """«de la empresa», «referenciales» o «mixtos», para decirlo en el informe (RF-39e)."""
        marcas = {f.referencial for f in (*self.bancos, *self.reguladores, self.precio_energia) if f is not None}
        if marcas == {True}:
            return 'referenciales'
        return 'mixtos' if len(marcas) > 1 else 'de la empresa'

    def faltan_para_bancos(self) -> list[str]:
        """Lo que impide la etapa 2: sin bancos o sin precio de la energía no se ejecuta (RF-39c)."""
        return ((['bancos de condensadores'] if not self.bancos else [])
                + (['precio de la energía'] if self.precio_energia is None else []))

    def faltan_para_reguladores(self) -> list[str]:
        return ['reguladores de tensión'] if not self.reguladores else []

    def con_precio_energia(self, valor: float, unidad: str, moneda: str, fuente: str, fecha: str,
                           referencial: bool = False) -> 'Costes':
        """La tabla con el precio de compra del libro de pérdidas, si ella no lo trae."""
        return replace(self, precio_energia=Coste(inversion=None, moneda=moneda, fuente=fuente, fecha=fecha,
                                                  referencial=referencial, valor=valor, unidad=unidad))


def leer_costes(ruta: Path) -> Costes:
    """Bancos, reguladores y precio de la energía de una tabla de costes JSON (RF-39b, RF-40c)."""
    try:
        datos = json.loads(Path(ruta).read_text(encoding='utf-8'))
    except (OSError, ValueError) as exc:
        raise LibroError(f'No se pudo leer la tabla de costes {Path(ruta).name}: {exc}') from exc
    if not isinstance(datos, Mapping):
        raise LibroError('La tabla de costes tiene que ser un objeto JSON.')
    bancos = tuple(_fila(f, f'bancos[{i}]', ('kvar', 'inversion'))
                   for i, f in enumerate(datos.get('bancos') or [], 1))
    reguladores = tuple(_fila(f, f'reguladores[{i}]', ('tipo', 'inversion'))
                        for i, f in enumerate(datos.get('reguladores') or [], 1))
    precio = datos.get('precio_energia')
    return Costes(bancos, reguladores,
                  _fila(precio, 'precio_energia', ('valor', 'unidad')) if precio is not None else None)
