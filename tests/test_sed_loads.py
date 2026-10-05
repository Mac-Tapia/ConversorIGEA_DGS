"""Cargas de SED por libro Excel/CSV, en el formato real de la empresa.

Formato nativo (``referencia/PA217.xlsx``): **una hoja por alimentador**, el nombre de la
hoja es el alimentador, y las columnas son ``SED · Kw · Kvar · (kVA) · FP``.

Lo que se fija aquí:

- Cada fila trae un par: ``Kw``/``Kvar``, ``Kw``/``FP`` o ``(kVA)``/``FP``. El último
  es el del fichero real, donde las dos primeras columnas vienen en blanco.
- Si trae más de un par y no coinciden, es un error de fila: nunca manda uno en silencio.
- La plantilla deja vacías las columnas de entrada; una fila en blanco no toca la SED.
- ``FP = 0`` y ``kVA = 0`` no es un error: es una SED sin dato, y se cuenta aparte.
- Una sola fila inválida bloquea toda la actualización.
- SED de la hoja que no están en el modelo → ``unknown``, nunca se inventan.
- SED del modelo que la hoja no menciona → ``untouched``.
"""

from __future__ import annotations

import csv
import math
from pathlib import Path

import pytest

from igea_dgs.loads import (
    ACTION_SKIP,
    EXTRA_COLUMNS,
    SHEET_COLUMNS,
    LoadTemplateError,
    build_plan,
    canonical_columns,
    model_sed_loads,
    plan_to_payload,
    read_sheet,
    read_workbook,
    write_template,
)
from igea_dgs.model import FeederModel, Load, Node, Sed

REAL_XLSX = Path(__file__).resolve().parents[1] / 'referencia' / 'PA217.xlsx'


def _model(n: int = 3, name: str = 'AL01') -> FeederModel:
    nodes = {f'N{i}': Node(f'N{i}', float(i) * 10, 0.0) for i in range(n + 1)}
    loads, seds = [], []
    for i in range(1, n + 1):
        code = f'SE{1000 + i}'
        section, device = f'SEC_{i}', f'DEV_{i}_{code}'
        loads.append(Load(
            section_id=section, device_number=device, customer_number=f'CUST_{code}',
            location='1', node_id=f'N{i}', p_mw=0.010 * i, q_mvar=0.003 * i,
            pf=0.95, connected_kva=50.0 * i, kwh=0.0, phase='ABC',
            sed_code=code, display_name=code,
        ))
        seds.append(Sed(
            code=code, loc_name=code, node_id=f'N{i}', design_kva=50.0 * i,
            section_id=section, device_number=device, load_key=(section, device),
        ))
    return FeederModel(
        name=name, network_id=f'NET_{name}', nominal_kv=10.0, source_node='N0',
        nodes=nodes, lines=[], loads=loads, devices=[], line_types={}, seds=seds,
    )


class TestColumnRecognition:
    def test_native_header_is_recognised(self):
        mapping = canonical_columns(list(SHEET_COLUMNS))
        assert set(mapping.values()) == set(SHEET_COLUMNS)

    def test_spelling_variants_are_tolerated(self):
        """Un fichero con «P (kW)» o «cos phi» debe entenderse igual."""
        mapping = canonical_columns(['SED', 'P (kW)', 'Q kvar', 'S kVA', 'cos phi'])
        assert mapping[1] == 'Kw'
        assert mapping[2] == 'Kvar'
        assert mapping[3] == '(kVA)'
        assert mapping[4] == 'FP'

    @pytest.mark.parametrize('cabecera', ['cos fi', 'Cos φ', 'FdP'])
    def test_power_factor_spellings(self, cabecera):
        assert canonical_columns(['SED', cabecera])[1] == 'FP'

    def test_unknown_columns_are_ignored(self):
        mapping = canonical_columns(['SED', 'Kw', 'zona', 'comentario'])
        assert set(mapping.values()) == {'SED', 'Kw'}

    def test_sheet_without_sed_column_is_rejected(self):
        with pytest.raises(LoadTemplateError, match='código de SED'):
            read_sheet(['nombre', 'Kw'], [['X', 1]], 'AL01')


class TestPowerDerivation:
    def test_kw_kvar_are_derived_from_kva_and_fp(self):
        """El caso del fichero real: Kw y Kvar vienen vacías."""
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '', '', '0.3633217774', '0.936182876']], 'AL01')
        assert sheet.errors == []
        row = sheet.rows[0]
        assert row.derived is True
        assert row.kw == pytest.approx(0.3401, abs=1e-4)
        assert row.kvar == pytest.approx(0.1277, abs=1e-4)
        # El módulo debe reconstruir el kVA de partida.
        assert math.hypot(row.kw, row.kvar) == pytest.approx(0.3633217774, rel=1e-9)

    def test_kw_and_power_factor_give_kvar(self):
        """Medidor que da kW y cos φ, sin reactiva."""
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '46', '', '', '0.92']], 'AL01')
        assert sheet.errors == []
        row = sheet.rows[0]
        assert row.derived is True
        assert row.kw == pytest.approx(46.0)
        assert row.kvar == pytest.approx(46.0 * math.tan(math.acos(0.92)))
        assert row.kva == pytest.approx(50.0)

    def test_kw_with_unit_power_factor_has_no_reactive(self):
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '30', '', '', '1']], 'AL01')
        assert sheet.rows[0].kvar == pytest.approx(0.0)

    def test_disagreeing_pairs_are_an_error_not_a_silent_choice(self):
        """Antes Kw/Kvar mandaban y el (kVA)/FP escrito se perdía sin aviso."""
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '100', '30', '999', '0.5']], 'AL01')
        assert sheet.rows == []
        assert len(sheet.errors) == 1
        assert 'fila 2' in sheet.errors[0] and 'no coincide' in sheet.errors[0]

    def test_kw_fp_against_a_different_kva_is_an_error(self):
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '46', '', '80', '0.92']], 'AL01')
        assert any('(kVA)' in e for e in sheet.errors)

    def test_agreeing_pairs_within_rounding_are_accepted(self):
        """kW 46 y kvar 19,6 dan FP 0,9200; una hoja con FP a dos decimales cuadra."""
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '46', '19.6', '50', '0.92']], 'AL01')
        assert sheet.errors == []
        assert sheet.rows[0].kw == pytest.approx(46.0)

    def test_only_kw_given_leaves_q_at_zero(self):
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '80', '', '', '']], 'AL01')
        row = sheet.rows[0]
        assert row.kw == pytest.approx(80.0)
        assert row.kvar == pytest.approx(0.0)

    def test_zero_kva_and_zero_fp_is_no_data_not_an_error(self):
        """En el fichero real son 59 filas: sin dato, no un fallo."""
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '', '', '0', '0']], 'AL01')
        assert sheet.errors == []
        assert sheet.no_data == ['SE1001']
        assert sheet.rows[0].kw == 0.0 and sheet.rows[0].kvar == 0.0

    def test_comma_decimal_separator(self):
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '', '', '1,5', '0,9']], 'AL01')
        assert sheet.errors == []
        assert sheet.rows[0].kw == pytest.approx(1.35)


class TestRowValidation:
    def _read(self, row):
        return read_sheet(list(SHEET_COLUMNS) + ['accion'], [row], 'AL01')

    def test_non_numeric_is_reported_with_sheet_and_row(self):
        sheet = self._read(['SE1001', 'mucho', '', '', '', ''])
        assert len(sheet.errors) == 1
        assert 'hoja AL01' in sheet.errors[0] and 'fila 2' in sheet.errors[0]
        assert 'Kw' in sheet.errors[0]

    def test_power_factor_out_of_range(self):
        sheet = self._read(['SE1001', '', '', '10', '1.4', ''])
        assert any('FP' in e and 'fuera de' in e for e in sheet.errors)

    def test_negative_value_is_reported(self):
        sheet = self._read(['SE1001', '-5', '', '', '', ''])
        assert any('negativo' in e for e in sheet.errors)

    def test_duplicate_sed_is_reported(self):
        """El fichero real trae dos: SE30800 y SE30583."""
        sheet = read_sheet(
            list(SHEET_COLUMNS),
            [['SE1001', '', '', '1', '0.9'], ['SE1001', '', '', '2', '0.9']],
            'AL01',
        )
        assert any('repetida' in e for e in sheet.errors)
        assert len(sheet.rows) == 1

    def test_blank_rows_are_skipped_silently(self):
        sheet = read_sheet(list(SHEET_COLUMNS), [['', '', '', '', ''], ['SE1001', '', '', '1', '0.9']], 'AL01')
        assert sheet.errors == []
        assert len(sheet.rows) == 1

    def test_invalid_action_is_reported(self):
        sheet = self._read(['SE1001', '', '', '1', '0.9', 'borrar'])
        assert any('accion' in e for e in sheet.errors)

    def test_skip_action_excludes_the_sed(self):
        sheet = self._read(['SE1001', '', '', '1', '0.9', ACTION_SKIP])
        assert sheet.rows == []
        assert sheet.skipped == ['SE1001']


class TestTemplateGeneration:
    def test_xlsx_has_one_sheet_per_feeder_in_the_native_format(self, tmp_path):
        from openpyxl import load_workbook

        feeders = {'PA217': model_sed_loads(_model(3, 'PA217')),
                   'IN111': model_sed_loads(_model(2, 'IN111'))}
        path = write_template(feeders, tmp_path / 'cargas.xlsx')
        wb = load_workbook(path)
        assert wb.sheetnames == ['PA217', 'IN111']
        header = [c.value for c in wb['PA217'][1]]
        assert header[:len(SHEET_COLUMNS)] == list(SHEET_COLUMNS)
        assert header[len(SHEET_COLUMNS):] == list(EXTRA_COLUMNS)
        assert wb['PA217'].max_row == 4          # cabecera + 3 SED
        assert wb['IN111'].max_row == 3

    def test_untouched_template_changes_nothing(self, tmp_path):
        """Subir la plantilla tal cual no puede poner a cero ninguna carga."""
        model = _model(4, 'PA217')
        path = write_template({'PA217': model_sed_loads(model)}, tmp_path / 'c.xlsx')
        libro = read_workbook(path)
        assert list(libro) == ['PA217']
        assert libro['PA217'].rows == [] and libro['PA217'].errors == []
        plan = build_plan(model, libro['PA217'])
        assert plan.updates == []
        assert plan.untouched == sorted(s.code for s in model.seds)

    def test_template_shows_the_current_load_outside_the_entry_columns(self, tmp_path):
        from openpyxl import load_workbook

        path = write_template({'AL01': model_sed_loads(_model(1))}, tmp_path / 'c.xlsx')
        ws = load_workbook(path)['AL01']
        fila = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
        assert all(fila[name] is None for name in ('Kw', 'Kvar', '(kVA)', 'FP'))
        assert fila['Kw_actual'] == pytest.approx(10.0)
        assert fila['Kvar_actual'] == pytest.approx(3.0)
        assert fila['FP_actual'] == pytest.approx(0.95)

    def test_filling_only_kva_and_fp_in_the_template_is_applied(self, tmp_path):
        """El caso que fallaba: el operador escribe kVA y FP y su dato se ignoraba."""
        from openpyxl import load_workbook

        model = _model(2)
        path = write_template({'AL01': model_sed_loads(model)}, tmp_path / 'c.xlsx')
        wb = load_workbook(path)
        ws = wb['AL01']
        ws['D2'], ws['E2'] = 50, 0.92                  # SE1001: (kVA) y FP
        wb.save(path)
        plan = build_plan(model, read_workbook(path)['AL01'])
        assert [n.sed_code for _a, n in plan.updates] == ['SE1001']
        assert plan.updates[0][1].kw == pytest.approx(46.0)
        assert plan.untouched == ['SE1002']

    def test_csv_carries_the_feeder_in_a_column(self, tmp_path):
        feeders = {'PA217': model_sed_loads(_model(2, 'PA217')),
                   'IN111': model_sed_loads(_model(2, 'IN111'))}
        path = write_template(feeders, tmp_path / 'cargas.csv')
        with path.open(encoding='utf-8-sig', newline='') as fh:
            rows = list(csv.reader(fh, delimiter=';'))
        assert rows[0][0] == 'feeder'
        libro = read_workbook(path)
        assert set(libro) == {'PA217', 'IN111'}

    def test_list_instead_of_dict_is_accepted(self, tmp_path):
        path = write_template(model_sed_loads(_model(2, 'AL01')), tmp_path / 'c.xlsx')
        assert list(read_workbook(path)) == ['AL01']

    def test_unsupported_extension_is_rejected(self, tmp_path):
        with pytest.raises(LoadTemplateError, match='no soportada'):
            write_template(model_sed_loads(_model(1)), tmp_path / 'c.txt')


class TestWorkbookReading:
    def test_empty_sheets_are_ignored(self, tmp_path):
        """El fichero real trae IN111 y PA219 vacías: no deben ser un error."""
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet('PA217')
        ws.append(list(SHEET_COLUMNS))
        ws.append(['SE1001', '', '', '1.5', '0.9'])
        wb.create_sheet('IN111')            # totalmente vacía
        solo_cabecera = wb.create_sheet('PA219')
        solo_cabecera.append(list(SHEET_COLUMNS))
        path = tmp_path / 'c.xlsx'
        wb.save(path)

        libro = read_workbook(path)
        assert list(libro) == ['PA217']

    def test_a_broken_sheet_does_not_sink_the_workbook(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        buena = wb.create_sheet('PA217')
        buena.append(list(SHEET_COLUMNS))
        buena.append(['SE1001', '', '', '1.5', '0.9'])
        ajena = wb.create_sheet('NOTAS')
        ajena.append(['esto', 'no', 'es', 'una', 'hoja de cargas'])
        ajena.append(['x', 'y', 'z', '', ''])
        path = tmp_path / 'c.xlsx'
        wb.save(path)

        libro = read_workbook(path)
        assert libro['PA217'].errors == []
        assert libro['NOTAS'].errors, 'la hoja ajena debe reportar su problema'

    def test_missing_file_is_reported(self, tmp_path):
        with pytest.raises(LoadTemplateError, match='No se encuentra'):
            read_workbook(tmp_path / 'no-existe.xlsx')

    def test_workbook_without_data_is_reported(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        path = tmp_path / 'vacio.xlsx'
        wb.save(path)
        with pytest.raises(LoadTemplateError, match='ninguna hoja con datos'):
            read_workbook(path)


class TestPlanDiff:
    def _sheet(self, rows, feeder='AL01'):
        return read_sheet(list(SHEET_COLUMNS), rows, feeder)

    def test_existing_seds_are_updated(self):
        model = _model(3)
        sheet = self._sheet([['SE1001', '120', '40', '', '']])
        plan = build_plan(model, sheet)
        assert len(plan.updates) == 1
        actual, nueva = plan.updates[0]
        assert actual.kw == pytest.approx(10.0)      # 0,010 MW
        assert nueva.kw == pytest.approx(120.0)
        # La identidad la pone el modelo, no la hoja.
        assert nueva.section_id == 'SEC_1'
        assert nueva.node_id == 'N1'

    def test_unknown_sed_is_not_invented(self):
        model = _model(2)
        sheet = self._sheet([['SE1001', '10', '3', '', ''], ['5SE9999', '25', '8', '', '']])
        plan = build_plan(model, sheet)
        assert [r.sed_code for r in plan.unknown] == ['5SE9999']
        assert len(plan.updates) == 1
        assert 'módulo de creación' in plan.report()

    def test_sed_absent_from_the_sheet_is_untouched(self):
        model = _model(3)
        plan = build_plan(model, self._sheet([['SE1001', '10', '3', '', '']]))
        assert plan.untouched == ['SE1002', 'SE1003']

    def test_one_bad_row_blocks_the_whole_update(self):
        model = _model(3)
        sheet = self._sheet([['SE1001', 'mucho', '', '', ''], ['SE1002', '10', '3', '', '']])
        plan = build_plan(model, sheet)
        assert plan.row_errors
        assert plan.is_applicable is False

    def test_summary_counts_every_category(self):
        model = _model(4)
        sheet = read_sheet(
            list(SHEET_COLUMNS) + ['accion'],
            [
                ['SE1001', '', '', '1', '0.9', ACTION_SKIP],
                ['SE1002', '50', '15', '', '', ''],
                ['SE1003', '', '', '0', '0', ''],          # sin datos
                ['SE9999', '10', '3', '', '', ''],          # no está en el modelo
            ],
            'AL01',
        )
        plan = build_plan(model, sheet)
        assert plan.summary() == {
            'feeder': 'AL01', 'updates': 2, 'unknown': 1, 'untouched': 1,
            'skipped': 1, 'no_data': 1, 'row_errors': 0, 'applicable': True,
        }


class TestPayloadForPowerFactory:
    def test_units_are_converted_to_mw_for_powerfactory(self):
        model = _model(2)
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '120', '50', '', '']], 'AL01')
        payload = plan_to_payload(build_plan(model, sheet))
        item = payload['updates'][0]
        assert item['plini_mw'] == pytest.approx(0.120)
        assert item['qlini_mvar'] == pytest.approx(0.050)
        assert item['slini_mva'] == pytest.approx(math.hypot(0.120, 0.050))
        assert item['previous']['plini_mw'] == pytest.approx(0.010)

    def test_derived_flag_travels_to_the_payload(self):
        model = _model(1)
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '', '', '10', '0.9']], 'AL01')
        payload = plan_to_payload(build_plan(model, sheet))
        assert payload['updates'][0]['derived_from_kva_fp'] is True

    def test_payload_is_json_serialisable(self):
        import json

        model = _model(3)
        sheet = read_sheet(list(SHEET_COLUMNS), [['SE1001', '1', '1', '', '']], 'AL01')
        payload = plan_to_payload(build_plan(model, sheet))
        assert json.loads(json.dumps(payload))['summary']['updates'] == 1


@pytest.mark.skipif(not REAL_XLSX.is_file(), reason='referencia/PA217.xlsx no está presente')
class TestRealCompanyWorkbook:
    """El libro real de la empresa, con sus problemas de datos incluidos."""

    def test_the_real_workbook_is_read(self):
        libro = read_workbook(REAL_XLSX)
        assert 'PA217' in libro
        sheet = libro['PA217']
        assert len(sheet.rows) > 100
        # Kw/Kvar vienen vacías: la mayoría de cargas deben salir derivadas.
        assert sum(1 for r in sheet.rows if r.derived) > 50

    def test_it_reports_the_duplicate_seds_it_contains(self):
        sheet = read_workbook(REAL_XLSX)['PA217']
        duplicados = [e for e in sheet.errors if 'repetida' in e]
        assert duplicados, 'el fichero real tiene SED repetidas y deben detectarse'

    def test_plan_against_the_real_model(self, ds):
        from igea_dgs.model import build_feeder_model

        try:
            model = build_feeder_model(ds, 'PA217', strict=False)
        except KeyError:
            pytest.skip('PA217 no está en el export cargado')
        plan = build_plan(model, read_workbook(REAL_XLSX)['PA217'])
        # La mayoría de las SED de la hoja existen en el modelo.
        assert len(plan.updates) > 150
        # Y el fichero trae una SED que no existe: '5SE31022', un SE31022 con un 5 de más.
        assert any(r.sed_code.startswith('5') for r in plan.unknown)
        # Las SED con prefijo M del modelo no están en la hoja.
        assert plan.untouched
        # Con filas repetidas, el plan NO es aplicable hasta corregirlas.
        assert plan.is_applicable is False


class TestSplitByFeeder:
    """Un solo libro con varios alimentadores: cada SED va al alimentador que la tiene."""

    def test_rows_follow_the_sed_not_the_sheet(self):
        from igea_dgs.loads import split_by_feeder

        libro = {
            'TODO': read_sheet(list(SHEET_COLUMNS), [
                ['SE1001', '10', '3', '', ''],          # es de AL01
                ['SE2001', '20', '5', '', ''],          # es de AL02
            ], 'TODO'),
        }
        out, sueltos = split_by_feeder(libro, {'SE1001': 'AL01', 'SE2001': 'AL02'}, ['AL02', 'AL01'])
        assert sueltos == []
        assert [r.sed_code for r in out['AL01'].rows] == ['SE1001']
        assert [r.sed_code for r in out['AL02'].rows] == ['SE2001']
        assert out['AL02'].rows[0].feeder == 'AL02'

    def test_unknown_sed_stays_in_the_feeder_of_its_sheet(self):
        from igea_dgs.loads import split_by_feeder

        libro = {'AL01': read_sheet(list(SHEET_COLUMNS), [['5SE9999', '1', '1', '', '']], 'AL01')}
        out, sueltos = split_by_feeder(libro, {}, ['AL01'])
        assert [r.sed_code for r in out['AL01'].rows] == ['5SE9999']
        assert sueltos == []

    def test_unknown_sed_in_a_foreign_sheet_cannot_be_assigned(self):
        from igea_dgs.loads import split_by_feeder

        libro = {'NOTAS': read_sheet(list(SHEET_COLUMNS), [['5SE9999', '1', '1', '', '']], 'NOTAS')}
        _out, sueltos = split_by_feeder(libro, {}, ['AL01'])
        assert len(sueltos) == 1 and '5SE9999' in sueltos[0]

    def test_same_sed_in_two_sheets_is_an_error(self):
        from igea_dgs.loads import split_by_feeder

        fila = ['SE1001', '10', '3', '', '']
        libro = {h: read_sheet(list(SHEET_COLUMNS), [fila], h) for h in ('A', 'B')}
        out, _ = split_by_feeder(libro, {'SE1001': 'AL01'}, ['AL01'])
        assert any('más de una hoja' in e for e in out['AL01'].errors)

    def test_split_plans_match_each_model(self):
        from igea_dgs.loads import split_by_feeder

        m1, m2 = _model(2, 'AL01'), _model(2, 'AL02')
        # Los dos modelos sintéticos usan los mismos códigos: se renombran los de AL02.
        for sed in m2.seds:
            object.__setattr__(sed, 'code', sed.code.replace('SE1', 'SE2'))
        sed_feeder = {s.code: m.name for m in (m1, m2) for s in m.seds}
        libro = {'X': read_sheet(list(SHEET_COLUMNS), [['SE2001', '40', '10', '', '']], 'X')}
        out, _ = split_by_feeder(libro, sed_feeder, ['AL01', 'AL02'])
        assert build_plan(m2, out['AL02']).updates[0][1].kw == pytest.approx(40.0)
        assert build_plan(m1, out['AL01']).updates == []
