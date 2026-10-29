"""022:RF-56 — ``tools/sistema_web.py``: un sistema eléctrico de punta a punta por la web."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from apoyo import cargar_herramienta
from openpyxl import Workbook


@pytest.fixture(scope='module')
def herramienta():
    return cargar_herramienta('sistema_web.py', 'sistema_web_prueba')


class WebFalsa:
    """Responde como la web: el espacio, los trabajos y la demanda de un sistema."""

    def __init__(self, proyecto=None, informes=None):
        self.llamadas: list[tuple[str, str]] = []
        self.proyecto = proyecto
        self.informes = informes or []
        self.esperados: list[str] = []

    def get(self, ruta, **kw):
        self.llamadas.append(('GET', ruta))
        if ruta == '':
            return {'loaded': True, 'jobs': [{'id': 'v', 'title': 'viejo', 'status': 'running'}],
                    'pf_projects': {'SISTEMA_PISCO': self.proyecto} if self.proyecto else {}}
        if ruta == '/demanda/sistema':
            return {'sistema': {'grupo': 'SISTEMA_PISCO', 'alimentadores': ['PI102', 'PI101'],
                                'medidores': {'PI102': ['L-PI102'], 'PI101': []}}}
        if ruta == '/demanda/centrales':                                 # 022:RF-88
            return {'centrales': [
                {'medidor': 'CT_X', 'kv': 22.9, 'set': 'SET 31', 'tramo': 'PI102'},
                {'medidor': 'CT_Y', 'kv': 10.0, 'set': 'SET 31', 'tramo': ''}]}
        raise AssertionError(ruta)

    def post(self, ruta, **kw):
        self.llamadas.append(('POST', ruta))
        if ruta == '/demanda/cabecera/buscar':
            # 022:RF-88: la central del perfil va en el tramo; la otra, en su barra.
            assert kw['json']['centrales_tramo'] == {'CT_X': 'PI102'}, kw['json']
        if ruta == '/powerfactory/flow':
            self.proyecto = 'P_PISCO'
        if ruta == '/demanda/ajuste/plan':
            assert kw['json']['medidores'] == ['L-PI102']
            return {'token': 't', 'avisos': ['PI101 sin medidor']}
        return {'id': ruta, 'title': ruta}

    def esperar(self, trabajo):
        self.esperados.append(trabajo['id'])
        if trabajo['id'] == '/convert/sistema':
            return {'status': 'done', 'result': {'status': 'ok', 'group': 'SISTEMA_PISCO'}}
        if trabajo['id'] == '/demanda/cabecera/buscar':
            return {'id': 'max', 'status': 'done', 'result': {
                'barras': [{'clave': 'SET 31 10 kV', 'kw': 5456.3, 'fecha': '2026-07-21 18:45',
                            'alimentadores': ['PI101', 'PI102']}],
                'resultados': [{'medidor': 'L-PI102', 'alimentador': 'PI102', 'kw': 1000.0},
                               {'medidor': 'L-PI141', 'error': 'sin dato'},
                               # 022:RF-88: la central en tramo, no al plan.
                               {'medidor': 'CT_X', 'alimentador': 'Central PI102', 'kw': 50.0}]}}
        if trabajo['id'].endswith('/aplicar'):
            return {'status': 'done', 'result': {
                'sistema': {'sin_medidor': ['PI101']}, 'informes': self.informes}}
        return {'status': 'done', 'result': {}}


LIBROS = {'medidores': Path('libros/m.xlsx'), 'cabecera': Path('libros/cab'),
          'clientes': Path('libros/cli'), 'suministros': Path('libros/cualquier-nombre.xlsx')}
T_2025 = datetime(2025, 3, 1, 10, 15)


def _correr(herramienta, web, **kw):
    return herramienta.correr(web, 'Pisco', demanda_dir=Path('libros'),
                              periodo=kw.get('periodo', '2026-07'),
                              reconvertir=kw.get('reconvertir', False), transformadores=False,
                              tension_pu=None, bancos=True, libros=LIBROS)


def test_con_un_anio_entero_no_cruza_un_mes_y_sube_cada_libro(herramienta):
    """025:RF-25 — ``--periodo 2025``: la máxima del año; el cruce del plan va por el mes de
    cada máxima, así que no se manda un cruce de mes; cada libro va a su hueco y la
    cabecera se indexa por año."""
    web = WebFalsa(informes=[VALIDO])
    assert _correr(herramienta, web, periodo='2025') == 0
    rutas = [r for _m, r in web.llamadas]
    assert '/demanda/clientes/cruzar' not in rutas
    assert all(f'/demanda/{t}/path' in rutas for t in LIBROS)
    assert '/demanda/cabecera/indexar' in rutas


def test_los_libros_se_reconocen_por_su_contenido(herramienta, tmp_path):
    """025:RF-25b, RF-25c — sin nombres de carpeta de ninguna empresa; lo dado por opción
    manda; sin medición de cabecera no se sigue."""
    def libro(path, filas):
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        for f in filas:
            wb.active.append(f)
        wb.save(path)

    libro(tmp_path / 'x' / 'rel.xlsx', [['Codigo Alimentador', 'Medidor'], ['F1', 'M1']])
    libro(tmp_path / 'y' / '2025' / 'm.xlsx', [['Local Time', 'kW tot'], [T_2025, 1.0]])
    (tmp_path / 'z').mkdir()
    _libro_suministros(tmp_path / 'z' / 'padron.xlsx', 2)
    libros, _avisos = herramienta.libros_de_la_demanda(tmp_path)
    assert libros['medidores'] == tmp_path / 'x' / 'rel.xlsx'
    assert libros['cabecera'] == tmp_path / 'y'
    assert libros['suministros'] == tmp_path / 'z' / 'padron.xlsx'
    otro = tmp_path / 'otra cabecera'
    libros, _ = herramienta.libros_de_la_demanda(tmp_path, {'cabecera': otro})
    assert libros['cabecera'] == otro
    vacia = tmp_path / 'vacia'
    vacia.mkdir()
    with pytest.raises(herramienta.DemandaError, match='medición de cabecera'):
        herramienta.libros_de_la_demanda(vacia)


def _libro_suministros(path: Path, cantidad: int) -> None:
    libro = Workbook()
    hoja = libro.active
    hoja.append(['Suministro Act', 'SED', 'RADIAL'])
    for indice in range(cantidad):
        hoja.append([f'NIS-{indice}', f'SED-{indice}', f'AL-{indice}'])
    libro.save(path)


def test_el_libro_de_suministros_se_descubre_por_contenido(herramienta, tmp_path):
    """004:RF-58/RF-59 — cualquier nombre sirve y, con varios compatibles, se usa el que contiene
    más suministros actuales en vez de fijar un nombre o bloquear la ejecución."""
    carpeta = tmp_path / '3.-suministrocliente'
    carpeta.mkdir()
    _libro_suministros(carpeta / 'copia-con-nombre-arbitrario.xlsx', 1)
    vigente = carpeta / 'entrega-sin-fecha.xlsx'
    _libro_suministros(vigente, 3)
    libro_invalido = Workbook()
    libro_invalido.save(carpeta / 'notas.xlsx')

    assert herramienta.resolver_libro_suministros(tmp_path) == vigente


VALIDO = {'feeder': 'PI102', 'cuadrado': True, 'status': 'OK', 'flujos': 5,
          'desvio': {'kw': 0.0}, 'validacion_posterior_nativa': {'status': 'VALIDO'}}


def test_convierte_importa_y_calibra_el_sistema(herramienta, capsys):
    """022:RF-56 — sin proyecto del sistema: espera lo que corre, convierte, importa y
    calibra con los medidores que dio la máxima; dice los sin medidor."""
    web = WebFalsa(informes=[VALIDO])
    assert _correr(herramienta, web) == 0
    assert web.esperados[0] == 'v', 'primero lo que estaba en curso'
    rutas = [r for _m, r in web.llamadas]
    assert rutas.index('/convert/sistema') < rutas.index('/powerfactory/flow') < rutas.index(
        '/demanda/ajuste/plan') < rutas.index('/demanda/ajuste/t/aplicar')
    salida = capsys.readouterr().out
    assert 'Proyecto: P_PISCO' in salida and 'Sin medidor (demanda del archivo): PI101' in salida
    assert '1 válido(s), 0 con fallos.' in salida


def test_retoma_el_proyecto_ya_importado(herramienta):
    """022:RF-56 — con el proyecto del sistema ya en el espacio no se vuelve a convertir; con
    ``--reconvertir``, sí."""
    web = WebFalsa(proyecto='P_PISCO', informes=[VALIDO])
    assert _correr(herramienta, web) == 0
    assert ('POST', '/convert/sistema') not in web.llamadas
    web = WebFalsa(proyecto='P_PISCO', informes=[VALIDO])
    _correr(herramienta, web, reconvertir=True)
    assert ('POST', '/convert/sistema') in web.llamadas


def test_con_fallos_lo_dice_y_sale_con_1(herramienta, capsys):
    """022:RF-56 — un alimentador con remedios y validación que falla da código 1 y su
    primer flujo con lo que dijo DIgSILENT."""
    malo = {'feeder': 'PI102', 'cuadrado': False, 'status': 'FALLO_VALIDACION_NATIVA',
            'flujos': 18, 'desvio': {'kw': 0.0},
            'remedios': [{'remedio': 'adaptación automática del modelo'}],
            'historial': [{'converge': False, 'motivo': 'Newton algorithm stagnated'}],
            'validacion_posterior_nativa': {'status': 'REQUIERE_DISTRIBUCION'}}
    assert _correr(herramienta, WebFalsa(proyecto='P', informes=[malo])) == 1
    salida = capsys.readouterr().out
    assert 'Newton algorithm stagnated' in salida and 'adaptación automática' in salida


def test_la_web_que_corta_se_reintenta(herramienta, monkeypatch):
    """022:RF-56 — un corte de la conexión mientras la web trabaja no corta la cadena
    (Pisco, 2026-10-18: «WinError 10054» a mitad de la conversión)."""
    import requests

    web = herramienta.Web('http://web', 'e', reintentos=3, pausa=0)
    respuestas = [requests.ConnectionError('10054'),
                  type('R', (), {'status_code': 200, 'json': lambda self: {'ok': 1}})()]
    monkeypatch.setattr(web.http, 'request', lambda *a, **k: (
        (_ for _ in ()).throw(respuestas.pop(0)) if isinstance(respuestas[0], Exception)
        else respuestas.pop(0)))
    assert web.get('') == {'ok': 1}


def test_solo_una_parte_del_sistema(herramienta):
    """022:RF-56 — ``--alimentadores`` calibra solo esos, para aislar el que no deja converger;
    sin medidores elegidos no hay qué calibrar."""
    web = WebFalsa(proyecto='P', informes=[VALIDO])
    assert herramienta.correr(web, 'Pisco', demanda_dir=Path('l'), periodo=None, reconvertir=False,
                              transformadores=False, tension_pu=None, bancos=True,
                              solo=['pi102'], libros={**LIBROS, 'suministros': Path('l/s.xlsx')}) == 0
    web = WebFalsa(proyecto='P', informes=[VALIDO])
    assert herramienta.correr(web, 'Pisco', demanda_dir=Path('l'), periodo=None, reconvertir=False,
                              transformadores=False, tension_pu=None, bancos=True,
                              solo=['PI101'], libros={**LIBROS, 'suministros': Path('l/s.xlsx')}) == 2
    assert ('POST', '/demanda/ajuste/plan') not in web.llamadas


def test_medida_para_reconstruir(herramienta):
    """022:RF-59 — con medidor, la suya; sin medidor, la de su barra menos las de los demás
    (TM101 en SET 22: 6 792 − 6 858 = −66 kW, que no pasa carga)."""
    maxima = {'resultados': [
        {'alimentador': 'TM105', 'kw': 3482.0}, {'alimentador': 'TM106', 'kw': 3133.0},
        {'alimentador': 'TM104', 'kw': 31.0}, {'alimentador': 'TM103', 'kw': 97.0},
        {'alimentador': 'TM102', 'kw': 115.0}, {'alimentador': 'X', 'error': 'sin dato'}]}
    medidas = herramienta.medidas_de(maxima)
    assert medidas['TM106'] == 3133.0 and 'X' not in medidas
    barras = [{'kw': 6792.0,
               'alimentadores': ['TM101', 'TM102', 'TM103', 'TM104', 'TM105', 'TM106']}]
    assert herramienta.medida_de('TM106', medidas, barras) == 3133.0
    assert herramienta.medida_de('TM101', medidas, barras) == pytest.approx(-66.0)
    assert herramienta.medida_de('ZZ', medidas, barras) == 0.0


def test_reconstruir_reconvierte_con_las_medidas(herramienta):
    """022:RF-61, 004:RF-60 — ``--reconstruir`` manda los pedidos y las medidas a la conversión por
    sistema, importa de nuevo y calibra el proyecto nuevo."""
    web = WebFalsa(proyecto='P', informes=[VALIDO])
    enviados, suministros = [], []
    original = web.post

    def post(ruta, **kw):
        if ruta == '/convert/sistema':
            enviados.append(kw['json'])
        if ruta == '/demanda/suministros/path':
            suministros.append(kw['json']['path'])
        return original(ruta, **kw)

    web.post = post
    assert herramienta.correr(web, 'Pisco', demanda_dir=Path('l'), periodo=None,
                              reconvertir=False, transformadores=False, tension_pu=None,
                              bancos=True, reconstruir=['PI101'],
                              libros={**LIBROS, 'suministros': Path('l/s.xlsx')}) == 0
    assert enviados and enviados[0]['reconstruir'] == {'PI101': pytest.approx(4456.3)}
    # 004:RF-60 — el libro explícito sigue en la vuelta tras reconvertir.
    assert suministros == [str(Path('l/s.xlsx'))] * 2
    assert ('POST', '/powerfactory/flow') in web.llamadas
    assert [r for _m, r in web.llamadas].count('/demanda/ajuste/plan') == 1


def test_escanea_antes_de_calibrar_y_la_central_toma_su_barra(herramienta, monkeypatch, capsys):
    """028:RF-15 — el sistema se escanea antes del plan; la central del perfil sin barra va al
    plan con la del generador que encontró el escaneo."""
    from igea_dgs import perfil_empresa

    monkeypatch.setattr(perfil_empresa, 'datos_perfil', lambda *a, **k: {})
    monkeypatch.setattr(perfil_empresa, 'centrales_termicas', lambda _d: {
        'PI102': {'generacion': 'CT_X', 'llegada': 'CT_L', 'salida': 'CT_S'}})

    class ConEscaneo(WebFalsa):
        planes: list = []

        def post(self, ruta, **kw):
            if ruta == '/demanda/ajuste/plan':
                self.llamadas.append(('POST', ruta))
                self.planes.append(kw['json'])
                return {'token': 't', 'avisos': []}
            return super().post(ruta, **kw)

        def esperar(self, trabajo):
            if trabajo['id'] == '/powerfactory/escanear':
                self.esperados.append(trabajo['id'])
                return {'status': 'done', 'result': {
                    'resumen': {'PI102': {'generadores': 1, 'bancos': 0, 'cargas_cambiadas': 0}},
                    'alimentadores': {'PI102': {'generadores': [
                        {'barra': 'N5', 'en_servicio': True, 'origen': 'operador'}]}}}}
            return super().esperar(trabajo)

    web = ConEscaneo(informes=[VALIDO])
    assert _correr(herramienta, web) == 0
    rutas = [r for _m, r in web.llamadas]
    assert rutas.index('/powerfactory/escanear') < rutas.index('/demanda/ajuste/plan')
    [plan] = web.planes
    assert plan['centrales']['PI102']['nodo'] == 'N5'
    assert 'escaneo PI102: 1 generador(es)' in capsys.readouterr().out


def test_no_espera_al_trabajo_que_lo_lanzo(herramienta, monkeypatch):
    """028:RF-21 — lanzado por el botón «Recalibrar», no se espera a sí mismo: el trabajo que lo
    lanzó sigue en curso mientras recalibra (Chincha no empezaba nunca, 2026-10-28)."""
    monkeypatch.setenv('IGEA_TRABAJO_PROPIO', 'v')
    web = WebFalsa(informes=[VALIDO])
    assert _correr(herramienta, web) == 0
    assert 'v' not in web.esperados


class WebRed(WebFalsa):
    def __init__(self, red, grupo, *, importado=False, fallo=None):
        super().__init__(proyecto='P_RED' if importado else None, informes=[VALIDO])
        self.red, self.grupo, self.fallo = red, grupo, fallo
        self.enviados = {}

    def get(self, ruta, **kw):
        if ruta == '':
            return {'loaded': True, 'jobs': [],
                    'pf_projects': {self.grupo: self.proyecto} if self.proyecto else {}}
        if ruta == '/demanda/sistema':
            assert kw['params']['todos'] is True
            return {'sistema': {'grupo': self.grupo, 'alimentadores': self.red,
                                'medidores': {a: [f'L-{a}'] for a in self.red}}}
        if ruta == '/demanda/centrales':
            return {'centrales': []}
        return super().get(ruta, **kw)

    def post(self, ruta, **kw):
        self.llamadas.append(('POST', ruta))
        self.enviados[ruta] = kw.get('json')
        if ruta == '/powerfactory/flow' and self.fallo != ruta:
            self.proyecto = 'P_RED'
        if ruta == '/demanda/ajuste/plan':
            return {'token': 't'}
        return {'id': ruta, 'title': ruta}

    def esperar(self, trabajo):
        ruta = trabajo['id']
        if self.fallo == ruta:
            return {'id': ruta, 'status': 'failed', 'error': 'fallo deliberado'}
        if ruta == '/convert':
            return {'status': 'done', 'result': (
                {'status': 'ok'} if len(self.red) > 1 else
                {'status': None, 'summary': {'status': 'completed', 'ok': 1, 'failed': 0, 'skipped': 0}})}
        return super().esperar(trabajo)


def _correr_red(herramienta, web, **kw):
    return herramienta.correr(web, None, demanda_dir=Path('libros'), periodo='2026-07',
                              reconvertir=False, transformadores=False, tension_pu=None,
                              bancos=True, libros=LIBROS, red=web.red, **kw)


@pytest.mark.parametrize('red,nombre,grupo', [
    (['PI102'], None, 'PI102'),
    (['PI102', 'PI101'], None, 'PI102_PI101'),
    (['PI102', 'PI101'], ' Red Pisco / prueba ', 'Red_Pisco___prueba'),
])
def test_red_convierte_importa_y_usa_maximas_individuales(herramienta, red, nombre, grupo):
    """028:RF-25 / 028:RF-24 — ``--red`` convierte, importa y calibra uno o varios
    alimentadores, unidos en un proyecto si son varios, con la máxima de cada uno."""
    web = WebRed(red, grupo)
    assert _correr_red(herramienta, web, grupo=nombre) == 0
    assert web.enviados['/convert'] == {
        'feeders': red, 'unir': len(red) > 1, 'nombre': grupo if len(red) > 1 else ''}
    assert web.enviados['/powerfactory/flow'] == {'feeders': [grupo]}
    assert web.enviados['/demanda/cabecera/buscar']['por_barra'] is False
    assert ('POST', '/convert/sistema') not in web.llamadas


def test_red_recalibra_sin_reconvertir(herramienta):
    """028:RF-25 — con el proyecto ya importado, recalibra sin reconvertir."""
    web = WebRed(['PI102', 'PI101'], 'RED', importado=True)
    assert _correr_red(herramienta, web, grupo='RED') == 0
    assert '/convert' not in web.enviados and '/powerfactory/flow' not in web.enviados


@pytest.mark.parametrize('ruta', ['/convert', '/powerfactory/flow',
                                  '/demanda/cabecera/indexar', '/demanda/clientes/indexar',
                                  '/powerfactory/escanear', '/demanda/cabecera/buscar',
                                  '/powerfactory/projects'])
def test_red_no_calibra_si_falla_un_paso_previo(herramienta, ruta):
    web = WebRed(['PI102'], 'PI102', fallo=ruta)
    assert _correr_red(herramienta, web) == 2
    assert '/demanda/ajuste/plan' not in web.enviados


def test_red_reconstruir_se_rechaza_antes_de_convertir(herramienta):
    web = WebRed(['PI102'], 'PI102')
    assert _correr_red(herramienta, web, reconstruir=['PI102']) == 2
    assert not web.enviados


def test_sin_validacion_nativa_o_trabajo_fallido_no_es_exito(herramienta):
    web = WebRed(['PI102'], 'PI102', fallo='/demanda/ajuste/t/aplicar')
    assert _correr_red(herramienta, web) == 2
    assert herramienta._resumen({'result': {'informes': [
        {**VALIDO, 'validacion_posterior_nativa': {}}]}}) == (0, 1)


@pytest.mark.parametrize('status,summary', [
    ('completed', {'ok': 0, 'failed': 1, 'skipped': 0}),
    ('completed', {'ok': 0, 'failed': 0, 'skipped': 1}),
    ('cancelled', {'ok': 1, 'failed': 0, 'skipped': 0}),
])
def test_red_no_importa_conversion_individual_sin_resultado_util(herramienta, status, summary):
    web = WebRed(['PI102'], 'PI102')
    esperar = web.esperar

    def resultado(trabajo):
        if trabajo['id'] == '/convert':
            return {'status': 'done', 'result': {'status': None, 'summary': {**summary, 'status': status}}}
        return esperar(trabajo)

    web.esperar = resultado
    assert _correr_red(herramienta, web) == 2
    assert '/powerfactory/flow' not in web.enviados
