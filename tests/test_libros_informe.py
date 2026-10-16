"""Spec 017: libro de pérdidas de la empresa y tabla de costes, con libros sintéticos."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from igea_dgs.libros_informe import (
    COSTES_REFERENCIALES, ColumnasPerdidas, LibroError, leer_costes, leer_libro_perdidas)

PERFIL = {
    'informes': {
        'libro_perdidas': {
            'hoja': 'Pérdidas MT',
            'columnas': {
                'alimentador': 'Alimentador',
                'sistema_electrico': 'Sistema eléctrico',
                'reconocido_kw_pct': '% reconocido (kW)',
                'reconocido_kwh_pct': '% reconocido (kWh)',
                'codigo_sistema': 'Código',
            },
            'precio_compra': {'hoja': 'Precios', 'celda': 'C3'},
        },
    },
}


def _libro(tmp_path: Path, filas=None, porcentaje_como_fraccion=True, titulo=True) -> Path:
    """Un libro como el de una empresa: título arriba, cabecera más abajo, % con formato."""
    wb = Workbook()
    hoja = wb.active
    hoja.title = 'Pérdidas MT'
    if titulo:
        hoja.append(['PÉRDIDAS TÉCNICAS RECONOCIDAS'])
        hoja.append([])
    hoja.append(['Ítem', ' alimentador ', 'Sistema  eléctrico', '% Reconocido (kW)', '% reconocido (kWh)',
                 'Código'])
    for fila in filas or [(1, 'IC106', 'Ica', 0.0512, 0.0431, 'SE0044'),
                          (2, 'CA101', 'Chincha', 0.0611, 0.0502, 'SE0050')]:
        hoja.append(list(fila))
        if porcentaje_como_fraccion:
            for celda in hoja[hoja.max_row][3:5]:
                celda.number_format = '0.00%'
    precios = wb.create_sheet('Precios')
    precios['B3'] = 'Precio de compra (S/./kWh)'
    precios['C3'] = 0.2715
    ruta = tmp_path / 'perdidas.xlsx'
    wb.save(ruta)
    return ruta


# --------------------------------------------------------------------------- RF-62, RF-63
def test_lee_sistema_reconocidas_y_precio_del_alimentador(tmp_path):
    """017:RF-62, 017:RF-63 — sistema, % reconocidos y precio, por los nombres del perfil."""
    columnas = ColumnasPerdidas.de_perfil(PERFIL)
    datos = leer_libro_perdidas(_libro(tmp_path), 'CA101', columnas)
    assert datos.sistema_electrico == 'Chincha'
    assert datos.reconocido_kw_pct == pytest.approx(6.11)
    assert datos.reconocido_kwh_pct == pytest.approx(5.02)
    assert datos.precio_compra == pytest.approx(0.2715)
    assert datos.codigo_sistema == 'SE0050'
    assert datos.faltan == ()


def test_alimentador_con_otra_escritura(tmp_path):
    """017:RF-62 — «ca101 » en el formulario es el CA101 del libro."""
    datos = leer_libro_perdidas(_libro(tmp_path), ' ca101 ', ColumnasPerdidas.de_perfil(PERFIL))
    assert datos.sistema_electrico == 'Chincha'


def test_porcentaje_escrito_como_numero(tmp_path):
    """017:RF-62 — un 6,11 sin formato de % es 6,11 %, no 611 %."""
    ruta = _libro(tmp_path, filas=[(1, 'CA101', 'Chincha', 6.11, 5.02)], porcentaje_como_fraccion=False)
    datos = leer_libro_perdidas(ruta, 'CA101', ColumnasPerdidas.de_perfil(PERFIL))
    assert datos.reconocido_kw_pct == pytest.approx(6.11)


def test_precio_en_una_columna_de_la_misma_hoja(tmp_path):
    """017:RF-63 — el precio también puede ser una columna de la fila del alimentador."""
    wb = Workbook()
    hoja = wb.active
    hoja.title = 'Pérdidas MT'
    hoja.append(['Alimentador', 'Sistema eléctrico', '% reconocido (kW)', '% reconocido (kWh)', 'Precio'])
    hoja.append(['CA101', 'Chincha', 6.11, 5.02, 0.30])
    ruta = tmp_path / 'p.xlsx'
    wb.save(ruta)
    perfil = json.loads(json.dumps(PERFIL))
    perfil['informes']['libro_perdidas']['precio_compra'] = {'columna': 'Precio'}
    datos = leer_libro_perdidas(ruta, 'CA101', ColumnasPerdidas.de_perfil(perfil))
    assert datos.precio_compra == pytest.approx(0.30)


def test_hoja_o_columna_que_no_estan_se_nombran(tmp_path):
    """017:RF-63 — con otras hojas o columnas, el error dice cuáles faltan y qué hay."""
    perfil = json.loads(json.dumps(PERFIL))
    perfil['informes']['libro_perdidas']['columnas']['sistema_electrico'] = 'Sistema'
    with pytest.raises(LibroError, match=r"'Sistema'.*Sistema  eléctrico"):
        leer_libro_perdidas(_libro(tmp_path), 'CA101', ColumnasPerdidas.de_perfil(perfil))
    perfil['informes']['libro_perdidas']['hoja'] = 'Otra'
    with pytest.raises(LibroError, match=r"'Otra'.*Pérdidas MT"):
        leer_libro_perdidas(_libro(tmp_path), 'CA101', ColumnasPerdidas.de_perfil(perfil))


def test_sin_libro_en_el_perfil_no_se_supone_nada():
    """017:RF-63 — sin la clave del perfil no hay nombres por defecto: se dice qué falta."""
    with pytest.raises(LibroError, match='informes.libro_perdidas'):
        ColumnasPerdidas.de_perfil({})
    perfil = json.loads(json.dumps(PERFIL))
    del perfil['informes']['libro_perdidas']['columnas']['reconocido_kwh_pct']
    with pytest.raises(LibroError, match='reconocido_kwh_pct'):
        ColumnasPerdidas.de_perfil(perfil)


# --------------------------------------------------------------------------- RF-64
def test_alimentador_que_no_esta_deja_el_hueco(tmp_path):
    """017:RF-64 — sin el alimentador no hay error: lo reconocido queda señalado con su motivo."""
    datos = leer_libro_perdidas(_libro(tmp_path), 'TA225', ColumnasPerdidas.de_perfil(PERFIL))
    assert datos.sistema_electrico is None
    assert datos.reconocido_kwh_pct is None
    assert datos.precio_compra == pytest.approx(0.2715)
    assert any('TA225' in m for m in datos.faltan)


def test_celda_vacia_es_un_hueco(tmp_path):
    """017:RF-64 — un % reconocido en blanco queda en None y se dice cuál."""
    ruta = _libro(tmp_path, filas=[(1, 'CA101', 'Chincha', None, 0.05)])
    datos = leer_libro_perdidas(ruta, 'CA101', ColumnasPerdidas.de_perfil(PERFIL))
    assert datos.reconocido_kw_pct is None
    assert datos.reconocido_kwh_pct == pytest.approx(5.0)
    assert any('% reconocido (kW)' in m for m in datos.faltan)


# --------------------------------------------------------------------------- RF-39b a RF-40c
def _costes(tmp_path, **cambios):
    datos = {
        'precio_energia': {'valor': 0.27, 'unidad': 'S/./kWh', 'moneda': 'PEN',
                           'fuente': 'Libro de la empresa', 'fecha': '2026-09', 'referencial': False},
        'bancos': [{'kvar': 300, 'tension_kv': 10, 'inversion': 9000, 'mantenimiento_anual': 200,
                    'moneda': 'USD', 'fuente': 'Cotización 123', 'fecha': '2026-08', 'referencial': False}],
        'reguladores': [{'tipo': 'Monofásico 32 pasos 167 kVA', 'clase_kv': 15, 'inversion': 60000,
                         'mantenimiento_anual': 900, 'moneda': 'USD', 'fuente': 'Cotización 124',
                         'fecha': '2026-08', 'referencial': False}],
    }
    datos.update(cambios)
    ruta = tmp_path / 'costes.json'
    ruta.write_text(json.dumps(datos), encoding='utf-8')
    return ruta


def test_bancos_reguladores_y_precio_con_fuente_y_fecha(tmp_path):
    """017:RF-39b, 017:RF-40c — cada fila con kvar o tipo, inversión, mantenimiento, fuente y fecha."""
    costes = leer_costes(_costes(tmp_path))
    banco, = costes.bancos
    assert (banco.kvar, banco.tension_kv, banco.inversion, banco.mantenimiento_anual) == (300, 10, 9000, 200)
    assert banco.fuente == 'Cotización 123' and banco.fecha == '2026-08'
    regulador, = costes.reguladores
    assert regulador.tipo == 'Monofásico 32 pasos 167 kVA' and regulador.inversion == 60000
    assert costes.precio_energia.valor == 0.27
    assert costes.faltan_para_bancos() == []
    assert costes.faltan_para_reguladores() == []


def test_fila_sin_fuente_o_fecha_no_vale(tmp_path):
    """017:RF-39b — un precio sin fuente o sin fecha no se usa: error con la fila."""
    sin_fuente = {'kvar': 300, 'inversion': 9000, 'moneda': 'USD', 'fecha': '2026-08', 'referencial': False}
    with pytest.raises(LibroError, match=r'bancos\[1\].*fuente'):
        leer_costes(_costes(tmp_path, bancos=[sin_fuente]))


def test_sin_bancos_o_sin_precio_no_hay_etapa_2(tmp_path):
    """017:RF-39c — sin bancos o sin precio de la energía, se dice qué falta para la etapa 2."""
    costes = leer_costes(_costes(tmp_path, bancos=[], precio_energia=None))
    assert costes.faltan_para_bancos() == ['bancos de condensadores', 'precio de la energía']
    assert costes.faltan_para_reguladores() == []
    con_precio = costes.con_precio_energia(0.2715, 'S/./kWh', 'PEN', 'Libro de pérdidas', '2026-09')
    assert con_precio.faltan_para_bancos() == ['bancos de condensadores']


def test_origen_de_los_costes(tmp_path):
    """017:RF-39e — el informe dice si los costes son de la empresa, referenciales o mixtos."""
    assert leer_costes(_costes(tmp_path)).origen == 'de la empresa'
    referencial = leer_costes(COSTES_REFERENCIALES)
    assert referencial.origen == 'referenciales'
    costes = leer_costes(_costes(tmp_path)).con_precio_energia(0.3, 'S/./kWh', 'PEN', 'x', '2026')
    mixto = leer_costes(_costes(tmp_path, bancos=[json.loads(json.dumps(
        {'kvar': 300, 'inversion': 1, 'moneda': 'USD', 'fuente': 'web', 'fecha': '2022', 'referencial': True}))]))
    assert costes.origen == 'de la empresa' and mixto.origen == 'mixtos'


def test_tabla_referencial_versionada(tmp_path):
    """017:RF-39d — la tabla de ejemplo trae precios de la web, todos referenciales con URL y fecha."""
    costes = leer_costes(COSTES_REFERENCIALES)
    filas = list(costes.bancos) + list(costes.reguladores)
    assert costes.bancos and costes.reguladores
    assert all(f.referencial and f.fuente.startswith('https://') and f.fecha for f in filas)
