"""Cargas de las SED: actualización masiva y creación, por libro Excel/CSV.

Formato nativo — el de ``referencia/PA217.xlsx``, que es el que usa la empresa:

- **Una hoja por alimentador**, y el **nombre de la hoja es el alimentador**
  (``PA217``, ``IN111``, ``PA219``…). Un solo libro cubre varios alimentadores.
- Columnas: ``SED`` · ``Kw`` · ``Kvar`` · ``(kVA)`` · ``FP``.

Las cargas llegan medidas de tres formas, y cada fila trae **un par**, el que tenga:

    Kw y Kvar   → se usan tal cual
    Kw y FP     → Kvar = Kw · √(1 − FP²) / FP
    (kVA) y FP  → Kw = kVA · FP,  Kvar = kVA · √(1 − FP²)

El fichero real trae el tercer caso: ``Kw``/``Kvar`` vacías. Comprobado contra él: kVA
0,3633 con FP 0,936183 da Kw 0,3401 y Kvar 0,1277, cuyo módulo devuelve exactamente
0,3633.

Si la fila trae más de un par, tienen que coincidir: si no, es un error de fila. Antes
mandaba ``Kw``/``Kvar`` en silencio, y la plantilla venía con ``Kw``/``Kvar`` ya
rellenas con la carga actual: quien escribía solo ``(kVA)`` y ``FP`` veía que su dato
se ignoraba y la SED seguía igual. Por eso la plantilla deja vacías las columnas de
entrada y muestra la carga actual aparte, en ``Kw_actual``/``Kvar_actual``/``FP_actual``.

Una fila sin ningún valor no cambia la SED (la plantilla trae todas, y se rellenan solo
las que llegan). Una fila con ``FP = 0`` y ``kVA = 0`` escritos, que en el fichero real
son 58, deja la carga en cero y se registra como fila sin datos, no como error.

Tres situaciones, que no son el mismo problema:

- ``updates``  — SED de la hoja que existen en el modelo: se actualizan.
- ``unknown``  — SED de la hoja que **no** están en el modelo: no se inventan. Salen
  listadas para el módulo de creación. (En el fichero real aparece ``5SE31022``, que es
  ``SE31022`` con un 5 de más: el que detecta esto es justamente este módulo.)
- ``untouched``— SED del modelo que la hoja no menciona: se dejan como están. (En el
  fichero real son las 84 SED con prefijo ``M``.)

Este módulo **no** habla con PowerFactory: produce y valida el plan. Aplicarlo es tarea
de ``tools/apply_sed_loads.py``, en un proceso aparte con el Python que exige la API.
"""

from __future__ import annotations

import csv
import math
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any, Sequence

# Columnas del formato nativo, en su orden y con su grafía exacta.
SHEET_COLUMNS = ('SED', 'Kw', 'Kvar', '(kVA)', 'FP')
# Columnas extra que añade la plantilla generada: informativas, no obligatorias.
# La carga actual va en columnas propias, no en Kw/Kvar: si fuera en las de entrada,
# mandaría sobre el (kVA)/FP que escribe el operador.
EXTRA_COLUMNS = (
    'kVA_instalado', 'Kw_actual', 'Kvar_actual', 'FP_actual', 'accion', 'observaciones',
)
ENTRY_COLUMNS = ('Kw', 'Kvar', '(kVA)', 'FP')

# Margen para dar por iguales dos pares de la misma fila: el redondeo de una hoja a
# dos decimales (FP 0,92 por 0,9196) no es una contradicción.
CONFLICT_REL_TOL = 0.01
CONFLICT_FP_TOL = 0.01

# Nombres tolerados para cada columna. La cabecera se normaliza (minúsculas, sin
# espacios ni paréntesis) antes de buscar aquí, así que un fichero con «P (kW)» o
# «cos phi» se sigue entendiendo.
COLUMN_ALIASES = {
    'sed': 'SED', 'sedcode': 'SED', 'sed_code': 'SED', 'codigo': 'SED', 'code': 'SED',
    'kw': 'Kw', 'pkw': 'Kw', 'p': 'Kw', 'p_kw': 'Kw', 'potencia': 'Kw',
    'kvar': 'Kvar', 'qkvar': 'Kvar', 'q': 'Kvar', 'q_kvar': 'Kvar', 'reactiva': 'Kvar',
    'kva': '(kVA)', 'skva': '(kVA)', 's': '(kVA)', 's_kva': '(kVA)', 'aparente': '(kVA)',
    'fp': 'FP', 'pf': 'FP', 'cos': 'FP', 'cosphi': 'FP', 'factordepotencia': 'FP',
    'cosfi': 'FP', 'cosφ': 'FP', 'cosϕ': 'FP', 'fdp': 'FP',
    'kvainstalado': 'kVA_instalado', 'kvainst': 'kVA_instalado',
    'accion': 'accion', 'observaciones': 'observaciones', 'obs': 'observaciones',
}

ACTION_UPDATE = 'actualizar'
ACTION_SKIP = 'omitir'
VALID_ACTIONS = (ACTION_UPDATE, ACTION_SKIP, '')

SUPPORTED_SUFFIXES = ('.xlsx', '.csv')


class LoadTemplateError(ValueError):
    """El libro no se puede leer o su contenido no es utilizable."""


def _normalise(header: Any) -> str:
    text = ('' if header is None else str(header)).strip().lower()
    for ch in ' ()[]-_./\\':
        text = text.replace(ch, '')
    return text


def canonical_columns(header: Sequence[Any]) -> dict[int, str]:
    """Índice de columna → nombre canónico, tolerando variantes de grafía."""
    mapping: dict[int, str] = {}
    for index, raw in enumerate(header):
        key = _normalise(raw)
        if key in COLUMN_ALIASES:
            mapping[index] = COLUMN_ALIASES[key]
    return mapping


@dataclass(frozen=True)
class SedLoad:
    """Carga de una SED: la del modelo, o la que pide la hoja."""

    sed_code: str
    feeder: str = ''
    kw: float = 0.0
    kvar: float = 0.0
    kva: float = 0.0
    fp: float = 0.0
    installed_kva: float = 0.0
    # Identidad dentro del modelo; la hoja no la trae y se completa al cruzar.
    section_id: str = ''
    device_number: str = ''
    node_id: str = ''
    derived: bool = False          # Kw/Kvar calculados desde (kVA) y FP
    no_data: bool = False          # ni Kw/Kvar ni kVA/FP utilizables

    @property
    def p_mw(self) -> float:
        return self.kw / 1000.0

    @property
    def q_mvar(self) -> float:
        return self.kvar / 1000.0


@dataclass
class SheetRead:
    """Resultado de leer una hoja: lo utilizable, lo omitido y lo que está mal."""

    feeder: str = ''
    rows: list[SedLoad] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    no_data: list[str] = field(default_factory=list)


@dataclass
class LoadUpdatePlan:
    """Qué se va a cambiar, qué no se reconoce y qué se queda igual."""

    feeder: str
    updates: list[tuple[SedLoad, SedLoad]] = field(default_factory=list)   # (actual, nueva)
    unknown: list[SedLoad] = field(default_factory=list)
    untouched: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    no_data: list[str] = field(default_factory=list)
    row_errors: list[str] = field(default_factory=list)

    @property
    def has_changes(self) -> bool:
        return bool(self.updates)

    @property
    def is_applicable(self) -> bool:
        """Una sola fila inválida bloquea todo: una actualización a medias es peor."""
        return not self.row_errors and self.has_changes

    def summary(self) -> dict:
        return {
            'feeder': self.feeder,
            'updates': len(self.updates),
            'unknown': len(self.unknown),
            'untouched': len(self.untouched),
            'skipped': len(self.skipped),
            'no_data': len(self.no_data),
            'row_errors': len(self.row_errors),
            'applicable': self.is_applicable,
        }

    def report(self) -> str:
        lines = [
            f'PLAN DE ACTUALIZACIÓN DE CARGAS — {self.feeder}',
            '=' * 70,
            f'  SED a actualizar                         : {len(self.updates)}',
            f'  SED sin datos en la hoja (kVA y FP en 0) : {len(self.no_data)}',
            f'  SED omitidas por el operador             : {len(self.skipped)}',
            f'  SED del modelo no mencionadas            : {len(self.untouched)}',
            f'  SED de la hoja que NO están en el modelo : {len(self.unknown)}',
            f'  Filas con error                          : {len(self.row_errors)}',
        ]
        if self.unknown:
            lines += ['', 'SED no encontradas en el modelo (van al módulo de creación):']
            for item in self.unknown[:20]:
                lines.append(f'  - {item.sed_code}  ({item.kw:g} kW / {item.kvar:g} kvar)')
            if len(self.unknown) > 20:
                lines.append(f'  … y {len(self.unknown) - 20} más')
        if self.row_errors:
            lines += ['', 'ERRORES (no se aplica ningún cambio hasta corregirlos):']
            for err in self.row_errors[:20]:
                lines.append(f'  - {err}')
            if len(self.row_errors) > 20:
                lines.append(f'  … y {len(self.row_errors) - 20} más')
        if self.updates:
            derivadas = sum(1 for _a, n in self.updates if n.derived)
            if derivadas:
                lines += ['', f'{derivadas} carga(s) calculadas desde Kw y FP, o desde (kVA) y FP.']
            lines += ['', 'Cambios mayores (primeros 10 por variación de P):']
            ordenados = sorted(self.updates, key=lambda p: abs(p[1].kw - p[0].kw), reverse=True)
            for actual, nueva in ordenados[:10]:
                lines.append(
                    f'  {actual.sed_code:14s} P {actual.kw:10.4f} → {nueva.kw:10.4f} kW   '
                    f'Q {actual.kvar:10.4f} → {nueva.kvar:10.4f} kvar'
                )
        return '\n'.join(lines)


# ------------------------------------------------------------------ desde el modelo

def model_sed_loads(model) -> list[SedLoad]:
    """Carga actual de cada SED del modelo, una fila por SED.

    La clave es el código de SED, que es también el ``loc_name`` del ``ElmSubstat`` y
    del ``ElmLod`` en el DGS: es lo que permite localizar la carga en PowerFactory sin
    depender de FID, que cambian en cada conversión.
    """
    by_key = {(load.section_id, load.device_number): load for load in model.loads}
    rows: list[SedLoad] = []
    for sed in model.seds:
        load = by_key.get(sed.load_key)
        kw = (load.p_mw * 1000.0) if load else 0.0
        kvar = (load.q_mvar * 1000.0) if load else 0.0
        rows.append(SedLoad(
            sed_code=sed.code,
            feeder=model.name,
            kw=kw, kvar=kvar,
            kva=math.hypot(kw, kvar),
            fp=load.pf if load else 0.0,
            installed_kva=sed.design_kva,
            section_id=sed.section_id,
            device_number=sed.device_number,
            node_id=sed.node_id,
        ))
    rows.sort(key=lambda r: r.sed_code)
    return rows


# ------------------------------------------------------------------ plantilla

def write_template(
    feeders: dict[str, Sequence[SedLoad]] | Sequence[SedLoad],
    path: Path | str,
    *,
    include_extras: bool = True,
) -> Path:
    """Escribe la plantilla: ``.xlsx`` con una hoja por alimentador, o ``.csv`` plano.

    ``feeders`` puede ser ``{alimentador: filas}`` o una lista suelta (un alimentador).
    """
    out = Path(path)
    suffix = out.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise LoadTemplateError(
            f'Extensión no soportada: {out.suffix!r}. Use ' + ' o '.join(SUPPORTED_SUFFIXES)
        )
    if not isinstance(feeders, dict):
        rows = list(feeders)
        name = rows[0].feeder if rows else 'CARGAS'
        feeders = {name: rows}
    out.parent.mkdir(parents=True, exist_ok=True)

    columns = list(SHEET_COLUMNS) + (list(EXTRA_COLUMNS) if include_extras else [])

    def record(row: SedLoad) -> list[Any]:
        # Las de entrada van vacías: el operador rellena el par que tenga.
        base: list[Any] = [row.sed_code, None, None, None, None]
        if include_extras:
            base += [
                round(row.installed_kva, 3),
                round(row.kw, 6), round(row.kvar, 6), round(row.fp, 6),
                ACTION_UPDATE, '',
            ]
        return base

    if suffix == '.csv':
        # CSV no tiene hojas: se añade la columna del alimentador para no perderlo.
        with out.open('w', encoding='utf-8-sig', newline='') as fh:
            writer = csv.writer(fh, delimiter=';')
            writer.writerow(['feeder'] + columns)
            for feeder, rows in feeders.items():
                for row in rows:
                    writer.writerow([feeder] + record(row))
        return out

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise LoadTemplateError(
            'La plantilla .xlsx necesita openpyxl. Instale con: '
            'pip install -r requirements.txt  — o genere la plantilla en .csv.'
        ) from exc

    wb = Workbook()
    wb.remove(wb.active)
    negrita = Font(bold=True)
    verde = PatternFill('solid', fgColor='DFF0D8')      # rellenar
    gris = PatternFill('solid', fgColor='DDDDDD')       # no editar
    for feeder, rows in feeders.items():
        # Excel limita el nombre de hoja a 31 caracteres y prohíbe : \ / ? * [ ]
        title = str(feeder)[:31]
        for ch in ':\\/?*[]':
            title = title.replace(ch, '-')
        ws = wb.create_sheet(title=title or 'CARGAS')
        ws.append(columns)
        for col, name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font = negrita
            cell.fill = verde if name in ENTRY_COLUMNS else gris
            ws.column_dimensions[cell.column_letter].width = max(len(str(name)) + 3, 13)
        for row in rows:
            ws.append(record(row))
        ws.freeze_panes = 'A2'
    if not wb.sheetnames:
        wb.create_sheet(title='CARGAS').append(columns)
    wb.save(out)
    return out


# ------------------------------------------------------------------ lectura

def _number(value: Any) -> float | None:
    text = ('' if value is None else str(value)).strip().replace(',', '.')
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return float('nan')
    return number if math.isfinite(number) else float('nan')


def _resolve_power(
    kw: float | None, kvar: float | None, kva: float | None, fp: float | None,
) -> tuple[float, float, float, float, bool, bool]:
    """Devuelve ``(kw, kvar, kva, fp, derivado, sin_datos)``.

    Pares, por orden: Kw y Kvar tal cual; Kw y FP; (kVA) y FP, que es el caso del
    fichero real. Si la fila trae más de un par, ``_power_conflict`` comprueba que
    coincidan. Sin ningún par, la carga queda en cero.
    """
    if kw is not None and kvar is None and fp is not None and 0 < fp <= 1:
        # Kw con cos φ: lo que da un medidor que no registra reactiva.
        q = kw * math.sqrt(max(1.0 - fp * fp, 0.0)) / fp
        return kw, q, math.hypot(kw, q), fp, True, kw == 0.0
    if kw is not None or kvar is not None:
        p = kw or 0.0
        q = kvar or 0.0
        s = math.hypot(p, q)
        factor = (p / s) if s > 0 else (fp or 0.0)
        return p, q, s, factor, False, (p == 0.0 and q == 0.0)
    s = kva or 0.0
    factor = fp or 0.0
    if s > 0 and 0 < factor <= 1:
        p = s * factor
        q = s * math.sqrt(max(1.0 - factor * factor, 0.0))
        return p, q, s, factor, True, False
    return 0.0, 0.0, s, factor, False, True


def _power_conflict(
    kva: float | None, fp: float | None, p: float, q: float,
) -> str | None:
    """Si la fila trae (kVA)/FP que no cuadran con la carga resuelta, dice en qué.

    Kw y Kvar ya están en ``p``/``q`` cuando vienen, así que basta mirar las otras dos.
    Un 0 en (kVA) o FP es «sin dato» en el fichero real, no un valor que comparar.
    """
    s = math.hypot(p, q)
    if s == 0.0:
        return None
    if kva and abs(s - kva) > CONFLICT_REL_TOL * max(s, kva):
        return f'(kVA)={kva:g} no coincide con Kw/Kvar, que dan {s:.4g} kVA'
    if fp and abs(p / s - fp) > CONFLICT_FP_TOL:
        return f'FP={fp:g} no coincide con Kw/Kvar, que dan FP {p / s:.4f}'
    return None


def read_sheet(header: Sequence[Any], data: Sequence[Sequence[Any]], feeder: str) -> SheetRead:
    """Lee una hoja ya materializada. Separada para poder probarla sin ficheros."""
    result = SheetRead(feeder=feeder)
    mapping = canonical_columns(header)
    if 'SED' not in mapping.values():
        raise LoadTemplateError(
            f'hoja {feeder!r}: no se encontró la columna de código de SED. '
            f'Cabecera leída: {[str(h) for h in header]}. '
            'Se esperaba una columna «SED».'
        )
    seen: dict[str, int] = {}
    for offset, values in enumerate(data):
        row_no = offset + 2                      # fila 1 = cabecera
        cell = {name: values[i] if i < len(values) else None for i, name in mapping.items()}
        code = ('' if cell.get('SED') is None else str(cell['SED'])).strip()
        if not code:
            continue                             # fila en blanco: no es un error
        if code in seen:
            result.errors.append(
                f'hoja {feeder}, fila {row_no}: SED {code!r} repetida '
                f'(ya en la fila {seen[code]})'
            )
            continue
        seen[code] = row_no
        action = ('' if cell.get('accion') is None else str(cell['accion'])).strip().lower()
        if action not in VALID_ACTIONS:
            result.errors.append(
                f'hoja {feeder}, fila {row_no}: accion={action!r} no válida '
                f'(use {ACTION_UPDATE!r}, {ACTION_SKIP!r} o vacío)'
            )
            continue
        if action == ACTION_SKIP:
            result.skipped.append(code)
            continue

        numbers = {}
        bad = False
        for name in ('Kw', 'Kvar', '(kVA)', 'FP', 'kVA_instalado'):
            value = _number(cell.get(name))
            if value is not None and math.isnan(value):
                result.errors.append(
                    f'hoja {feeder}, fila {row_no}: {name}={cell.get(name)!r} no es un número'
                )
                bad = True
                break
            numbers[name] = value
        if bad:
            continue
        fp = numbers['FP']
        if fp is not None and not (0.0 <= fp <= 1.0):
            result.errors.append(
                f'hoja {feeder}, fila {row_no}: FP={fp} fuera de [0, 1]'
            )
            continue
        for name in ('Kw', 'Kvar', '(kVA)', 'kVA_instalado'):
            if numbers[name] is not None and numbers[name] < 0:
                result.errors.append(
                    f'hoja {feeder}, fila {row_no}: {name}={numbers[name]} negativo'
                )
                bad = True
                break
        if bad:
            continue

        if all(numbers[name] is None for name in ENTRY_COLUMNS):
            # La plantilla trae todas las SED: la que se deja en blanco no se toca.
            # Ponerla en cero sería borrar la carga de todo lo que no se midió.
            continue
        kw, kvar, kva, factor, derived, no_data = _resolve_power(
            numbers['Kw'], numbers['Kvar'], numbers['(kVA)'], fp,
        )
        conflict = _power_conflict(numbers['(kVA)'], fp, kw, kvar)
        if conflict:
            result.errors.append(
                f'hoja {feeder}, fila {row_no}: {code}: {conflict}. '
                'Deje un solo par: Kw y Kvar, Kw y FP, o (kVA) y FP.'
            )
            continue
        if no_data:
            result.no_data.append(code)
        result.rows.append(SedLoad(
            sed_code=code, feeder=feeder,
            kw=kw, kvar=kvar, kva=kva, fp=factor,
            installed_kva=numbers['kVA_instalado'] or 0.0,
            derived=derived, no_data=no_data,
        ))
    return result


def read_workbook(path: Path | str) -> dict[str, SheetRead]:
    """Lee el libro completo: ``{alimentador: SheetRead}``.

    En ``.xlsx`` el alimentador es el **nombre de la hoja**; en ``.csv``, la columna
    ``feeder`` si existe, y si no, todas las filas van a una sola clave.
    """
    src = Path(path)
    if not src.is_file():
        raise LoadTemplateError(f'No se encuentra el fichero: {src}')
    suffix = src.suffix.lower()

    if suffix == '.csv':
        with src.open('r', encoding='utf-8-sig', newline='') as fh:
            sample = fh.read(4096)
            fh.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
            except csv.Error:
                dialect = csv.excel
                dialect.delimiter = ';'
            reader = csv.reader(fh, dialect=dialect)
            try:
                header = next(reader)
            except StopIteration:
                raise LoadTemplateError(f'{src.name}: fichero vacío.') from None
            body = [row for row in reader if any((c or '').strip() for c in row)]
        feeder_col = next(
            (i for i, h in enumerate(header) if _normalise(h) in ('feeder', 'alimentador')),
            None,
        )
        if feeder_col is None:
            return {src.stem: read_sheet(header, body, src.stem)}
        grupos: dict[str, list[list[Any]]] = {}
        for row in body:
            name = (row[feeder_col] if feeder_col < len(row) else '') or src.stem
            grupos.setdefault(str(name).strip(), []).append(row)
        return {name: read_sheet(header, rows, name) for name, rows in grupos.items()}

    if suffix != '.xlsx':
        raise LoadTemplateError(
            f'Extensión no soportada: {src.suffix!r}. Use ' + ' o '.join(SUPPORTED_SUFFIXES)
        )
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise LoadTemplateError(
            'Leer .xlsx necesita openpyxl. Instale con: pip install -r requirements.txt'
        ) from exc

    wb = load_workbook(src, data_only=True, read_only=True)
    out: dict[str, SheetRead] = {}
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            it = ws.iter_rows(values_only=True)
            try:
                header = list(next(it))
            except StopIteration:
                continue                          # hoja vacía: se ignora en silencio
            body = [
                list(values) for values in it
                if values and any(v is not None and str(v).strip() for v in values)
            ]
            if not body:
                continue                          # hoja de solo cabecera
            try:
                out[name] = read_sheet(header, body, name)
            except LoadTemplateError as exc:
                # Una hoja ajena no debe tumbar el libro entero.
                out[name] = SheetRead(feeder=name, errors=[str(exc)])
    finally:
        wb.close()
    if not out:
        raise LoadTemplateError(
            f'{src.name}: ninguna hoja con datos. Se esperaba una hoja por alimentador '
            'con columnas SED, Kw, Kvar, (kVA) y FP.'
        )
    return out


# ------------------------------------------------------------------ varios alimentadores

def split_by_feeder(
    libro: dict[str, SheetRead], sed_feeder: dict[str, str], feeders: Sequence[str],
) -> tuple[dict[str, SheetRead], list[str]]:
    """Reparte un libro con varios alimentadores: cada SED, al alimentador que la tiene.

    La hoja no basta: el operador junta en una hoja SED de varios alimentadores, o
    pega una en la hoja equivocada, y entonces la SED «no existe» en ese alimentador
    aunque sí exista en el de al lado. Por eso manda ``sed_feeder`` (código → alimentador,
    sacado de los modelos elegidos); la hoja solo decide cuando la SED no está en
    ninguno, y entonces sale como desconocida en el alimentador de su hoja.

    Devuelve ``({alimentador: SheetRead}, errores_sin_alimentador)``. Los errores de
    una hoja que no es ningún alimentador elegido no se pueden asignar y van aparte.
    """
    elegidos = list(feeders)
    out = {f: SheetRead(feeder=f) for f in elegidos}
    sueltos: list[str] = []

    def destino(code: str, hoja: str) -> str | None:
        return sed_feeder.get(code) or (hoja if hoja in out else None)

    for hoja, read in libro.items():
        if read.errors:
            if hoja in out:
                out[hoja].errors.extend(read.errors)
            else:
                sueltos.extend(read.errors)
        for row in read.rows:
            f = destino(row.sed_code, hoja)
            if f is None:
                sueltos.append(
                    f'hoja {hoja}: {row.sed_code} no está en ningún alimentador elegido '
                    f'({", ".join(elegidos)}) y la hoja no es uno de ellos.'
                )
                continue
            out[f].rows.append(replace(row, feeder=f))
            if row.no_data:
                out[f].no_data.append(row.sed_code)
        for code in read.skipped:
            f = destino(code, hoja)
            if f is not None:
                out[f].skipped.append(code)
    # La misma SED en dos hojas acaba en el mismo alimentador: sería aplicarla dos
    # veces con valores quizá distintos.
    for f, read in out.items():
        vistos: set[str] = set()
        for row in read.rows:
            if row.sed_code in vistos:
                read.errors.append(f'{f}: SED {row.sed_code} aparece en más de una hoja')
            vistos.add(row.sed_code)
    return out, sueltos


# ------------------------------------------------------------------ plan

def build_plan(model, read: SheetRead) -> LoadUpdatePlan:
    """Cruza una hoja con el modelo y devuelve el plan, sin aplicar nada."""
    current = {row.sed_code: row for row in model_sed_loads(model)}
    plan = LoadUpdatePlan(feeder=getattr(model, 'name', read.feeder))
    plan.row_errors = list(read.errors)
    plan.skipped = list(read.skipped)
    plan.no_data = list(read.no_data)

    mentioned: set[str] = set(plan.skipped)
    for row in read.rows:
        mentioned.add(row.sed_code)
        existing = current.get(row.sed_code)
        if existing is None:
            plan.unknown.append(row)
            continue
        # La identidad la pone el modelo: la hoja solo trae el código y la carga.
        plan.updates.append((existing, SedLoad(
            sed_code=existing.sed_code,
            feeder=existing.feeder,
            kw=row.kw, kvar=row.kvar, kva=row.kva,
            fp=row.fp if row.fp else existing.fp,
            installed_kva=row.installed_kva or existing.installed_kva,
            section_id=existing.section_id,
            device_number=existing.device_number,
            node_id=existing.node_id,
            derived=row.derived,
            no_data=row.no_data,
        )))
    plan.untouched = sorted(code for code in current if code not in mentioned)
    return plan


def plan_to_payload(plan: LoadUpdatePlan) -> dict:
    """Plan serializable, para el proceso que habla con PowerFactory."""
    return {
        'feeder': plan.feeder,
        'summary': plan.summary(),
        'updates': [
            {
                'sed_code': nueva.sed_code,
                'section_id': nueva.section_id,
                'device_number': nueva.device_number,
                'plini_mw': nueva.p_mw,
                'qlini_mvar': nueva.q_mvar,
                'coslini': nueva.fp,
                'slini_mva': math.hypot(nueva.p_mw, nueva.q_mvar),
                'derived_from_kva_fp': nueva.derived,
                'previous': {'plini_mw': actual.p_mw, 'qlini_mvar': actual.q_mvar},
            }
            for actual, nueva in plan.updates
        ],
        'unknown': [
            {'sed_code': r.sed_code, 'kw': r.kw, 'kvar': r.kvar, 'kva': r.kva, 'fp': r.fp}
            for r in plan.unknown
        ],
        'untouched': plan.untouched,
        'skipped': plan.skipped,
        'no_data': plan.no_data,
        'row_errors': plan.row_errors,
    }
