"""Spec 017: el informe de estudio desde la web, con lo que guardaron los estudios."""
from __future__ import annotations

import json
import time
from pathlib import Path

import docx
import pytest

pytest.importorskip('fastapi')

from fastapi.testclient import TestClient  # noqa: E402

from apoyo import esperar_trabajo  # noqa: E402
from igea_dgs.web.app import create_app  # noqa: E402


def _medida(vmin, kvar_perdidas=55.0):
    return {'converge': True, 'barras': [['CAB', 1.0], ['N1', vmin]], 'elementos': [['L1', 'ElmLne', 60.0]],
            'vm_min': vmin, 'vm_min_barra': 'N1', 'vm_max': 1.0, 'vm_max_barra': 'CAB', 'mas_cargado': 'L1',
            'cargabilidad_max': 60.0, 'cabecera_kw': 6000.0, 'cabecera_kvar': 3000.0, 'perdidas_kw': 300.0,
            'perdidas_kvar': kvar_perdidas,
            'vm_min_fases': {'A': [vmin, 'N1'], 'B': [vmin + 0.01, 'N1'], 'C': [vmin + 0.02, 'N1']}}


@pytest.fixture()
def client(tmp_path, monkeypatch):
    modelo = docx.Document()
    modelo.sections[0].header.add_table(1, 1, width=modelo.sections[0].page_width).cell(0, 0).text = \
        'INFORME TECNICO Nº :\nGT-019-PR'
    modelo.add_paragraph('Texto del modelo: AGROINDUSTRIAS DEL MODELO.')
    modelo.save(tmp_path / 'modelo.docx')
    perfil = tmp_path / 'perfil_empresa.json'
    perfil.write_text(json.dumps({'informes': {
        'modelo_estudio': str(tmp_path / 'modelo.docx'),
        'modelo_factibilidad': str(tmp_path / 'modelo.docx'), 'bandas_tension_pct': [95, 105],
        'bandas_cargabilidad_pct': [80, 100], 'cabecera': {'numero': 'INFORME TECNICO'}, 'norma': 'NTCSE'}}),
        encoding='utf-8')
    monkeypatch.setenv('IGEA_PERFIL_EMPRESA', str(perfil))
    monkeypatch.setenv('IGEA_WEB_FRONTEND', str(tmp_path / 'sin_front'))
    with TestClient(create_app(tmp_path / 'web')) as c:
        yield c


def _espacio(client, tmp_path, *, con=True) -> tuple[str, Path]:
    wid = client.post('/api/workspaces').json()['id']
    salida = tmp_path / 'web' / wid / 'salida'
    salida.mkdir(parents=True, exist_ok=True)
    (salida / 'estudio_sin_flujo_desequilibrado_a.json').write_text(
        json.dumps({'medidas': {'AL01': _medida(0.90)}}), encoding='utf-8')
    time.sleep(0.01)
    if con:
        (salida / 'estudio_con_flujo_desequilibrado_b.json').write_text(
            json.dumps({'medidas': {'AL01': _medida(0.96)}}), encoding='utf-8')
    return wid, salida


FORMULARIO = {'alimentador': 'AL01', 'fecha': '2026-10-15', 'lugar': 'Ica', 'numero': 'GT-300-PR',
              'asunto': 'CAÍDA DE TENSIÓN DEL CLIENTE X', 'destinatario': 'ING. Z',
              'antecedentes': ['El cliente X reclama baja tensión.'],
              'limites': {'tension_min_pu': 0.95, 'tension_max_pu': 1.05, 'cargabilidad_max_pct': 100}}


def test_informe_de_estudio_desde_los_estudios_guardados(client, tmp_path):
    """017:RF-1b, 017:RF-26, 017:RF-29, 017:RF-30 — el Word sale de los estudios de los dos casos y del formulario."""
    wid, salida = _espacio(client, tmp_path)
    obra = {'clase': 'condensador', 'alimentador': 'AL01', 'casos': ['Con proyecto'], 'barra': 'N7',
            'q_kvar': 300}
    r = client.post(f'/api/workspaces/{wid}/obras', json=obra)
    if r.status_code != 201:  # la validación de obras pide la red de DIgSILENT: se guarda a mano
        datos = json.loads((tmp_path / 'web' / wid / 'workspace.json').read_text(encoding='utf-8'))
        datos['obras'] = [obra]
        (tmp_path / 'web' / wid / 'workspace.json').write_text(json.dumps(datos), encoding='utf-8')
    job = client.post(f'/api/workspaces/{wid}/informes/estudio', json=FORMULARIO)
    assert job.status_code == 202
    fin = esperar_trabajo(client, job.json())
    assert fin['status'] == 'done', fin
    resultado = fin['result']
    informe = salida / resultado['informe']
    texto = '\n'.join(p.text for p in docx.Document(informe).paragraphs)
    assert 'CAÍDA DE TENSIÓN DEL CLIENTE X' in texto and 'El cliente X reclama baja tensión.' in texto
    assert 'AGROINDUSTRIAS DEL MODELO' not in texto and 'Sin proyecto' not in texto
    assert '0,900 p.u.' in texto
    assert resultado['estudios'][0] == 'estudio_sin_flujo_desequilibrado_a.json'


def test_sin_estudios_de_sin_proyecto_lo_dice(client, tmp_path):
    """017:RF-29 — sin el flujo de «Sin proyecto» guardado, el trabajo dice qué estudio falta."""
    wid = client.post('/api/workspaces').json()['id']
    fin = esperar_trabajo(client, client.post(f'/api/workspaces/{wid}/informes/estudio', json=FORMULARIO).json())
    assert fin['status'] == 'failed' and 'estudio de flujo' in (fin['error'] or '')


def _carga_aplicada(client, wid: str, salida: Path, *, hoja: str = 'formulario') -> None:
    """Lo que deja el paso 3 al aplicar una carga nueva: su informe y su plan."""
    for rol in ('sin', 'con'):
        fichero = next(salida.glob(f'estudio_{rol}_*.json'))
        datos = json.loads(fichero.read_text(encoding='utf-8'))
        datos['project'] = 'PRY'
        if rol == 'con':
            datos['medidas']['AL01']['barras'].append(['B_NUEVA', 0.97])
        fichero.write_text(json.dumps(datos), encoding='utf-8')
    (salida / 'cargas_pf_aplicacion_tok1.json').write_text(json.dumps({
        'project': 'PRY', 'cargas': [{'sed': 'SN1', 'alimentador': 'AL01', 'barra': 'B_NUEVA',
                                      'variacion': 'CN_SN1', 'distancia_m': 120.0}]}), encoding='utf-8')
    planes = salida.parent / 'planes'
    planes.mkdir(parents=True, exist_ok=True)
    (planes / 'cargas_pf_tok1.json').write_text(json.dumps({'nuevas': [
        {'sed': 'SN1', 'hoja': hoja, 'alimentador': 'AL01', 'kw': 800.0, 'kvar': 300.0, 'kva_instalado': 1000}]}),
        encoding='utf-8')


FACTIBILIDAD = {'project': 'PRY', 'alimentador': 'AL01', 'sed': 'SN1', 'fecha': '2026-10-15',
                'asunto': 'FACTIBILIDAD DE SUMINISTRO', 'factor_carga': 0.6,
                'limites': FORMULARIO['limites'],
                'solicitud': {'cliente': 'EMPRESA X', 'ubicacion': 'Fundo Y, Ica', 'tipo_solicitud': 'factibilidad',
                              'potencia_kw': 800, 'fecha_solicitud': '2026-09-01', 'set_nombre': 'SET Z',
                              'tension_kv': 10, 'transformador': '25 MVA 60/10 kV', 'expediente': 'EXP-1',
                              'en_concesion': True},
                'perdidas_reconocidas': {'reconocido_kw_pct': 5, 'reconocido_kwh_pct': 4, 'precio_compra': 0.2},
                'transformador_set': {'kv': 10, 'sin_pct': 70, 'con_pct': 74, 'caso': '2026'}}


def test_contexto_dice_lo_que_ya_guardo_la_web(client, tmp_path):
    """017:RF-75, 017:RF-89, 017:RF-104 — estudios, cargas aplicadas e informes, sin DIgSILENT."""
    wid, salida = _espacio(client, tmp_path)
    _carga_aplicada(client, wid, salida)
    r = client.get(f'/api/workspaces/{wid}/informes', params={'project': 'PRY', 'feeder': ['AL01']})
    assert r.status_code == 200, r.text
    ctx = r.json()
    assert ctx['faltan'] == {'factibilidad': [], 'estudio': []}
    assert ctx['estudios']['sin']['alimentadores'] == {'AL01': True}
    assert [(c['sed'], c['barra'], c['sola'], c['kw']) for c in ctx['cargas']] == [('SN1', 'B_NUEVA', True, 800.0)]
    otro = client.get(f'/api/workspaces/{wid}/informes', params={'project': 'OTRO', 'feeder': ['AL01']}).json()
    assert otro['cargas'] == [] and otro['estudios']['sin'] is None, 'los flujos de otro proyecto no valen'


def test_informe_de_factibilidad_con_lo_aplicado(client, tmp_path):
    """017:RF-47, 017:RF-52, 017:RF-53, 017:RF-59, 017:RF-71 — punto de diseño en la barra
    de la carga y veredicto con el transformador escrito por el operador."""
    wid, salida = _espacio(client, tmp_path)
    _carga_aplicada(client, wid, salida)
    job = client.post(f'/api/workspaces/{wid}/informes/factibilidad', json=FACTIBILIDAD)
    assert job.status_code == 202, job.text
    fin = esperar_trabajo(client, job.json())
    assert fin['status'] == 'done', fin
    res = fin['result']
    assert res['factible'] is True and res['excesos'] == []
    assert res['punto_diseno']['u_pu'] == 0.97 and round(res['punto_diseno']['kv_ll'], 2) == 9.7
    texto = '\n'.join(p.text for p in docx.Document(salida / res['informe']).paragraphs)
    assert 'EMPRESA X' in texto and 'B_NUEVA' in texto
    generados = client.get(f'/api/workspaces/{wid}/informes', params={'project': 'PRY'}).json()['generados']
    assert [(g['tipo'], g['docx']) for g in generados] == [('factibilidad', res['informe'])]


def test_factibilidad_rechaza_carga_en_bloque_y_formulario_incompleto(client, tmp_path):
    """017:RF-47, 017:RF-96 — antes de empezar: una carga del libro no vale y la solicitud
    sin cliente se rechaza con el campo."""
    wid, salida = _espacio(client, tmp_path)
    _carga_aplicada(client, wid, salida, hoja='Hoja1')
    r = client.post(f'/api/workspaces/{wid}/informes/factibilidad', json=FACTIBILIDAD)
    assert r.status_code == 400 and 'en bloque' in r.json()['detail']
    incompleto = {**FACTIBILIDAD, 'solicitud': {**FACTIBILIDAD['solicitud'], 'cliente': ''}}
    assert client.post(f'/api/workspaces/{wid}/informes/factibilidad', json=incompleto).status_code == 422


def test_formulario_se_conserva_al_reiniciar(client, tmp_path):
    """017:RF-82 — lo escrito en 4.4 vuelve con el contexto, también tras reiniciar la web."""
    wid, _ = _espacio(client, tmp_path)
    datos = {'cab': {'asunto': 'MI ASUNTO'}, 'antecedentes': 'Texto largo.'}
    r = client.put(f'/api/workspaces/{wid}/informes/formulario',
                   json={'project': 'PRY', 'tipo': 'estudio', 'datos': datos})
    assert r.status_code == 200, r.text
    assert client.put(f'/api/workspaces/{wid}/informes/formulario',
                      json={'project': 'PRY', 'tipo': 'otro', 'datos': {}}).status_code == 422
    with TestClient(create_app(tmp_path / 'web')) as otra:
        ctx = otra.get(f'/api/workspaces/{wid}/informes', params={'project': 'PRY'}).json()
    assert ctx['formularios'] == {'estudio': datos}


def test_libro_de_perdidas_se_lee_y_entra_en_el_informe(client, tmp_path, monkeypatch):
    """017:RF-62, 017:RF-93, 017:RF-94 — el libro subido da lo reconocido del alimentador y
    el informe lo usa si no se escribió a mano; sin columnas en el perfil, dice cuáles faltan."""
    import io

    import openpyxl

    wid, salida = _espacio(client, tmp_path)
    _carga_aplicada(client, wid, salida)
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Perdidas'
    hoja.append(['Alimentador', 'Sistema', 'Pérd kW', 'Pérd kWh', 'Precio'])
    hoja.append(['AL01', 'ICA URBANO', 5.1, 4.2, 0.25])
    contenido = io.BytesIO()
    libro.save(contenido)
    r = client.post(f'/api/workspaces/{wid}/informes/perdidas',
                    files={'file': ('perdidas.xlsx', contenido.getvalue())})
    assert r.status_code == 200, r.text
    sin_perfil = client.get(f'/api/workspaces/{wid}/informes', params={'project': 'PRY', 'feeder': ['AL01']})
    assert 'libro_perdidas' in sin_perfil.json()['libro_perdidas']['error']

    perfil = Path(tmp_path / 'perfil_empresa.json')
    datos = json.loads(perfil.read_text(encoding='utf-8'))
    datos['informes']['libro_perdidas'] = {
        'hoja': 'Perdidas', 'columnas': {'alimentador': 'Alimentador', 'sistema_electrico': 'Sistema',
                                         'reconocido_kw_pct': 'Pérd kW', 'reconocido_kwh_pct': 'Pérd kWh'},
        'precio_compra': {'columna': 'Precio'}}
    perfil.write_text(json.dumps(datos), encoding='utf-8')
    ctx = client.get(f'/api/workspaces/{wid}/informes', params={'project': 'PRY', 'feeder': ['AL01']}).json()
    lectura = ctx['libro_perdidas']['lecturas']['AL01']
    assert (lectura['sistema_electrico'], lectura['reconocido_kwh_pct'], lectura['precio_compra']) == \
        ('ICA URBANO', 4.2, 0.25)

    sin_mano = {**FACTIBILIDAD, 'perdidas_reconocidas': {}}
    fin = esperar_trabajo(client, client.post(f'/api/workspaces/{wid}/informes/factibilidad', json=sin_mano).json())
    assert fin['status'] == 'done', fin
    resumen = json.loads((salida / fin['result']['json']).read_text(encoding='utf-8'))
    assert resumen['libro']['sistema_electrico'] == 'ICA URBANO'
    assert resumen['perdidas']['actual']['no_reconocida_soles'] is not None


def test_transformador_con_el_delta_de_los_flujos_guardados(client, tmp_path, monkeypatch):
    """017:RF-65, 017:RF-67, 017:RF-67b — el Δ de cabecera «Con − Sin» de 4.2 va al guion de
    transmisión; nada se vuelve a simular en el proyecto del alimentador."""
    from igea_dgs.web import services

    wid, salida = _espacio(client, tmp_path)
    con = next(salida.glob('estudio_con_*.json'))
    datos = json.loads(con.read_text(encoding='utf-8'))
    datos['medidas']['AL01'].update(cabecera_kw=6800.0, cabecera_kvar=3300.0)
    con.write_text(json.dumps(datos), encoding='utf-8')
    planes = []

    def guion(_ws, _ctx, script, plan, **_kw):
        planes.append((script, plan))
        t = {'nombre': 'TR1', 'kv': 10.0, 'lado': 'buslv'}
        return 0, {'caso': plan.get('caso'), 'sin': {'converge': True, 'trafos': [{**t, 'devanado_pct': 40.0}]},
                   'con': {'converge': True, 'trafos': [{**t, 'devanado_pct': 45.0}]}, 'restituida': True}, \
            salida / 'trafo.json'

    monkeypatch.setattr(services, '_guion_pf', guion)
    job = client.post(f'/api/workspaces/{wid}/informes/transformador',
                      json={'project': '', 'alimentador': 'AL01', 'carga': 'LElCarmen10', 'caso': '2026_D',
                            'transmision': 'AD08'})
    assert job.status_code == 422, 'el proyecto del alimentador es obligatorio'
    job = client.post(f'/api/workspaces/{wid}/informes/transformador',
                      json={'project': 'PRY', 'alimentador': 'AL01', 'carga': 'LElCarmen10', 'caso': '2026_D',
                            'transmision': 'AD08'})
    fin = esperar_trabajo(client, job.json())
    assert fin['status'] == 'done', fin
    script, plan = planes[0]
    assert script == 'transformador_set_pf.py' and plan['project'] == 'AD08'
    assert (plan['delta_kw'], plan['delta_kvar']) == (800.0, 300.0)
    assert fin['result']['devanado'] == {'kv': 10.0, 'sin_pct': 40.0, 'con_pct': 45.0}


def test_informe_de_estudio_con_las_obras_ya_aplicadas(client, tmp_path):
    """017:RF-35 / 019:RF-12 — CA101, 2026-10-15: al aplicarse una obra sale de la lista del
    espacio, y el informe salía solo con el estado actual. Ahora toma las obras aplicadas
    cuya variación está activa en el flujo de «Con proyecto», con la consigna de cabecera
    como etapa de regulación."""
    from igea_dgs.web import services

    wid, salida = _espacio(client, tmp_path)
    con = salida / 'estudio_con_flujo_desequilibrado_b.json'
    con.write_text(json.dumps({'medidas': {'AL01': _medida(0.99)},
                               'obras_activas': ['Obra_seccion_AL01', 'Obra_consigna_AL01']}),
                   encoding='utf-8')
    ws = client.app.state.store.get(wid)
    services.registrar_obra_aplicada(ws, '', 'Obra_seccion_AL01', {
        'id': 'x', 'clase': 'seccion', 'alimentador': 'AL01', 'casos': ['Con proyecto'],
        'tramos': ['L1'], 'tipo': {'nombre': 'AA12003D', 'clase': 'TypLne'}, 'km': 6.58})
    services.registrar_obra_aplicada(ws, '', 'Obra_consigna_AL01', {
        'clase': 'consigna', 'alimentador': 'AL01', 'casos': ['Con proyecto'], 'consigna_pu': 1.05})
    services.registrar_obra_aplicada(ws, '', 'Obra_vieja_AL01', {
        'clase': 'condensador', 'alimentador': 'AL01', 'casos': ['Con proyecto'], 'barra': 'BARRA_VIEJA',
        'q_kvar': 300})
    ws.save()
    fin = esperar_trabajo(client, client.post(f'/api/workspaces/{wid}/informes/estudio',
                                              json=FORMULARIO).json())
    assert fin['status'] == 'done', fin
    assert fin['result']['etapa'] == 'conductor'
    texto = '\n'.join(p.text for p in docx.Document(salida / fin['result']['informe']).paragraphs)
    assert 'AA12003D' in texto and '1,05 p.u.' in texto
    assert 'BARRA_VIEJA' not in texto, 'una obra cuya variación no está activa no entra'
