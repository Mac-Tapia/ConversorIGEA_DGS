"""Dataset PIDE universal: contrato, custodia, extracción y publicación para LLM."""
from __future__ import annotations

from pathlib import Path
import json


CONTRATO = Path(__file__).parents[1] / 'config' / 'pide_universal_2017.json'


def test_contrato_cubre_articulos_formatos_y_etapas_sin_empresa():
    """Falla si se pierde un artículo/formato o una empresa contamina la metodología."""
    from igea_dgs.pide_dataset import cargar_contrato

    contrato = cargar_contrato(CONTRATO)

    assert [a.numero for a in contrato.articulos] == list(range(1, 21))
    assert [f.codigo for f in contrato.formatos] == [f'F{i:02d}' for i in range(1, 16)]
    assert {a.etapa for a in contrato.articulos} == {
        'general', 'planeamiento_largo_plazo', 'pide', 'entrega'
    }
    assert contrato.validar() == []
    serializado = contrato.as_dict()
    assert 'ELOR' not in str(serializado).upper()
    assert all(a.entradas and a.procedimiento and a.evidencias and a.entregables
               for a in contrato.articulos)
    assert {'reconfiguracion', 'cambio_conductor', 'nuevo_alimentador',
            'nueva_subestacion', 'compensacion_reactiva'} <= {
                a.codigo for a in contrato.alternativas
            }


def test_inventario_conserva_custodia_clasifica_y_deduplica(tmp_path):
    """Falla si una copia se pierde, si falta SHA-256 o si se filtra la ruta absoluta."""
    from igea_dgs.pide_dataset import cargar_contrato, inventariar

    raiz = tmp_path / 'PIDE EMPRESA'
    entrega = raiz / 'PIDE 2026' / '4.2 INFORME PIDE' / '11.1 ANEXO 01'
    entrega.mkdir(parents=True)
    (entrega / 'FORMATO 08.xlsx').write_bytes(b'PK\x03\x04contenido-igual')
    copia = raiz / 'PIDE 2021' / 'INformes' / '17.Levantamiento de obs. quinto informe'
    copia.mkdir(parents=True)
    (copia / 'FORMATO 08 copia.xlsx').write_bytes(b'PK\x03\x04contenido-igual')
    (entrega / '~$temporal.xlsx').write_bytes(b'lock')
    (raiz / 'ELOR25_V2').mkdir()
    (raiz / 'ELOR25_V2' / 'ELOR25_V2').write_bytes(b'TAPE' + bytes(20))

    archivos = inventariar(raiz, 'EMPRESA DEMO', cargar_contrato(CONTRATO))

    assert len(archivos) == 4
    assert len({a.relative_path for a in archivos}) == 4
    assert all(len(a.sha256) == 64 and a.empresa == 'EMPRESA DEMO' for a in archivos)
    assert sum(a.duplicate_of is not None for a in archivos) == 1
    temporal = next(a for a in archivos if a.temporal)
    assert 'temporal_office' in temporal.advertencias
    backup = next(a for a in archivos if a.signature_kind == 'sqlserver_backup')
    assert backup.extension == '' and backup.extractability == 'native_required'
    clasificado = next(a for a in archivos if a.relative_path.endswith('FORMATO 08.xlsx'))
    assert clasificado.periodo == '2026' and clasificado.entregable == '4.2 INFORME PIDE'
    assert clasificado.anexo == '11.1 ANEXO 01'
    serializado = '\n'.join(json.dumps(a.as_dict(), ensure_ascii=False) for a in archivos)
    assert str(tmp_path) not in serializado


def test_inventario_rechaza_archivo_que_cambia_durante_huella(tmp_path, monkeypatch):
    """Falla si se publica una huella cuando tamaño o fecha cambian durante la lectura."""
    from igea_dgs import pide_dataset

    ruta = tmp_path / 'cambia.bin'
    ruta.write_bytes(b'abc')
    real = Path.stat
    llamadas = {'n': 0}

    def stat_cambiante(self, *args, **kwargs):
        resultado = real(self, *args, **kwargs)
        if self == ruta:
            llamadas['n'] += 1
            if llamadas['n'] == 2:
                ruta.write_bytes(b'abcd')
                resultado = real(self, *args, **kwargs)
        return resultado

    monkeypatch.setattr(Path, 'stat', stat_cambiante)
    sha, error = pide_dataset.sha256_estable(ruta)
    assert sha == ''
    assert error == 'archivo_modificado_durante_lectura'


def test_extrae_docx_xlsx_csv_pdf_y_geojson_reales(tmp_path):
    """Falla si un formato legible no produce texto localizado o activos estructurados."""
    from docx import Document
    from openpyxl import Workbook
    from pypdf import PdfWriter
    from igea_dgs.pide_dataset import cargar_contrato, extraer_archivo, inventariar

    raiz = tmp_path / 'corpus'
    raiz.mkdir()
    doc = Document()
    doc.add_heading('Diagnóstico', level=1)
    doc.add_paragraph('El alimentador presenta caída de tensión y requiere alternativas.')
    doc.save(raiz / 'informe.docx')
    wb = Workbook()
    ws = wb.active
    ws.title = 'F08 Alternativas'
    ws.append(['Código', 'Alternativa', 'Costo'])
    ws.append(['P1', 'Nuevo alimentador', 100])
    wb.save(raiz / 'formatos.xlsx')
    (raiz / 'demanda.csv').write_text('sistema,anio,demanda\nS1,2026,10.5\n', encoding='utf-8')
    (raiz / 'densidad.geojson').write_text(json.dumps({
        'type': 'FeatureCollection', 'features': [{
            'type': 'Feature', 'properties': {'densidad': 2.5},
            'geometry': {'type': 'Point', 'coordinates': [-73.2, -3.7]}
        }]
    }), encoding='utf-8')
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    with (raiz / 'plano.pdf').open('wb') as stream:
        pdf.write(stream)

    inventario = inventariar(raiz, 'EMPRESA DEMO', cargar_contrato(CONTRATO))
    extraidas = {a.extension: extraer_archivo(a, raiz, max_chars=200) for a in inventario}

    assert extraidas['.docx'].status == 'ok'
    assert 'caída de tensión' in ' '.join(f['text'] for f in extraidas['.docx'].fragments)
    xlsx = extraidas['.xlsx']
    assert xlsx.status == 'ok' and xlsx.assets[0]['sheet'] == 'F08 Alternativas'
    assert xlsx.assets[0]['dimension'] == 'A1:C2'
    # 027:RF-48: el contenido de la hoja y su cabecera, no solo su nombre.
    assert xlsx.assets[0]['cabecera'] == ['Código', 'Alternativa', 'Costo']
    assert 'Nuevo alimentador' in ' '.join(f['text'] for f in xlsx.fragments)
    assert extraidas['.csv'].assets[0]['columns'] == ['sistema', 'anio', 'demanda']
    assert extraidas['.geojson'].assets[0]['features'] == 1
    assert extraidas['.pdf'].status == 'ok'
    assert extraidas['.pdf'].assets[0]['pages'] == 1


def test_extraccion_no_finge_lectura_de_duplicados_corruptos_o_propietarios(tmp_path):
    """Falla si un binario, copia o ZIP roto se presenta como contenido aprendido."""
    from igea_dgs.pide_dataset import cargar_contrato, extraer_archivo, inventariar

    raiz = tmp_path / 'corpus'
    raiz.mkdir()
    (raiz / 'a.pdf').write_bytes(b'%PDF-roto')
    (raiz / 'b.pdf').write_bytes(b'%PDF-roto')
    (raiz / 'modelo.pfd').write_bytes(b'binario-nativo')
    inventario = inventariar(raiz, 'EMPRESA DEMO', cargar_contrato(CONTRATO))

    resultados = {a.relative_path: extraer_archivo(a, raiz) for a in inventario}

    assert resultados['a.pdf'].status == 'error'
    assert resultados['b.pdf'].status == 'duplicate'
    assert resultados['b.pdf'].fragments == ()
    assert resultados['modelo.pfd'].status == 'native_required'
    assert resultados['modelo.pfd'].fragments == ()


def test_publicacion_llm_es_atomica_determinista_y_tiene_custodia(tmp_path):
    """Falla si el corpus cambia sin cambiar entradas o publica texto sin procedencia."""
    from docx import Document
    from igea_dgs.pide_dataset import construir_dataset

    raiz = tmp_path / 'corpus'
    raiz.mkdir()
    doc = Document()
    doc.add_paragraph('Diagnóstico de caída de tensión en el alimentador de prueba.')
    doc.save(raiz / 'informe.docx')
    (raiz / 'copia.docx').write_bytes((raiz / 'informe.docx').read_bytes())
    destino = tmp_path / 'dataset'

    resumen1 = construir_dataset(raiz, destino, 'EMPRESA DEMO', CONTRATO)
    snapshot1 = {p.name: p.read_bytes() for p in destino.iterdir() if p.is_file()}
    resumen2 = construir_dataset(raiz, destino, 'EMPRESA DEMO', CONTRATO)
    snapshot2 = {p.name: p.read_bytes() for p in destino.iterdir() if p.is_file()}

    assert resumen1 == resumen2
    assert snapshot1 == snapshot2
    assert set(snapshot1) == {'manifest.jsonl', 'fragments.jsonl', 'assets.jsonl',
                              'training.jsonl', 'quality.jsonl', 'summary.json'}
    assert resumen1['files'] == 2 and resumen1['unique_contents'] == 1
    assert resumen1['articles'] == 20 and resumen1['training_records'] == 20
    fragments = [json.loads(x) for x in snapshot1['fragments.jsonl'].decode().splitlines()]
    assert fragments and all(f['file_id'] and len(f['sha256']) == 64 for f in fragments)
    assert all(f['estado'] == 'context' and f['empresa'] == 'EMPRESA DEMO' for f in fragments)
    assert str(tmp_path) not in '\n'.join(x.decode(errors='ignore') for x in snapshot1.values())
    training = [json.loads(x) for x in snapshot1['training.jsonl'].decode().splitlines()]
    assert all(r['estado'] == 'context' and len(r['messages']) == 3 for r in training)


def test_cli_construye_fixture_y_permite_omitir_texto(tmp_path):
    """Falla si el operador no puede generar solo custodia sin extraer texto empresarial."""
    from tools.construir_dataset_pide import main

    raiz = tmp_path / 'corpus'
    raiz.mkdir()
    (raiz / 'nota.txt').write_text('evidencia', encoding='utf-8')
    destino = tmp_path / 'salida'
    rc = main(['--raiz', str(raiz), '--destino', str(destino),
               '--empresa', 'EMPRESA DEMO', '--contrato', str(CONTRATO), '--sin-texto'])
    assert rc == 0
    assert (destino / 'manifest.jsonl').is_file()
    assert (destino / 'fragments.jsonl').read_text(encoding='utf-8') == ''


def test_dataset_real_pideelor_cubre_cada_archivo_y_articulo(tmp_path):
    """Aceptación local: si existe PIDEELOR, su publicación no puede quedar parcial.

    Se construye aquí mismo y se compara con los ficheros de la fuente: leía un
    ``output/pide_dataset`` que solo existía si alguien lo había generado antes, y fijaba
    1397 ficheros, que cambian con cada expediente (2026-10-27).
    """
    from igea_dgs.pide_dataset import construir_dataset

    repo = Path(__file__).parents[1]
    fuente = repo / 'referencia' / 'PIDEELOR'
    if not fuente.is_dir():
        return  # El corpus empresarial es privado y no forma parte del repositorio.
    destino = tmp_path / 'PIDEELOR'
    resumen = construir_dataset(fuente, destino, 'PIDEELOR', CONTRATO, sin_texto=True)
    n = sum(1 for p in fuente.rglob('*') if p.is_file())
    assert resumen['files'] == n > 0
    assert resumen['articles'] == 20 and resumen['formats'] == 15
    manifest = [json.loads(x) for x in (destino / 'manifest.jsonl').read_text(
        encoding='utf-8').splitlines()]
    assert len(manifest) == n
    assert len({r['relative_path'] for r in manifest}) == n
    assert all(len(r['sha256']) == 64 for r in manifest)
    publicados = '\n'.join(p.read_text(encoding='utf-8', errors='ignore')
                            for p in destino.glob('*.json*'))
    assert str(repo) not in publicados
    assert '"estado":"gold"' not in publicados


def test_docx_que_python_docx_no_abre_se_lee_igual(tmp_path, monkeypatch):
    """027:RF-46, 027:RF-51 — un Word antiguo con «\» en sus rutas internas (P&C-059-2022 de
    PIDEELOR) hacía fallar a python-docx: su texto se lee de word/document.xml y se avisa."""
    import docx
    from docx import Document
    from igea_dgs import pide_dataset as pd

    doc = Document()
    doc.add_paragraph('Alternativa de inversión por cambio de conductor.')
    doc.save(tmp_path / 'viejo.docx')

    def roto(*_a, **_k):
        raise KeyError("There is no item named 'word/media\\image1.jpg' in the archive")

    monkeypatch.setattr(docx, 'Document', roto)
    fragments, assets, avisos = pd._extraer_docx(tmp_path / 'viejo.docx', 200)
    assert 'cambio de conductor' in ' '.join(f['text'] for f in fragments)
    assert avisos == ['docx_sin_python_docx'] and assets[0]['kind'] == 'docx'
