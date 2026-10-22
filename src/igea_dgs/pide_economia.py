"""Evaluación económica y priorización del PIDE (spec 024, T14; RF-33, RF-34, RF-54 a RF-59).

Art. 14 de la norma prepublicada con la RM 428-2017-MEM/DM: costo de inversión (CI),
de operación y mantenimiento (COyM) y de pérdidas (CP); costo anual = anualidad del CI
+ COyM + CP, con 30 años de vida útil y la tasa del art. 79 de la LCE; VAN y TIR. Además,
el flujo de caja de la empresa (privada) con las tarifas del pliego y su tasa, porque el
usuario lo pidió el 2026-10-21 («tarifas, costos de energía vendida por la potencia nueva
instalada de nuevos clientes para un horizonte de 20 años, tasa a privados»).

Art. 15: entre las alternativas factibles técnica y económicamente, primero las de menor
costo/kW, costo/kW.h, costo/usuario y costo/km con rentabilidad aceptable; después los
criterios del art. 15.3 con los pesos del perfil. Puro: no lee DIgSILENT.
"""
from __future__ import annotations

import re
import unicodedata
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from statistics import median
from typing import Any

from .pide import RUBRO_DE_ALTERNATIVA
from .util import finite_or_none

#: Palabras de la denominación de una obra del CAPEX → alternativa del PIDE (RF-54).
#: Se prueban en este orden: «transformador de potencia» es de una SET, no de una SED.
PALABRAS_ALTERNATIVA = (
    ('nueva_set', ('subestacion de transformacion', 'transformador de potencia', ' set ', 'nueva set')),
    ('bancos', ('banco de condensador', 'bancos de condensador', 'compensacion reactiva')),
    ('reguladores', ('regulador',)),
    ('nueva_terna', ('terna',)),
    ('cambio_seccion', ('cambio de seccion', 'cambio de conductor', 'reconductor', 'repotenci')),
    ('refuerzo_red', ('enlace', 'interconexion', 'migracion', 'transferencia', 'refuerzo')),
    ('nuevo_alimentador', ('nuevo alimentador', 'alimentador nuevo', 'nuevos alimentadores')),
    ('nueva_sed', ('sed', 'subestacion de distribucion', 'transformador de distribucion')),
)


class EconomiaPideError(ValueError):
    """Falta un dato del perfil o del libro: se dice cuál, nunca se supone (RF-58e)."""


def _plano(texto: Any) -> str:
    s = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    return f" {re.sub(r'[^a-z0-9]+', ' ', s.lower()).strip()} "


# ------------------------------------------------------------------ libro CAPEX (RF-54 a 56)
@dataclass(frozen=True)
class ObraCapex:
    nombre: str
    zona: str
    objetivo: str
    descripcion: str
    unidad: str
    cantidad: float | None
    costo_unitario: float | None
    inversion: float | None
    alternativa: str | None


def alternativa_de(*textos: str) -> str | None:
    """La alternativa del PIDE que nombra una obra del CAPEX, o ``None``."""
    plano = _plano(' '.join(textos))
    for alternativa, palabras in PALABRAS_ALTERNATIVA:
        # Al principio de palabra: «repotenci» vale para «repotenciación».
        if any(re.search(rf'\b{re.escape(p.strip())}', plano) for p in palabras):
            return alternativa
    return None




def leer_capex(ruta: Path | str, config: Mapping[str, Any]) -> tuple[list[ObraCapex], list[str]]:
    """Obras del libro CAPEX con su alternativa del PIDE, por los nombres del perfil.

    Solo se leen las columnas del perfil: responsables, líderes y contratos del libro no
    salen de aquí (RF-56). Devuelve las obras y los avisos (RF-55): hoja o columna que no
    está y alternativas del PIDE sin ninguna obra en el libro.
    """
    from openpyxl import load_workbook

    columnas = dict(config.get('columnas') or {})
    hoja_nombre = config.get('hoja')
    libro = load_workbook(ruta, read_only=True, data_only=True)
    if hoja_nombre not in libro.sheetnames:
        raise EconomiaPideError(f'el libro no tiene la hoja «{hoja_nombre}» del perfil')
    hoja = libro[hoja_nombre]
    filas = hoja.iter_rows(values_only=True)
    marca = _plano(config.get('fila_cabecera_contiene') or columnas.get('costo_unitario'))
    indice: dict[str, int] = {}
    for fila in filas:
        planos = [_plano(v) for v in fila]
        if marca in planos:
            for clave, nombre in columnas.items():
                if _plano(nombre) in planos:
                    indice[clave] = planos.index(_plano(nombre))
            break
    faltan = [c for c in columnas if c not in indice]
    if 'nombre' in faltan or 'costo_unitario' in faltan:
        raise EconomiaPideError(f'no se encontraron las columnas {faltan} en «{hoja_nombre}»')
    avisos = [f'columna «{columnas[c]}» del perfil no encontrada' for c in faltan]

    def celda(fila, clave):
        i = indice.get(clave)
        return fila[i] if i is not None and i < len(fila) else None

    obras = []
    for fila in filas:
        nombre = celda(fila, 'nombre')
        if not nombre or _plano(nombre) == _plano(columnas['nombre']):
            continue
        textos = [str(celda(fila, k) or '') for k in ('nombre', 'objetivo', 'descripcion')]
        obras.append(ObraCapex(str(nombre).strip(), str(celda(fila, 'zona') or ''),
                               str(celda(fila, 'objetivo') or ''), str(celda(fila, 'descripcion') or ''),
                               str(celda(fila, 'unidad') or ''), finite_or_none(celda(fila, 'cantidad')),
                               finite_or_none(celda(fila, 'costo_unitario')), finite_or_none(celda(fila, 'inversion')),
                               alternativa_de(*textos)))
    libro.close()
    sin = sorted(set(RUBRO_DE_ALTERNATIVA) - {o.alternativa for o in obras if o.alternativa})
    if sin:
        avisos.append(f'alternativas del PIDE sin obra en el libro CAPEX: {", ".join(sin)}')
    return obras, avisos


def costos_por_alternativa(obras: Iterable[ObraCapex]) -> dict[str, dict[str, Any]]:
    """Costo unitario de cada alternativa: la mediana de sus obras, con su unidad y cuántas."""
    grupos: dict[str, list[ObraCapex]] = {}
    for o in obras:
        if o.alternativa and o.costo_unitario:
            grupos.setdefault(o.alternativa, []).append(o)
    out = {}
    for alt, lista in grupos.items():
        unidades = sorted({o.unidad for o in lista if o.unidad})
        out[alt] = {'costo_unitario': median(o.costo_unitario for o in lista),
                    'unidad': unidades[0] if len(unidades) == 1 else '/'.join(unidades),
                    'obras': len(lista), 'denominacion': lista[0].nombre}
    return out


# ------------------------------------------------------------------ art. 14 (RF-57, RF-58)
def frc(tasa: float, anios: int) -> float:
    """Factor de recuperación de capital: anualidad de 1 en ``anios`` a ``tasa``."""
    if tasa == 0:
        return 1 / anios
    f = (1 + tasa) ** anios
    return tasa * f / (f - 1)


def costo_anual(ci: float, coym: float | None, cp: float | None, *, tasa: float, vida: int) -> dict[str, float | None]:
    """Costo anual = anualidad del CI + COyM + CP (art. 14.4); lo que falta queda dicho."""
    anualidad = ci * frc(tasa, vida)
    return {'anualidad': anualidad, 'coym': coym, 'cp': cp,
            'costo_anual': anualidad + (coym or 0.0) + (cp or 0.0),
            'completo': coym is not None and cp is not None}


def van(flujos: Sequence[float], tasa: float) -> float:
    """Valor actual neto de flujos anuales desde el año 0."""
    return sum(f / (1 + tasa) ** t for t, f in enumerate(flujos))


def tir(flujos: Sequence[float]) -> float | None:
    """Tasa interna de retorno por bisección; ``None`` si no cambia de signo."""
    lo, hi = -0.99, 10.0
    f_lo, f_hi = van(flujos, lo), van(flujos, hi)
    if f_lo * f_hi > 0:
        return None
    for _ in range(200):
        mid = (lo + hi) / 2
        f_mid = van(flujos, mid)
        if abs(f_mid) < 1e-9:
            return mid
        if f_lo * f_mid < 0:
            hi, f_hi = mid, f_mid
        else:
            lo, f_lo = mid, f_mid
    return (lo + hi) / 2


def retorno(flujos: Sequence[float]) -> float | None:
    """Años hasta recuperar la inversión (flujo acumulado ≥ 0), con interpolación."""
    acumulado = 0.0
    for t, f in enumerate(flujos):
        anterior = acumulado
        acumulado += f
        if t > 0 and acumulado >= 0 > anterior:
            return t - 1 + (-anterior / f if f else 0)
    return None


# ------------------------------------------------------------------ flujo privado (RF-58c, d)
def ingreso_anual_kw(kw: float, tarifa: Mapping[str, Any], *, factor_carga: float,
                     fraccion_punta: float) -> float:
    """Lo que factura en un año 1 kW nuevo con la opción tarifaria dada (sin IGV).

    Energía: kW × factor de carga × 8760 h a su cargo (en MT3, punta y fuera de punta con
    ``fraccion_punta``). Potencia: generación y distribución, por mes.
    """
    energia = kw * factor_carga * 8760
    if 'energia_kwh' in tarifa:
        por_energia = energia * tarifa['energia_kwh']
    else:
        por_energia = energia * (fraccion_punta * tarifa['energia_punta_kwh']
                                 + (1 - fraccion_punta) * tarifa['energia_fuera_punta_kwh'])
    por_potencia = 12 * kw * (tarifa.get('potencia_generacion_kw_mes', 0.0)
                              + tarifa.get('potencia_distribucion_kw_mes', 0.0))
    return por_energia + por_potencia


@dataclass
class Evaluacion:
    ci: float
    regulada: dict[str, Any]
    privada: dict[str, Any] | None
    indicadores: dict[str, float | None]
    avisos: list[str] = field(default_factory=list)


def evaluar(ci: float, *, kw_nuevos: float, usuarios: int, km: float, coym_anual: float | None,
            cp_anual: float | None, plan: Mapping[str, Any], factor_carga: float | None,
            precio_compra_kwh: float | None, opcion: str, fraccion_punta: float = 5 / 24,
            crecimiento: float = 0.0) -> Evaluacion:
    """Evaluación regulada (art. 14) y privada (flujo de 20 años) de una alternativa."""
    avisos = []
    vida, tasa_r = int(plan['vida_util_anios']), float(plan['tasa_regulada'])
    horizonte = int(plan['horizonte_anios'])
    regulada = costo_anual(ci, coym_anual, cp_anual, tasa=tasa_r, vida=vida)
    flujo_costos = [-ci] + [-(coym_anual or 0) - (cp_anual or 0)] * horizonte
    regulada['van_costos'] = van(flujo_costos, tasa_r)
    if not regulada['completo']:
        avisos.append('costo anual sin COyM o sin pérdidas: el perfil o el flujo no los dan')
    energia_anual = kw_nuevos * (factor_carga or 0) * 8760
    indicadores = {
        'costo_kw': ci / kw_nuevos if kw_nuevos else None,
        'costo_kwh': ci / energia_anual if energia_anual else None,
        'costo_usuario': ci / usuarios if usuarios else None,
        'costo_km': ci / km if km else None,
    }
    tarifa = (plan.get('tarifas') or {}).get(opcion)
    tasa_p = plan.get('tasa_privada')
    if tarifa is None or tasa_p is None or factor_carga is None or precio_compra_kwh is None:
        avisos.append('evaluación privada sin hacer: faltan tarifa, tasa privada, factor de carga '
                      'o precio de compra (RF-58e)')
        return Evaluacion(ci, regulada, None, indicadores, avisos)
    flujos = [-ci]
    for t in range(1, horizonte + 1):
        kw = kw_nuevos * (1 + crecimiento) ** t
        ingresos = ingreso_anual_kw(kw, tarifa, factor_carga=factor_carga, fraccion_punta=fraccion_punta)
        compra = kw * factor_carga * 8760 * precio_compra_kwh
        flujos.append(ingresos - compra - (coym_anual or 0) - (cp_anual or 0))
    beneficios = sum(f / (1 + tasa_p) ** t for t, f in enumerate(flujos) if t > 0)
    privada = {'van': van(flujos, float(tasa_p)), 'tir': tir(flujos), 'retorno_anios': retorno(flujos),
               'beneficio_costo': beneficios / ci if ci else None, 'flujos': flujos}
    return Evaluacion(ci, regulada, privada, indicadores, avisos)


# ------------------------------------------------------------------ art. 15 (RF-33, RF-34)
@dataclass
class Candidata:
    codigo: str
    alternativa: str
    problema: str
    cumple_norma: bool
    evaluacion: Evaluacion
    criterios: dict[str, float] = field(default_factory=dict)   # 0 a 1, del art. 15.3
    puntaje: float = 0.0
    sobredimensionada: bool = False
    descartada: str | None = None

    @property
    def rubro(self) -> str:
        return RUBRO_DE_ALTERNATIVA.get(self.alternativa, 'ampliacion_reforzamiento')


INDICADORES = ('costo_kw', 'costo_kwh', 'costo_usuario', 'costo_km')


def priorizar(candidatas: Sequence[Candidata], pesos: Mapping[str, float]) -> list[Candidata]:
    """Orden del art. 15 (RF-33, RF-34, RF-58b).

    1. Fuera: las que no cumplen la norma y las de rentabilidad no aceptable (VAN privado
       negativo cuando lo hay).
    2. Puntaje: por cada indicador de costo, el menor de todas entre el suyo (1 para la más
       barata); más los criterios del art. 15.3 por sus pesos. Mayor puntaje, antes.
    3. Para un mismo problema, las que cumplen con mayor costo que la elegida quedan
       marcadas como sobredimensionadas.
    """
    validas = []
    for c in candidatas:
        if not c.cumple_norma:
            c.descartada = 'no cumple la norma'
        elif (c.rubro == 'ampliacion_reforzamiento' and c.evaluacion.privada is not None
              and c.evaluacion.privada['van'] < 0):
            # Solo la ampliación se paga con la energía que vende; el mejoramiento se
            # justifica por la calidad que exige la norma (art. 12.4 b).
            c.descartada = 'rentabilidad no aceptable (VAN privado negativo)'
        else:
            validas.append(c)
    minimos = {i: min((v for c in validas if (v := c.evaluacion.indicadores.get(i))), default=None)
               for i in INDICADORES}
    for c in validas:
        p = 0.0
        for i, m in minimos.items():
            v = c.evaluacion.indicadores.get(i)
            if m and v:
                p += m / v
        p += sum(float(pesos.get(k, 0)) * float(v) for k, v in c.criterios.items())
        c.puntaje = p
    orden = sorted(validas, key=lambda c: (-c.puntaje, c.evaluacion.ci, c.codigo))
    elegida: dict[str, Candidata] = {}
    for c in orden:
        if c.problema not in elegida:
            elegida[c.problema] = c
        elif c.evaluacion.ci > elegida[c.problema].evaluacion.ci:
            c.sobredimensionada = True
    return orden + [c for c in candidatas if c.descartada]


# ------------------------------------------------------------------ propuesta (RF-45 a RF-47)
#: Causas de calidad del art. 15.3 («subsanación de deficiencias por calidad de producto»).
CAUSAS_DE_CALIDAD = ('caida_de_tension', 'factor_de_potencia', 'sobretension')


def _problemas_por_anio(informe: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Cada problema con el primer año en que aparece en DIgSILENT (RF-51, RF-53)."""
    primero: dict[tuple, dict] = {}
    for fila in sorted(informe.get('anios') or [], key=lambda a: a['anio']):
        for p in fila.get('problemas') or []:
            clave = (p.get('causa'), p.get('elemento'))
            if clave not in primero:
                primero[clave] = {**p, 'anio': fila['anio']}
    return list(primero.values())


def _obra_011(alternativa: str, alimentador: str | None, elemento: str, kv: float | None,
              bancos_kvar: Mapping[float, float] | None) -> tuple[dict | None, str | None]:
    """La obra de la spec 011 que construye la alternativa, o el motivo de no tenerla."""
    if not alimentador:
        return None, 'no se sabe el alimentador del problema'
    if alternativa == 'bancos':
        if kv is None or not bancos_kvar:
            return None, 'falta el tamaño de banco por tensión en el perfil'
        kvar = min(bancos_kvar.items(), key=lambda x: abs(x[0] - kv))[1]
        return {'clase': 'condensador', 'alimentador': alimentador, 'barra': elemento,
                'q_kvar': kvar, 'escalones': 1, 'casos': ['Con proyecto']}, None
    if alternativa == 'nueva_terna':
        return {'clase': 'terna', 'alimentador': alimentador,
                'tramos': [elemento.removesuffix('.ElmLne')], 'casos': ['Con proyecto']}, None
    return None, (f'«{alternativa}» necesita elegir un tipo de la biblioteca o un diseño: queda '
                  'propuesta para el operador')


def propuesta_pide(informe: Mapping[str, Any], plan: Mapping[str, Any], *,
                   costos: Mapping[str, Mapping[str, Any]], bancos_kvar: Mapping[float, float] | None,
                   factor_carga: float | None = None, precio_compra_kwh: float | None = None
                   ) -> dict[str, Any]:
    """Del diagnóstico año a año a las obras priorizadas, con su evaluación (RF-45, RF-52,
    RF-57 a RF-59), su obra de la spec 011 si se puede construir sola (RF-47) y las filas
    de los formatos (RF-63)."""
    from .pide import (
        ALTERNATIVAS_POR_CAUSA,
        CARACTER_DE_ALTERNATIVA,
        Problema,
        plan_por_etapas,
    )

    avisos: list[str] = []
    problemas = _problemas_por_anio(informe)
    anios = sorted(informe.get('anios') or [], key=lambda a: a['anio'])
    cab0 = (anios[0].get('cabeceras_kw') or {}) if anios else {}
    # Una caída de tensión en baja se resuelve en la SED, no con un banco de media.
    for p in problemas:
        if p.get('causa') == 'caida_de_tension' and p.get('kv') is not None and p['kv'] < 1:
            p['causa'] = 'transformador_sobrecargado'
    por_alim: dict[str | None, list[dict]] = {}
    for p in problemas:
        por_alim.setdefault(p.get('alimentador'), []).append(p)
    costo_rel = {a: c['costo_unitario'] for a, c in costos.items()}
    if not costos:
        avisos.append('sin libro CAPEX: las obras van sin costo y se eligen por menor intervención')
    candidatas, filas_obra = [], []
    n = 0
    for alimentador, lista in sorted(por_alim.items(), key=lambda x: str(x[0])):
        obras = plan_por_etapas([Problema(p['tipo'], p['elemento'], p.get('valor'), p.get('limite'),
                                          p['causa'], p.get('evidencia', ''), int(p['anio']))
                                 for p in lista if p.get('causa') in ALTERNATIVAS_POR_CAUSA],
                                anios_pide=plan.get('anios_pide') or [4], costo_relativo=costo_rel or None)
        for o in obras:
            n += 1
            codigo = f'AL{n:02d}'
            grupo = [p for p in lista if p['causa'] == o.causa and any(
                x.startswith(f"{p['elemento']}:") for x in o.problemas)]
            # El banco va en la barra de media más baja del grupo, donde más sube la tensión.
            mt = [p for p in grupo if (p.get('kv') or 0) >= 1 and p.get('valor') is not None]
            sitio = min(mt, key=lambda p: p['valor']) if mt else (grupo[0] if grupo else {})
            elemento = sitio.get('elemento', o.elemento)
            costo = costos.get(o.alternativa)
            ci = float(costo['costo_unitario']) if costo else 0.0
            if costo and str(costo.get('unidad', '')).lower().startswith('km'):
                avisos.append(f'{codigo}: costo por km con 1 km supuesto; el largo lo da el diseño')
            cab_obra = next((a.get('cabeceras_kw') or {} for a in anios if a['anio'] >= o.anio), cab0)
            kw = 0.0
            if o.rubro == 'ampliacion_reforzamiento' and alimentador:
                kw = max(0.0, (cab_obra.get(alimentador) or 0) - (cab0.get(alimentador) or 0))
            ev = evaluar(ci, kw_nuevos=kw, usuarios=0, km=0, coym_anual=None, cp_anual=None, plan=plan,
                         factor_carga=factor_carga, precio_compra_kwh=precio_compra_kwh,
                         opcion=str(plan.get('opcion_tarifaria_mt') or 'MT3'))
            criterios = {'subsanacion_calidad': 1.0 if o.causa in CAUSAS_DE_CALIDAD else 0.0,
                         'crecimiento_demanda': 1.0 if kw > 0 else 0.0}
            c = Candidata(codigo, o.alternativa, f'{alimentador}:{o.causa}', True, ev, criterios)
            obra, motivo = _obra_011(o.alternativa, alimentador, elemento, sitio.get('kv'), bancos_kvar)
            candidatas.append(c)
            filas_obra.append({'codigo': codigo, 'anio': o.anio, 'alternativa': o.alternativa,
                               'rubro': o.rubro, 'caracter': CARACTER_DE_ALTERNATIVA[o.alternativa],
                               'alimentador': alimentador, 'elemento': elemento,
                               'causa': o.causa, 'problemas': o.problemas[:20],
                               'denominacion': (costo or {}).get('denominacion'), 'ci': ci,
                               'obra_011': obra, 'motivo_sin_obra': motivo, 'candidata': c})
    orden = priorizar(candidatas, plan.get('pesos_art_15_3') or {})
    prioridad = {c.codigo: i + 1 for i, c in enumerate(orden)}
    for f in filas_obra:
        c = f.pop('candidata')
        f.update(prioridad=prioridad[c.codigo], puntaje=c.puntaje, sobredimensionada=c.sobredimensionada,
                 descartada=c.descartada, regulada=c.evaluacion.regulada, privada=c.evaluacion.privada,
                 indicadores=c.evaluacion.indicadores, avisos=c.evaluacion.avisos)
    filas_obra.sort(key=lambda f: f['prioridad'])
    return {'obras': filas_obra, 'avisos': avisos, 'formatos': _filas_formatos(filas_obra, anios)}


def _filas_formatos(obras: Sequence[Mapping[str, Any]], anios: Sequence[Mapping[str, Any]]) -> dict[str, list]:
    """Filas de los formatos 03, 07, 08, 09, 10 y 12 con lo calculado (RF-63)."""
    f = {}
    if anios:
        f['F03 Tensión MT'] = [['', alim, '', kw / 1000 if kw else None, '', '', '',
                                round(100 * (1 - (anios[0].get('vm_min') or 1)), 2)]
                               for alim, kw in sorted((anios[0].get('cabeceras_kw') or {}).items())]
    f['F07 Alternativas estudio'] = [[o['anio'], o['rubro'], o['codigo'], o['alternativa'], o['ci'],
                                      o['indicadores'].get('costo_kw'), o['indicadores'].get('costo_kwh'),
                                      None, None] for o in obras]
    f['F08 Alternativas PIDE'] = [[o['anio'], o['rubro'], o['codigo'], o['alternativa'], None, None,
                                   None, None, None, None] for o in obras]
    f['F09 Evaluación económica'] = [[o['anio'], o['rubro'], o['codigo'], o['alternativa'], o['ci'],
                                      o['indicadores'].get('costo_kw'), o['indicadores'].get('costo_kwh'),
                                      None, None, (o['privada'] or {}).get('van'),
                                      o['regulada'].get('costo_anual'), (o['privada'] or {}).get('tir')]
                                     for o in obras]
    resumen: dict[tuple, list[float]] = {}
    for o in obras:
        if o['descartada'] or o['sobredimensionada']:
            continue
        r = resumen.setdefault((o['anio'], o['rubro']), [0.0, 0.0])
        r[0] += o['ci']
        r[1] += o['regulada'].get('anualidad') or 0.0
    f['F10 Propuesta empresa'] = [[a, rubro, None, None, None, ci, anu, None, None]
                                  for (a, rubro), (ci, anu) in sorted(resumen.items())]
    f['F12 Propuesta proyecto'] = [[o['alimentador'], o['anio'], o['rubro'], o['codigo'],
                                    o['denominacion'] or o['alternativa'], o['ci'],
                                    o['regulada'].get('anualidad'), None, o['regulada'].get('costo_anual')]
                                   for o in obras if not (o['descartada'] or o['sobredimensionada'])]
    return f
