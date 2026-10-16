"""Spec 017: los dos informes en Word sobre un modelo sintético, como esquemas generales."""
from __future__ import annotations

import json
import struct
import zipfile
import zlib
from datetime import date
from pathlib import Path

import docx
import pytest

from igea_dgs.informe_docx import (
    DatosEstudio, DatosFactibilidad, abrir_modelo, escribir_informe_estudio, escribir_informe_factibilidad)
from igea_dgs.informes import (
    Caso, ConfigInforme, Devanado, Encabezado, Etapa, Extremo, InformeError, Potencia, Solicitud,
    perdidas_anuales)
from igea_dgs.libros_informe import DatosPerdidas

#: Datos del modelo que no pueden aparecer en ningún informe generado.
DEL_MODELO = ('GT-019-PR', 'CAIDA DE TENSIÓN DEL MODELO', '9999', '3 de 6', 'JAVIER')


def _png(ruta: Path) -> Path:
    """Un PNG de 1×1 válido, sin depender de Pillow."""
    def bloque(tipo, datos):
        return struct.pack('>I', len(datos)) + tipo + datos + struct.pack('>I', zlib.crc32(tipo + datos))
    ruta.write_bytes(b'\x89PNG\r\n\x1a\n' + bloque(b'IHDR', struct.pack('>IIBBBBB', 1, 1, 8, 2, 0, 0, 0))
                     + bloque(b'IDAT', zlib.compress(b'\x00\xff\x00\x00')) + bloque(b'IEND', b''))
    return ruta


def _modelo(ruta: Path) -> Path:
    """Un modelo como los de la empresa: tabla en la cabecera, firmas y «de 6» fijo en el pie."""
    d = docx.Document()
    s = d.sections[0]
    cab = s.header.add_table(1, 2, width=s.page_width)
    cab.cell(0, 0).text = 'INFORME TECNICO Nº :\nGT-019-PR'
    cab.cell(0, 1).text = 'REFERENCIA:\nCAIDA DE TENSIÓN DEL MODELO'
    pie = s.footer.add_table(1, 2, width=s.page_width)
    pie.cell(0, 0).text = 'ELABORADO POR :'
    pie.cell(0, 1).paragraphs[0].add_run('3')
    pie.cell(0, 1).paragraphs[0].add_run(' de 6')
    d.add_paragraph('PARA : JAVIER, texto del modelo que no debe quedar: 9999 kW.')
    d.save(ruta)
    return ruta


def _config(tmp_path: Path, **cambios) -> ConfigInforme:
    (tmp_path / 'modelos').mkdir(exist_ok=True)
    for tipo in ('estudio', 'factibilidad'):
        _modelo(tmp_path / 'modelos' / f'{tipo}.docx')
    perfil = {
        'estudios': {'tension_min_pu': 0.95, 'tension_max_pu': 1.05, 'cargabilidad_max_pct': 100},
        'informes': {'modelo_estudio': 'modelos/estudio.docx', 'modelo_factibilidad': 'modelos/factibilidad.docx',
                     'bandas_tension_pct': [85, 90, 95, 105], 'bandas_cargabilidad_pct': [80, 90, 95, 105, 150],
                     'cabecera': {'numero': 'INFORME TECNICO', 'referencia': 'REFERENCIA'},
                     'norma': 'NTCSE', 'regulador': 'Osinergmin', 'proyecto_transmision': 'TRANSMISION'},
    }
    perfil['informes'].update(cambios)
    return ConfigInforme.de_perfil(perfil, tmp_path)


def _caso(nombre, minima, carga=60.0, converge=True, fases=False):
    if not converge:
        return Caso(nombre, False)
    tension = ({f: Extremo(minima + i / 100, f'N{i}') for i, f in enumerate('ABC')} if fases
               else {'trifásica': Extremo(minima, 'N0')})
    return Caso(nombre, True, fuente=Potencia(3000, 900), perdidas=Potencia(150, 120),
                tension_min=tension, tension_max=Extremo(1.0, 'CAB'),
                nodos_fuera=0 if minima >= 0.95 else 12, cargabilidad_max=Extremo(carga, 'T1'))


def _imagenes(tmp_path, *estados):
    img = _png(tmp_path / 'img.png')
    return {e: {'ubicacion': img, 'tension': img, 'cargabilidad': img} for e in estados}


def _estudio(tmp_path, actual=None, etapas=None, imagenes=None, **extra):
    datos = dict(
        alimentadores=['AL01'],
        encabezado=Encabezado(date(2026, 10, 15), 'CAÍDA DE TENSIÓN EN AL01', 'ING. X\nPROYECTOS Y OBRAS',
                              lugar='Ica'),
        numero='GT-100-PR',
        antecedentes=['El cliente reclama baja tensión.'],
        descripcion=['Se usó el modelo de la red del alimentador.'],
        actual=actual or _caso('Sin proyecto', 0.91, fases=True),
        etapas=etapas if etapas is not None else [
            Etapa('bancos_existentes', motivo='los bancos ya están en servicio'),
            Etapa('bancos_nuevos', caso=_caso('Con proyecto', 0.96),
                  obras=({'tipo': 'banco', 'nodo': 'N15', 'kvar': 300},), anio=2027),
        ],
        imagenes=imagenes if imagenes is not None else _imagenes(tmp_path, 'actual', 'bancos_nuevos'),
    )
    datos.update(extra)
    return DatosEstudio(**datos)


SOLICITUD = Solicitud(cliente='CLIENTE PRUEBA', ubicacion='el distrito X',
                      tipo_solicitud='Factibilidad y Punto de Diseño', potencia_kw=430,
                      fecha_solicitud=date(2026, 8, 1), alimentador='AL02', set_nombre='SET1', tension_kv=10,
                      transformador='50/30/30 MVA', expediente='EXP-1', en_concesion=True)


def _factibilidad(tmp_path, **cambios):
    datos = dict(
        encabezado=Encabezado(date(2026, 10, 15), 'FACTIBILIDAD CLIENTE PRUEBA', 'ING. Y', lugar='Ica'),
        actual=_caso('Sin proyecto', 0.97), proyectado=_caso('Con proyecto', 0.96, carga=70.0, fases=True),
        solicitud=SOLICITUD, codigo_punto='NMT 123',
        punto_diseno={'u_pu': 0.97, 'kv_ll': 9.7, 'kv_ln': 5.6, 'i_a': 27.1, 'kva': 455, 'kw': 430, 'kvar': 142},
        solicitudes_vigentes=True,
        perdidas_actual=perdidas_anuales(2791, 40.10, 0.74, 0.79, 0.27),
        perdidas_proyectado=perdidas_anuales(3233, 49.86, 0.74, 0.79, 0.27),
        libro=DatosPerdidas('AL02', 'ICA', 1.2, 0.79, 0.27, (), 'SE0044'),
        devanado=Devanado(10, 72.28, 73.87), caso_transmision='2026',
        imagenes={**_imagenes(tmp_path, 'actual', 'proyectado'),
                  'transformador': {'cargabilidad': _png(tmp_path / 'trafo.png')}},
        fuentes=['el modelo de la red en DIgSILENT PowerFactory'], numero='GT-200-PR')
    datos.update(cambios)
    return DatosFactibilidad(**datos)


def _texto(ruta: Path) -> str:
    """Texto del cuerpo en orden, y después el de la cabecera y el pie."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(ruta)
    partes = []
    for el in d.element.body.iterchildren():
        if el.tag.endswith('}p'):
            partes.append(Paragraph(el, d).text)
        elif el.tag.endswith('}tbl'):
            partes += [c.text for r in Table(el, d).rows for c in r.cells]
    for s in d.sections:
        for parte in (s.header, s.footer):
            partes += [p.text for p in parte.paragraphs]
            partes += [c.text for t in parte.tables for r in t.rows for c in r.cells]
    return '\n'.join(partes)


def _filas(ruta: Path) -> list[list[str]]:
    return [[c.text for c in r.cells] for t in docx.Document(ruta).tables for r in t.rows]


# --------------------------------------------------------------------------- RF-2, RF-3
def test_estudio_con_las_secciones_del_modelo_en_orden(tmp_path):
    """017:RF-2, 017:RF-3 — un .docx con las secciones de Casa Blanca en su orden."""
    destino = tmp_path / 'salida' / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path))
    texto = _texto(destino)
    orden = ['1. ANTECEDENTES', '2. DESCRIPCIÓN', '2.1 VISTA TOPOLÓGICA Y UBICACIÓN', 'EN ESTADO ACTUAL',
             '2.2.1.1 ANÁLISIS DE CAÍDA DE TENSIÓN', '2.2.1.2 ANÁLISIS DE CARGABILIDAD', 'ESTADO PROYECTADO',
             '2.4 DETALLE DE LAS OBRAS', 'PRIMERA ETAPA – AÑO 2027', 'ESCENARIO ACTUAL Y PROYECTADO',
             '3. CONCLUSIONES Y RECOMENDACIONES']
    posiciones = [texto.find(t) for t in orden]
    assert -1 not in posiciones and posiciones == sorted(posiciones), dict(zip(orden, posiciones))


def test_factibilidad_con_las_secciones_del_modelo_en_orden(tmp_path):
    """017:RF-2, 017:RF-3, 017:RF-106 — el orden del modelo EMAPICA."""
    destino = tmp_path / 'factibilidad.docx'
    escribir_informe_factibilidad(_config(tmp_path), destino, _factibilidad(tmp_path))
    texto = _texto(destino)
    orden = ['1. ANTECEDENTES', '2. DESCRIPCIÓN', '2.1 VISTA TOPOLÓGICA Y UBICACIÓN', 'EN ESTADO ACTUAL',
             '2.2.1.1 ANÁLISIS DE CAÍDA DE TENSIÓN', '2.2.1.2 ANÁLISIS DE CARGABILIDAD', '2.3 ESTADO PROYECTADO',
             'Datos en el punto de diseño', '2.3.2 CARGA PARA CONSIDERAR', '2.4 ANÁLISIS DE PÉRDIDAS TÉCNICAS',
             '2.5 ANÁLISIS DEL TRANSFORMADOR DE POTENCIA', '3. CONCLUSIONES']
    posiciones = [texto.find(t) for t in orden]
    assert -1 not in posiciones and posiciones == sorted(posiciones), dict(zip(orden, posiciones))


@pytest.mark.parametrize('tipo', ['estudio', 'factibilidad'])
def test_ningun_dato_del_modelo_pasa_al_informe(tmp_path, tipo):
    """017:RF-3, 017:RF-6 — esquema general: del modelo solo la forma; n.º, referencia y cuerpo son nuevos."""
    destino = tmp_path / f'{tipo}.docx'
    if tipo == 'estudio':
        escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path))
    else:
        escribir_informe_factibilidad(_config(tmp_path), destino, _factibilidad(tmp_path))
    texto = _texto(destino)
    for dato in DEL_MODELO:
        assert dato not in texto, dato
    d = docx.Document(destino)
    assert d.sections[0].header.tables[0].cell(0, 0).text.startswith('INFORME TECNICO Nº :\nGT-')
    assert d.sections[0].footer.tables[0].cell(0, 0).text == 'ELABORADO POR :'
    with zipfile.ZipFile(destino) as z:
        pie = next(z.read(n).decode() for n in z.namelist() if n.startswith('word/footer'))
    assert 'NUMPAGES' in pie


def test_sin_datos_del_operador_quedan_vacios(tmp_path):
    """017:RF-22, 017:RF-23 — n.º, asunto y destinatario sin dato: vacíos para rellenar, nunca los del modelo."""
    datos = _estudio(tmp_path, numero=None, antecedentes=(), descripcion=(),
                     encabezado=Encabezado(date(2026, 10, 15), '', ''))
    destino = tmp_path / 'informe.docx'
    huecos = escribir_informe_estudio(_config(tmp_path), destino, datos)
    cab = docx.Document(destino).sections[0].header.tables[0]
    assert cab.cell(0, 0).text == 'INFORME TECNICO Nº :\n' + ' ' * 24
    assert 'GT-019' not in _texto(destino)
    assert {'asunto del formulario', 'destinatario del formulario', 'antecedentes del formulario'} <= set(huecos)


def test_sin_rotulos_de_cabecera_en_el_perfil_no_se_toca(tmp_path):
    """017:RF-3 — sin los rótulos del perfil la cabecera no se reconoce y queda como está."""
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path, cabecera=None), destino, _estudio(tmp_path))
    assert docx.Document(destino).sections[0].header.tables[0].cell(0, 0).text.endswith('GT-019-PR')


# --------------------------------------------------------------------------- RF-6 a RF-9, RF-20
def test_nada_del_proyecto_ni_del_conversor(tmp_path):
    """017:RF-6, 017:RF-8, 017:RF-9, 017:RF-20 — ni casos de DIgSILENT, ni etapas evaluadas, ni recibo."""
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path))
    texto = _texto(destino)
    for palabra in ('Sin proyecto', 'Con proyecto', 'caso', 'variación', 'evaluada', 'sumada', 'RECIBO',
                    'CURVA', '[FALTA'):
        assert palabra not in texto, palabra


def test_las_imagenes_son_las_de_digsilent(tmp_path):
    """017:RF-7, 017:RF-10, 017:RF-46 — una figura por imagen exportada."""
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path))
    with zipfile.ZipFile(destino) as z:
        cuerpo = z.read('word/document.xml').decode('utf-8')
    # Ubicación, tensión y cargabilidad de los dos estados, y otra vez la ubicación actual
    # junto a la proyectada en el escenario actual y proyectado, como en Casa Blanca.
    assert cuerpo.count('<pic:pic') == 7


# --------------------------------------------------------------------------- RF-15 a RF-19, RF-118
def test_datos_de_la_fuente_y_tension_por_fase(tmp_path):
    """017:RF-15, 017:RF-17, 017:RF-118 — tabla de la fuente y tensión por fase, con coma decimal."""
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path))
    filas = _filas(destino)
    assert ['Potencia desde la fuente', '3000,00', '900,00', '3132,09', '95,78'] in filas
    assert ['Pérdidas en las líneas', '150,00', '120,00', '192,09', '78,09'] in filas
    assert ['A', '91,00 %', 'N0'] in filas and ['C', '93,00 %', 'N2'] in filas


def test_equilibrado_una_sola_tension_trifasica(tmp_path):
    """017:RF-18 — en el flujo equilibrado la tensión va una vez, como trifásica."""
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path, actual=_caso('x', 0.91)))
    filas = _filas(destino)
    assert ['Trifásica', '91,00 %', 'N0'] in filas and not any(f[0] == 'A' for f in filas)


def test_leyendas_y_norma(tmp_path):
    """017:RF-13, 017:RF-113, 017:RF-114 — leyendas con las bandas y la norma del perfil."""
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path))
    texto = _texto(destino)
    assert 'menos de 85 %' in texto and 'más de 105 %' in texto and 'de 105 a 150 %' in texto
    assert 'NTCSE' in texto and '±5 % de la tensión nominal' in texto


# --------------------------------------------------------------------------- RF-23, RF-24
def test_imagen_que_falta_queda_vacia_con_el_motivo_en_el_json(tmp_path):
    """017:RF-23 — sin la imagen, un espacio vacío; el motivo va al JSON y el resto sale."""
    img = _png(tmp_path / 'img.png')
    datos = _estudio(tmp_path, imagenes={'actual': {'tension': img, 'cargabilidad': img}})
    destino = tmp_path / 'informe.docx'
    huecos = escribir_informe_estudio(_config(tmp_path), destino, datos)
    assert any('ubicación' in h and 'estado actual' in h for h in huecos)
    assert json.loads(destino.with_suffix('.json').read_text(encoding='utf-8'))['huecos'] == huecos
    assert '3. CONCLUSIONES Y RECOMENDACIONES' in _texto(destino)


def test_estado_que_no_converge_sin_valores(tmp_path):
    """017:RF-24 — la sección dice que no converge y no trae tablas de valores."""
    etapas = [Etapa('bancos_nuevos', caso=_caso('x', 0, converge=False),
                    obras=({'tipo': 'banco', 'nodo': 'N1', 'kvar': 300},))]
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path, etapas=etapas))
    texto = _texto(destino)
    assert 'no converge' in texto and 'Ninguna de las obras evaluadas' in texto


# --------------------------------------------------------------------------- RF-25, RF-27
def test_json_con_los_valores_junto_al_word(tmp_path):
    """017:RF-25 — junto al .docx, un JSON con estados, obras, imágenes y huecos."""
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path))
    datos = json.loads(destino.with_suffix('.json').read_text(encoding='utf-8'))
    assert datos['recomendada'] == 'bancos_nuevos' and datos['obras_propuestas'] == ['bancos_nuevos']
    assert datos['imagenes']['actual']['tension'].endswith('img.png') and datos['huecos'] == []


def test_un_fallo_no_deja_word_a_medias(tmp_path, monkeypatch):
    """017:RF-27 — si la escritura falla, ni .docx ni .tmp en la carpeta."""
    import igea_dgs.informe_docx as modulo

    def revienta(*_a, **_k):
        raise RuntimeError('cancelado')
    monkeypatch.setattr(modulo, '_conclusiones', revienta)
    destino = tmp_path / 'salida' / 'informe.docx'
    with pytest.raises(RuntimeError):
        escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path))
    assert not (tmp_path / 'salida').exists() or not list((tmp_path / 'salida').glob('*'))


# --------------------------------------------------------------------------- RF-31, RF-44 a RF-46
def test_detalle_de_obras_por_etapa_con_su_anio(tmp_path):
    """017:RF-44, 017:RF-45 — por etapa: obras con nodo, kvar, km y transformadores, y su año."""
    etapas = [
        Etapa('bancos_nuevos', caso=_caso('b', 0.93), obras=({'tipo': 'banco', 'nodo': 'N15', 'kvar': 300},),
              anio=2026),
        Etapa('reguladores', caso=_caso('r', 0.94), obras=({'tipo': 'regulador', 'nodo': 'N40'},)),
        Etapa('conductor', caso=_caso('c', 0.94),
              obras=({'tipo': 'conductor', 'km': 4.25, 'conductor': 'AAAC 120 mm²'},)),
        Etapa('migracion', caso=_caso('m', 0.99),
              obras=({'tipo': 'migracion', 'kv': 22.9, 'transformadores': 37},), anio=2028),
    ]
    destino = tmp_path / 'informe.docx'
    datos = _estudio(tmp_path, etapas=etapas, imagenes=_imagenes(tmp_path, 'actual', 'migracion'))
    escribir_informe_estudio(_config(tmp_path), destino, datos)
    texto = _texto(destino)
    assert 'PRIMERA ETAPA – AÑO 2026' in texto and 'CUARTA ETAPA – AÑO 2028' in texto
    assert 'Instalar un banco de condensadores de 300 kvar en el nodo N15.' in texto
    assert 'regulador de tensión en el nodo N40' in texto
    assert '4,25 km de la troncal con conductor AAAC 120 mm²' in texto
    assert '37 transformadores de distribución' in texto and '22,9 kV' in texto


def test_estado_actual_que_cumple_sin_obras(tmp_path):
    """017:RF-31 — si el estado actual cumple, el informe lo dice y no trae obras."""
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path, actual=_caso('x', 0.97), etapas=[]))
    texto = _texto(destino)
    assert 'no se requieren obras de mejora' in texto and 'DETALLE DE LAS OBRAS' not in texto


# --------------------------------------------------------------------------- factibilidad
def test_factibilidad_punto_de_diseno_perdidas_y_transformador(tmp_path):
    """017:RF-50, 017:RF-52, 017:RF-59, 017:RF-62, 017:RF-120, 017:RF-122 — tablas y figura como el modelo."""
    destino = tmp_path / 'factibilidad.docx'
    escribir_informe_factibilidad(_config(tmp_path), destino, _factibilidad(tmp_path))
    filas = _filas(destino)
    texto = _texto(destino)
    assert ['0,970', '9,70', '5,60', '27,1', '455', '430', '142'] in filas
    assert ['AL02', 'ICA', 'SE0044', '1,20', '0,79'] in filas
    assert any(f[0] == 'ESTADO ACTUAL AL02' and f[1] == '2791,00' for f in filas)
    assert 'solicitudes que se encuentran vigentes' in texto
    assert 'reconocido por Osinergmin' in texto
    assert 'Figura 1: Cargabilidad de 73,87 % en el devanado de 10 kV de la SET SET1' in texto
    assert 'NMT 123' in texto and 'a) La solicitud' in texto


def test_factibilidad_sin_solicitud_queda_vacia(tmp_path):
    """017:RF-23 — sin los datos del cliente, sus secciones quedan vacías para rellenar."""
    destino = tmp_path / 'factibilidad.docx'
    datos = _factibilidad(tmp_path, solicitud=None, libro=None)
    huecos = escribir_informe_factibilidad(_config(tmp_path), destino, datos)
    assert 'datos de la solicitud del formulario' in huecos and 'libro de pérdidas de la empresa' in huecos
    assert 'EXP-1' not in _texto(destino) and '430 kW' not in _texto(destino)


# --------------------------------------------------------------------------- modelo
def test_modelo_con_imagenes_rotas_se_abre(tmp_path):
    """017:RF-4 — un modelo que Word dejó con imágenes apuntando a «NULL» se abre igual."""
    sano = _modelo(tmp_path / 'sano.docx')
    roto = tmp_path / 'roto.docx'
    with zipfile.ZipFile(sano) as z, zipfile.ZipFile(roto, 'w') as w:
        for i in z.infolist():
            b = z.read(i)
            if i.filename == 'word/_rels/document.xml.rels':
                b = b.replace(b'</Relationships>', b'<Relationship Id="rId99" Type="http://schemas.'
                              b'openxmlformats.org/officeDocument/2006/relationships/image" '
                              b'Target="../NULL"/></Relationships>')
            w.writestr(i, b)
    with pytest.raises(KeyError):
        docx.Document(roto)
    assert abrir_modelo(roto).sections[0].header.tables


def test_modelo_que_no_es_un_word(tmp_path):
    """017:RF-5 — un modelo que no se puede abrir es un error que lo nombra."""
    malo = tmp_path / 'malo.docx'
    malo.write_bytes(b'PK\x03\x04 truncado')
    with pytest.raises(InformeError, match='malo.docx'):
        abrir_modelo(malo)


def test_total_de_paginas_partido_en_dos_fragmentos(tmp_path):
    """017:RF-6 — «3» « de » «8» del pie de EMAPICA: el 8 fijo pasa a ser el campo de páginas."""
    config = _config(tmp_path)
    modelo = docx.Document(config.modelo('estudio'))
    celda = modelo.sections[0].footer.tables[0].cell(0, 1)
    celda.paragraphs[0].runs[1].text = ' de '
    celda.paragraphs[0].add_run('8')
    modelo.save(config.modelo('estudio'))
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(config, destino, _estudio(tmp_path))
    pie = docx.Document(destino).sections[0].footer.tables[0].cell(0, 1)
    assert '8' not in pie.text and 'NUMPAGES' in pie._tc.xml


def test_bancos_existentes_tal_como_estan_en_digsilent(tmp_path):
    """017:RF-36, 017:RF-37 — los bancos del alimentador, con nodo, kvar y estado, en el estado actual."""
    bancos = ({'nombre': 'BC1', 'barra': 'N7', 'kvar': 300.0, 'en_servicio': True},
              {'nombre': 'BC2', 'barra': 'N9', 'kvar': 150.0, 'en_servicio': False})
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path, bancos=bancos))
    filas = _filas(destino)
    assert ['BC1', 'N7', '300', 'En servicio'] in filas and ['BC2', 'N9', '150', 'Fuera de servicio'] in filas
    assert json.loads(destino.with_suffix('.json').read_text(encoding='utf-8'))['bancos'][0]['barra'] == 'N7'


def test_clientes_fijos_con_fp_real_y_corregido(tmp_path):
    """015:RF-65 — usuario, 2026-10-15: el informe modelo lleva el factor de potencia real
    (EA/ERI) y el corregido de cada cliente fijo, y dice cuáles se suponen compensados."""
    fijas = ({'sed': 'SE1', 'cliente': 'C1', 'p_kw': 100.0, 'fp_real': 0.43, 'fp_corregido': 0.98,
              'q_real_kvar': 210.0, 'q_kvar': 20.3, 'compensado': True},
             {'sed': 'SE2', 'cliente': 'C2', 'p_kw': 200.0, 'fp_real': 0.982, 'fp_corregido': 0.982,
              'q_real_kvar': 38.0, 'q_kvar': 38.0, 'compensado': False})
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, _estudio(tmp_path, fijas=fijas))
    filas = _filas(destino)
    assert ['SED', 'Cliente', 'P (kW)', 'FP real', 'FP corregido', 'Q real (kvar)',
            'Q modelada (kvar)'] in filas
    assert ['SE1', 'C1', '100,0', '0,430', '0,980', '210,0', '20,3'] in filas \
        or ['SE1', 'C1', '100.0', '0.430', '0.980', '210.0', '20.3'] in filas
    texto = '\n'.join(p.text for p in __import__('docx').Document(destino).paragraphs)
    assert 'banco de condensadores propio' in texto and 'SE1' in texto
    assert json.loads(destino.with_suffix('.json').read_text(encoding='utf-8'))['fijas'][0]['sed'] == 'SE1'
