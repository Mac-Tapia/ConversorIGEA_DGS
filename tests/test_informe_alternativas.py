"""Spec 017, RF-128 a RF-135: alternativas de obras en el informe de estudio.

El usuario, 2026-10-16, sobre el informe de CA101: primero la alternativa de solo los
bancos, después el cambio de sección con el tramo desde–hasta y su plano, y los bancos
existentes con el nodo donde están.
"""
from __future__ import annotations

import random
from datetime import date
from pathlib import Path

import docx
from PIL import Image

import fake_pf
from apoyo import cargar_herramienta, contexto_nativo
from igea_dgs.informe_docx import DatosEstudio, escribir_informe_estudio
from igea_dgs.informes import (
    Caso, Encabezado, Etapa, Extremo, Potencia, alternativas_de_obras, describir_obra,
    trechos_de_conductor)
from igea_dgs.mapa_informe import COLOR_BANCO, COLOR_CONDUCTOR, dibujar

casos_pf = cargar_herramienta('casos_pf.py', 'casos_alt', tools_en_path=True)
guion = cargar_herramienta('estudios_pf.py', 'estudios_alt', tools_en_path=True)


# --------------------------------------------------------------------------- RF-128, RF-129
def _aplicada(variacion, clase):
    return {'variacion': variacion, 'obra': {'clase': clase, 'alimentador': 'AL01'}}


def test_una_alternativa_por_etapa_sumada_a_las_anteriores():
    """017:RF-128 — bancos primero y conductor después, cada una con las anteriores."""
    alts = alternativas_de_obras([_aplicada('Obra_seccion', 'seccion'),
                                  _aplicada('Obra_condensador', 'condensador'),
                                  _aplicada('Obra_condensador_2', 'condensador')])
    assert alts == [
        {'clave': 'bancos_nuevos', 'variaciones': ['Obra_condensador', 'Obra_condensador_2']},
        {'clave': 'conductor', 'variaciones': ['Obra_condensador', 'Obra_condensador_2', 'Obra_seccion']}]


def test_la_consigna_no_es_alternativa_propia():
    """017:RF-129 — la consigna va con la alternativa siguiente (CA101), o con la última."""
    ca101 = alternativas_de_obras([_aplicada('C1', 'condensador'), _aplicada('C2', 'condensador'),
                                   _aplicada('K', 'consigna'), _aplicada('S', 'seccion')])
    assert [a['clave'] for a in ca101] == ['bancos_nuevos', 'conductor']
    assert ca101[0]['variaciones'] == ['C1', 'C2'] and 'K' in ca101[1]['variaciones']
    solo_bancos = alternativas_de_obras([_aplicada('C1', 'condensador'), _aplicada('K', 'consigna')])
    assert solo_bancos == [{'clave': 'bancos_nuevos', 'variaciones': ['C1', 'K']}]
    assert alternativas_de_obras([]) == []


# --------------------------------------------------------------------------- RF-130, RF-131
def _variacion_que_cambia(app, nombre, carga, instante):
    """Variación de obra en el caso activo que deja una carga fuera de servicio.

    Cada etapa con su fecha: con la misma, DIgSILENT no graba (memoria del proyecto).
    """
    v = app.carpetas['scheme'].CreateObject('IntScheme', nombre)
    v.NewStage(f'Etapa_{nombre}', instante, 1)
    app.grabando = v.GetContents('*.IntSstage', 1)[0]
    carga.outserv = 1
    app.grabando = None
    return v


def _red():
    alims = ('F1',)
    app = fake_pf.proyecto(alimentadores=alims, cargas_por_alim=3, bt=True, flujo=True)
    app.atributos_en_variacion = True
    v = app.carpetas['scheme'].CreateObject('IntScheme', 'Ajuste_F1')
    v.NewStage('Etapa_ajuste', 0, 1)
    app.grabando = None
    base = {'variaciones': ['Ajuste_F1'], 'feeders': list(alims),
            'ajustes': [{'feeder': 'F1', 'variacion': 'Ajuste_F1', 'cuadrado': True}], 'excluidos': []}
    informe, codigo = casos_pf.preparar(app, {'project': 'P', 'feeders': list(alims), 'base': base,
                                              'caso_base_sin_verificar': {}})
    assert codigo == 0, informe
    for barra in app.proyecto.GetContents('*.ElmTerm', 1):
        barra.GPSlat, barra.GPSlon = -14.0 - random.random() / 100, -75.7 + random.random() / 100
    return app


def test_cada_alternativa_se_mide_con_solo_sus_obras():
    """017:RF-130 — en la misma ejecución, cada alternativa con sus variaciones y su
    geografía; al terminar siguen activas las que estaban."""
    app = _red()
    casos = casos_pf.casos_del_proyecto(app)
    casos['con'].Activate()
    cargas = app.proyecto.GetContents('*.ElmLod', 1)
    _variacion_que_cambia(app, 'Obra_condensador_F1', cargas[0], 1)
    _variacion_que_cambia(app, 'Obra_seccion_F1', cargas[1], 2)
    plan = {'project': 'P', 'feeders': ['F1'], 'estudio': 'flujo', 'caso': 'con',
            'stamp': '20261016_120000', 'native_action_context': contexto_nativo('P', ['F1'], 'estudios'),
            'alternativas': [
                {'clave': 'bancos_nuevos', 'variaciones': ['Obra_condensador_F1']},
                {'clave': 'conductor', 'variaciones': ['Obra_condensador_F1', 'Obra_seccion_F1']},
                {'clave': 'migracion', 'variaciones': ['Obra_migracion_F1']}]}
    informe, rc = guion.ejecutar(app, plan)
    assert rc == 0, informe
    alts = {a['clave']: a for a in informe['alternativas']}
    assert alts['bancos_nuevos']['variaciones'] == ['Obra_condensador_F1']
    kw = {k: a['medidas']['F1']['cabecera_kw'] for k, a in alts.items()}
    assert kw['bancos_nuevos'] > kw['conductor']
    assert abs(kw['conductor'] - informe['medidas']['F1']['cabecera_kw']) < 1e-6
    assert alts['bancos_nuevos']['geo']['F1']['tramos']
    assert alts['migracion']['faltan'] == ['Obra_migracion_F1']
    assert any('Obra_migracion_F1' in a for a in informe['avisos'])
    casos['con'].Activate()
    assert {'Obra_condensador_F1', 'Obra_seccion_F1'} <= set(guion.variaciones_activas(app))


def test_la_geografia_de_cada_tramo_trae_barras_tipo_y_longitud():
    """017:RF-131 — ``[nombre, pts, carga, barra1, barra2, tipo, km]``."""
    app = _red()
    informe, rc = guion.ejecutar(app, {'project': 'P', 'feeders': ['F1'], 'estudio': 'flujo',
                                       'caso': 'sin', 'stamp': 's',
                                       'native_action_context': contexto_nativo('P', ['F1'], 'estudios')})
    assert rc == 0, informe
    tramo = informe['geo']['F1']['tramos'][0]
    assert len(tramo) == 7 and tramo[3] and tramo[4] and tramo[5] == 'AAAC_F1'


# --------------------------------------------------------------------------- RF-132
#: Cadena CAB–N1–N2–N3–N4 y un ramal N2–R1.
GEO = {'cabecera': 'CAB',
       'barras': [['CAB', -14.000, -75.700, 1.0, 10.0], ['N1', -14.001, -75.700, 0.99, 10.0],
                  ['N2', -14.002, -75.700, 0.98, 10.0], ['N3', -14.003, -75.700, 0.97, 10.0],
                  ['N4', -14.004, -75.700, 0.96, 10.0], ['R1', -14.002, -75.699, 0.97, 10.0]],
       'tramos': [['L1', [[-14.000, -75.7], [-14.001, -75.7]], 50.0, 'CAB', 'N1', 'AA120', 0.11],
                  ['L2', [[-14.001, -75.7], [-14.002, -75.7]], 50.0, 'N1', 'N2', 'AA120', 0.11],
                  ['L3', [[-14.002, -75.7], [-14.003, -75.7]], 50.0, 'N3', 'N2', 'AA120', 0.11],
                  ['L4', [[-14.003, -75.7], [-14.004, -75.7]], 50.0, 'N3', 'N4', 'AA120', 0.11],
                  ['L5', [[-14.002, -75.7], [-14.002, -75.699]], 20.0, 'N2', 'R1', 'AA35', 0.11]]}
GEO_SIN = {**GEO, 'tramos': [t[:5] + ['AA35', t[6]] for t in GEO['tramos']]}


def test_trechos_desde_hasta_de_la_cabecera_a_la_cola():
    """017:RF-132 — trechos continuos con inicio, fin, tramos, km y conductor actual."""
    trechos = trechos_de_conductor(GEO, ['L4', 'L1', 'L3', 'L2'], GEO_SIN)
    assert len(trechos) == 1
    t = trechos[0]
    assert (t.desde, t.hasta) == ('CAB', 'N4')
    assert t.tramos == ('L1', 'L2', 'L3', 'L4')
    assert abs(t.km - 0.44) < 1e-9 and t.actuales == ('AA35',)
    assert t.desde_gps == (-14.0, -75.7)
    dos = trechos_de_conductor(GEO, ['L1', 'L4'], GEO_SIN)
    assert [(x.desde, x.hasta) for x in dos] == [('CAB', 'N1'), ('N3', 'N4')]
    assert trechos_de_conductor({'tramos': [['L1', [], 1.0]]}, ['L1']) == []


def test_un_acople_no_corta_el_trecho():
    """017:RF-132 — un seccionador entre dos tramos (CA101: 10 trechos falsos) los deja
    en un solo trecho, y el punto se nombra por su barra propia, no por «Terminal»."""
    tramos = [['L1', [[-14.0, -75.7], [-14.001, -75.7]], 50.0, 'CAB', 'N1', 'AA35', 0.1],
              ['L2', [[-14.0011, -75.7], [-14.002, -75.7]], 50.0, 'N1b', 'Terminal', 'AA35', 0.1],
              ['L3', [[-14.002, -75.7], [-14.003, -75.7]], 50.0, 'Terminal', 'N3', 'AA35', 0.1]]
    geo = {'cabecera': 'CAB', 'tramos': tramos, 'acoples': [['S1', 'N1', 'N1b']],
           'barras': [['CAB', -14.0, -75.7, 1, 10], ['N1', -14.001, -75.7, 1, 10],
                      ['N3', -14.003, -75.7, 1, 10]]}
    trechos = trechos_de_conductor(geo, ['L1', 'L2', 'L3'], geo)
    assert [(t.desde, t.hasta, t.tramos) for t in trechos] == [('CAB', 'N3', ('L1', 'L2', 'L3'))]
    sin_acople = trechos_de_conductor({**geo, 'acoples': []}, ['L1', 'L2', 'L3'], geo)
    assert len(sin_acople) == 2


def test_la_obra_de_conductor_dice_desde_donde_hasta_donde():
    """017:RF-132 — la descripción nombra el inicio, el fin y el conductor de hoy."""
    from dataclasses import asdict
    trechos = [asdict(t) for t in trechos_de_conductor(GEO, ['L1', 'L2'], GEO_SIN)]
    texto = describir_obra({'tipo': 'conductor', 'km': 0.22, 'conductor': 'AA120', 'trechos': trechos})
    assert 'desde la barra CAB hasta la barra N2' in texto and 'hoy con AA35' in texto


# --------------------------------------------------------------------------- RF-133
def test_el_plano_resalta_el_conductor_nuevo_y_los_bancos(tmp_path):
    """017:RF-133 — tramos cambiados en trazo grueso, inicio y fin, bancos señalados."""
    ruta = dibujar(GEO, 'ubicacion', tmp_path / 'p.png', titulo='Alternativa 2',
                   obras={'tramos': ['L1', 'L2'], 'trechos': [('CAB', 'N2')], 'bancos': [('N3', 300.0)]})
    with Image.open(ruta) as img:
        colores = {c for _, c in img.getcolors(1 << 20)}
    assert COLOR_CONDUCTOR in colores and COLOR_BANCO in colores
    ruta = dibujar(GEO, 'ubicacion', tmp_path / 'q.png', titulo='Actual')
    with Image.open(ruta) as img:
        assert COLOR_CONDUCTOR not in {c for _, c in img.getcolors(1 << 20)}


# --------------------------------------------------------------------------- RF-134, RF-135
def _modelo(ruta: Path) -> Path:
    d = docx.Document()
    s = d.sections[0]
    s.header.add_table(1, 2, width=s.page_width).cell(0, 0).text = 'INFORME TECNICO Nº :\nX'
    d.save(ruta)
    return ruta


def _caso(minima, carga=60.0):
    return Caso('Con proyecto', True, fuente=Potencia(3000, 900), perdidas=Potencia(150, 120),
                tension_min={'trifásica': Extremo(minima, 'N4')}, tension_max=Extremo(1.0, 'CAB'),
                nodos_fuera=0 if minima >= 0.95 else 12, cargabilidad_max=Extremo(carga, 'L1'))


def _config(tmp_path):
    from igea_dgs.informes import ConfigInforme
    perfil = {'estudios': {'tension_min_pu': 0.95, 'tension_max_pu': 1.05, 'cargabilidad_max_pct': 100},
              'informes': {'modelo_estudio': str(_modelo(tmp_path / 'm.docx')),
                           'modelo_factibilidad': str(_modelo(tmp_path / 'f.docx')),
                           'bandas_tension_pct': [85, 90, 95, 105],
                           'bandas_cargabilidad_pct': [80, 90, 95, 105, 150]}}
    return ConfigInforme.de_perfil(perfil, tmp_path)


def _texto(ruta):
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(ruta)
    partes = []
    for el in d.element.body.iterchildren():
        if el.tag.endswith('}p'):
            partes.append(Paragraph(el, d).text)
        elif el.tag.endswith('}tbl'):
            partes += [' | '.join(c.text for c in r.cells) for r in Table(el, d).rows]
    return '\n'.join(partes)


def test_informe_con_dos_alternativas_cuadro_y_recomendacion(tmp_path):
    """017:RF-128, 017:RF-132, 017:RF-134, 017:RF-135 — sección por alternativa, trecho
    desde–hasta, cuadro comparativo, la primera que cumple recomendada y los bancos
    existentes con su nodo."""
    from dataclasses import asdict
    trechos = [{**asdict(t), 'desde_utm': [400000.0, 8450000.0], 'hasta_utm': [400100.0, 8449600.0]}
               for t in trechos_de_conductor(GEO, ['L1', 'L2', 'L3', 'L4'], GEO_SIN)]
    etapas = [Etapa('bancos_nuevos', caso=_caso(0.93), obras=({'tipo': 'banco', 'nodo': 'N3', 'kvar': 300},)),
              Etapa('conductor', caso=_caso(0.98), obras=(
                  {'tipo': 'conductor', 'km': 0.44, 'conductor': 'AA120', 'trechos': trechos},))]
    png = tmp_path / 'p.png'
    Image.new('RGB', (4, 4), 'white').save(png)
    imagenes = {k: {'ubicacion': png} for k in ('bancos_nuevos', 'conductor')}
    bancos = ({'nombre': 'Shunt/Filter', 'barra': 'Terminal', 'nodo': 'SE20701', 'kvar': 300.0,
               'en_servicio': True, 'utm': [401000.0, 8451000.0], 'sed_cercana': ['SE20701', 0]},)
    datos = DatosEstudio(alimentadores=['AL01'], encabezado=Encabezado(date(2026, 10, 16), 'A', 'B'),
                         actual=Caso('Sin proyecto', True, fuente=Potencia(3100, 1500),
                                     tension_min={'trifásica': Extremo(0.90, 'N4')}, nodos_fuera=20),
                         etapas=etapas, bancos=bancos, imagenes=imagenes)
    destino = tmp_path / 'informe.docx'
    escribir_informe_estudio(_config(tmp_path), destino, datos)
    texto = _texto(destino)
    orden = ['2.3 ALTERNATIVAS DE SOLUCIÓN', '2.3.1 ALTERNATIVA 1: BANCOS DE CONDENSADORES PROPUESTOS',
             'PLANO DE LA ALTERNATIVA 1', '2.3.2 ALTERNATIVA 2: CAMBIO DE SECCIÓN DEL CONDUCTOR',
             'Incluye las obras de la alternativa 1', 'PLANO DE LA ALTERNATIVA 2', '2.4 CUADRO COMPARATIVO',
             'Se recomienda la alternativa 2', '2.5 DETALLE DE LAS OBRAS', '3. CONCLUSIONES']
    posiciones = [texto.find(t) for t in orden]
    assert -1 not in posiciones and posiciones == sorted(posiciones), dict(zip(orden, posiciones))
    assert '1 | CAB | E 400000,0 / N 8450000,0 | N4 |' in texto
    assert 'Tramos del trecho 1, de la cabecera hacia la cola: L1, L2, L3, L4.' in texto
    assert 'Alternativa 1 | 0,930 | N4' in texto and 'Alternativa 2 | 0,980 | N4' in texto
    assert 'Shunt/Filter | SE20701 | Terminal | 300 | E 401000,0 / N 8451000,0 | En servicio' in texto
    assert 'Se mantiene en servicio el banco existente: 300 kvar en el nodo SE20701.' in texto
    # 017:RF-133: el plano de solo bancos no habla de conductor nuevo.
    assert ('Plano de la alternativa 1: se señalan los bancos de condensadores nuevos (triángulo lleno) '
            'y los bancos existentes (triángulo hueco).') in texto
    assert 'Plano de la alternativa 2: se señalan los tramos con el conductor nuevo en trazo grueso' in texto


class _Nodo:
    """Barra o elemento mínimo para recorrer la red desde un banco."""

    def __init__(self, nombre, clase='ElmTerm', padre=None):
        self.loc_name, self._clase, self._padre, self.vecinos = nombre, clase, padre, []

    def GetClassName(self):
        return self._clase

    def GetParent(self):
        return self._padre

    def GetConnectedElements(self, *_a):
        return self.vecinos


def test_el_banco_se_nombra_por_su_nodo_de_la_red():
    """017:RF-135 — una barra «Terminal» se nombra por la subestación o el nodo vecino."""
    assert guion._nodo_de_banco(_Nodo('SE20730(1)'))[0] == 'SE20730(1)'
    assert guion._nodo_de_banco(_Nodo('Terminal', padre=_Nodo('SE20701', 'ElmTrfstat')))[0] == 'SE20701'
    aux, red, linea = _Nodo('Terminal(1)', padre=_Nodo('Red', 'ElmNet')), _Nodo('NODE_42'), _Nodo('L', 'ElmLne')
    aux.vecinos, linea.vecinos = [linea], [aux, red]
    assert guion._nodo_de_banco(aux)[0] == 'NODE_42'
    assert guion._nodo_de_banco(None) == (None, None)


def test_el_nodo_numerico_dice_la_sed_mas_cercana(tmp_path):
    """017:RF-135 — «14717» con la SED a su lado; dos bancos del mismo nodo, juntos."""
    from igea_dgs.informe_docx import _nodo_banco, _siguen_en_servicio
    from igea_dgs.web.services import _sed_cercana
    geo = {'barras': [['SE20730_BT', -13.4693, -76.01466, 0.9, 0.38], ['14717', -13.4694, -76.01466, 0.9, 10.0]]}
    cerca = _sed_cercana([-13.46937, -76.01466], geo)
    assert cerca[0] == 'SE20730' and cerca[1] < 20
    banco = {'nodo': '14717', 'kvar': 300.0, 'en_servicio': True, 'sed_cercana': cerca}
    assert _nodo_banco(banco).startswith('14717 (a ') and 'de la SED SE20730' in _nodo_banco(banco)
    assert _siguen_en_servicio([banco, dict(banco)]).endswith('2 de 300 kvar en el nodo ' + _nodo_banco(banco) + '.')
