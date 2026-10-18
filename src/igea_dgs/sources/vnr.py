"""Adaptador de un paquete VNR-GIS custodiado."""

from __future__ import annotations

from pathlib import Path

from .base import SourceLoadResult, provenance, validate_snapshot
from .snapshot import SourceSnapshot


class VnrSourceAdapter:
    mode = 'vnr'

    def __init__(self, *, company: str | None = None, period: str | None = None) -> None:
        self.company = company
        self.period = period

    def load(self, snapshot: SourceSnapshot, *, aliases: dict[str, str]) -> SourceLoadResult:
        del aliases
        validate_snapshot(
            snapshot,
            mode=self.mode,
            allowed={'vnr_package'},
            required={'vnr_package'},
        )
        from vnr_etl.application.cymdist_bridge import canonical_to_cymdist
        from vnr_etl.pipeline import canonicalize_vnr_package, safe_extract_package

        package = snapshot.files['vnr_package'].path
        if Path(package).suffix.lower() == '.bak':
            return self._cargar_copia_sql(snapshot, Path(package))
        extracted = safe_extract_package(package, snapshot.root / 'derivados' / 'vnr')
        canonical, scope = canonicalize_vnr_package(
            extracted.destination,
            company=self.company,
            period=self.period,
        )
        bridge = canonical_to_cymdist(canonical.model, source_path=package)
        source = provenance(snapshot)
        source.update({
            'company': scope['company'],
            'period': scope['period'],
            'package_format': extracted.format,
            'extracted_files': len(extracted.files),
            'canonical_counts': canonical.model.counts(),
            'bridge_issues': bridge.issues,
        })
        return SourceLoadResult(
            dataset=bridge.dataset,
            catalog_report=None,
            readiness=bridge.readiness,
            provenance=source,
        )

    def _cargar_copia_sql(self, snapshot: SourceSnapshot, bak: Path) -> SourceLoadResult:
        """Spec 023: la copia de SQL Server de VNR-GIS de cualquier concesionaria."""
        import os

        from vnr_etl.application.cymdist_bridge import canonical_to_cymdist
        from vnr_etl.connectors.sqlserver_vnr import SERVIDOR, cargar_copia

        servidor = os.environ.get('IGEA_SQLSERVER', SERVIDOR)
        original = Path(snapshot.files['vnr_package'].source_path or bak)
        modelo, avisos, alcance = cargar_copia(bak, servidor=servidor,
                                               vecinas=[original.parent])
        fichas = fichas_del_catalogo(modelo)
        bridge = canonical_to_cymdist(modelo, source_path=bak)
        source = provenance(snapshot)
        source.update({
            'company': alcance['company'],
            'period': alcance['period'],
            'package_format': 'sqlserver_bak',
            'sqlserver': alcance,
            'canonical_counts': modelo.counts(),
            'vnr_avisos': avisos,
            'fichas_del_catalogo': fichas,
            'bridge_issues': bridge.issues,
        })
        return SourceLoadResult(
            dataset=bridge.dataset,
            catalog_report=None,
            readiness=bridge.readiness,
            provenance=source,
        )


def fichas_del_catalogo(modelo, catalogo: Path | None = None) -> dict[str, int]:
    """R, X y amperios de cada tipo de conductor del modelo, del catálogo del proyecto (023:RF-10).

    VNR-GIS no trae parámetros eléctricos, solo la norma (material y sección). El catálogo
    de parámetros del proyecto los tiene por el mismo código: la ficha si su estado lo
    permite y, si no, el valor del modelo de la empresa. Sin catálogo, el resolvedor sigue
    buscando por material y sección. Devuelve cuántos tipos se completaron y cuántos no.
    """
    from ..catalog import ESTADOS_QUE_CORRIGEN
    from ..reglas import catalogo_del_proyecto

    ruta = catalogo or catalogo_del_proyecto()
    tipos = list(getattr(modelo, 'conductor_types', []) or [])
    if ruta is None or not tipos:
        return {'completados': 0, 'sin_ficha': len(tipos)}
    from openpyxl import load_workbook

    filas: dict[str, dict] = {}
    wb = load_workbook(ruta, read_only=True, data_only=True)
    for hoja in ('conductores_aereos', 'cables_subterraneos'):
        if hoja not in wb.sheetnames:
            continue
        valores = list(wb[hoja].iter_rows(values_only=True))
        if not valores:
            continue
        cabecera = [str(c or '') for c in valores[0]]
        for v in valores[1:]:
            fila = dict(zip(cabecera, v))
            if fila.get('codigo'):
                filas[str(fila['codigo']).strip()] = fila

    def valor(fila, nombre, unidad):
        ficha = fila.get(f'{nombre}_ficha_{unidad}')
        if ficha is not None and str(fila.get('estado') or '') in ESTADOS_QUE_CORRIGEN:
            return float(ficha)
        modelo_ = fila.get(f'{nombre}_modelo_{unidad}')
        return None if modelo_ is None else float(modelo_)

    # R, X y amperios dependen del material y la sección, no de la tensión ni de las fases:
    # en la norma VNR son sus cinco primeros caracteres («AA050», «CU016», «N2035»). El
    # catálogo de Electro Dunas pone «D» también a las redes de 22,9 kV, así que «AA05003G»
    # (o el monofásico «AA05001D») toma la ficha de «AA05003D» (Ica, 2026-10-18).
    por_familia: dict[str, dict] = {}
    for codigo in sorted(filas):
        por_familia.setdefault(codigo[:5], filas[codigo])
    completados = 0
    for tipo in tipos:
        codigo = tipo.conductor_code
        fila = filas.get(codigo) or filas.get(codigo[:-1]) or por_familia.get(codigo[:5])
        if fila is None:
            continue
        tipo.r1_ohm_km = valor(fila, 'R1', 'ohm_km')
        tipo.x1_ohm_km = valor(fila, 'X1', 'ohm_km')
        tipo.r0_ohm_km = valor(fila, 'R0', 'ohm_km')
        tipo.x0_ohm_km = valor(fila, 'X0', 'ohm_km')
        tipo.ampacity_a = valor(fila, 'ampacidad', 'A')
        completados += all(x is not None for x in (tipo.r1_ohm_km, tipo.x1_ohm_km,
                                                   tipo.r0_ohm_km, tipo.x0_ohm_km))
    return {'completados': completados, 'sin_ficha': len(tipos) - completados}
